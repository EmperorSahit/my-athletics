import sys
import sqlite3
import os
import shutil

from datetime import datetime

from PySide6.QtWidgets import (
    QApplication,
    QMainWindow,
    QWidget,
    QLabel,
    QPushButton,
    QVBoxLayout,
    QHBoxLayout,
    QStackedWidget,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMessageBox,
    QComboBox,
    QDateEdit,
    QTableWidget,
    QTableWidgetItem,
    QInputDialog,
    QFileDialog
)

from PySide6.QtCore import Qt, QDate

from database import (
    init_db,
    set_active_meeting,
    get_active_meeting,
    set_user_mode,
    get_user_mode
)

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from zipfile import BadZipFile

DB_NAME = "my_athletics.db"

# =========================
# RESOURCE PATH
# =========================
def resource_path(relative_path):

    # When packaged as an EXE,
    # PyInstaller extracts bundled files
    # into a temporary folder.
    if hasattr(sys, "_MEIPASS"):

        base_path = sys._MEIPASS

    else:

        base_path = os.path.dirname(
            os.path.abspath(__file__)
        )

    return os.path.join(
        base_path,
        relative_path
    )

# =========================
# CREATE MEETING SCREEN
# =========================
class CreateMeetingScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        title = QLabel("Create Meeting")
        title.setAlignment(Qt.AlignCenter)

        self.name_input = QLineEdit()
        self.name_input.setPlaceholderText("Meeting Name")
        self.date_input = QDateEdit()
        self.date_input.setDate(QDate.currentDate())
        self.date_input.setCalendarPopup(True)

        self.team_mode = QComboBox()
        self.team_mode.addItems([
            "School",
            "House"
        ])

        save_button = QPushButton("Save Meeting")
        save_button.clicked.connect(self.save_meeting)

        layout.addWidget(title)
        layout.addWidget(self.name_input)
        layout.addWidget(self.date_input)
        layout.addWidget(self.team_mode)
        layout.addWidget(save_button)
        layout.addStretch()

        self.setLayout(layout)

    def save_meeting(self):

        name = self.name_input.text().strip()

        if not name:
            QMessageBox.warning(
                self,
                "Missing Name",
                "Enter a meeting name."
            )
            return

        meeting_date = self.date_input.date().toString(
            "yyyy-MM-dd"
        )

        team_mode = self.team_mode.currentText()

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        INSERT INTO meetings
        (name, date, team_mode, status)
        VALUES (?, ?, ?, 'Draft')
        """, (
            name,
            meeting_date,
            team_mode
        ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Saved",
            "Meeting created successfully."
        )

        self.name_input.clear()

# =========================
# LOAD MEETING SCREEN
# =========================
class LoadMeetingScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        title = QLabel("Load Meeting")
        title.setAlignment(Qt.AlignCenter)

        self.list_widget = QListWidget()

        refresh_button = QPushButton("Refresh")
        refresh_button.clicked.connect(self.load_meetings)

        activate_button = QPushButton("Set Active Meeting")
        activate_button.clicked.connect(self.activate_meeting)

        layout.addWidget(title)
        layout.addWidget(self.list_widget)
        layout.addWidget(refresh_button)
        layout.addWidget(activate_button)

        self.setLayout(layout)

        self.load_meetings()

    def load_meetings(self):

        self.list_widget.clear()

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT id, name, date, team_mode
        FROM meetings
        ORDER BY id DESC
        """)

        rows = cur.fetchall()

        conn.close()

        for meeting_id, name, date, mode in rows:
            self.list_widget.addItem(
                f"{meeting_id} - {name} | {date} | {mode}"
            )
    def activate_meeting(self):

        item = self.list_widget.currentItem()

        if not item:
            QMessageBox.warning(
                self,
                "Select Meeting",
                "Choose a meeting first."
            )
            return

        meeting_id = int(
            item.text().split(" - ")[0]
        )

        set_active_meeting(meeting_id)

        QMessageBox.information(
            self,
            "Active Meeting",
            f"Meeting {meeting_id} activated."
        )

# =========================
# TEAM MANAGEMENT SCREEN
# =========================
class TeamManagementScreen(QWidget):

    def __init__(self):
        super().__init__()

        self.selected_team_id = None

        layout = QVBoxLayout()

        title = QLabel("Team Management")
        title.setAlignment(Qt.AlignCenter)

        self.team_name = QLineEdit()
        self.team_name.setPlaceholderText("Team Name")

        self.team_type = QComboBox()
        self.team_type.addItems([
            "School",
            "House"
        ])

        add_button = QPushButton(
            "Add Team"
        )
        add_button.clicked.connect(
            self.add_team
        )

        update_button = QPushButton(
            "Update Selected Team"
        )
        update_button.clicked.connect(
            self.update_team
        )

        delete_button = QPushButton(
            "Delete Selected Team"
        )
        delete_button.clicked.connect(
            self.delete_team
        )

        clear_button = QPushButton(
            "Clear Selection"
        )
        clear_button.clicked.connect(
            self.clear_selection
        )

        self.team_list = QListWidget()

        self.team_list.itemClicked.connect(
            self.load_selected_team
        )

        refresh_button = QPushButton(
            "Refresh Teams"
        )
        refresh_button.clicked.connect(
            self.load_teams
        )

        layout.addWidget(title)

        layout.addWidget(
            QLabel("Team / House Name")
        )
        layout.addWidget(self.team_name)

        layout.addWidget(
            QLabel("Type")
        )
        layout.addWidget(self.team_type)

        layout.addWidget(add_button)
        layout.addWidget(update_button)
        layout.addWidget(delete_button)
        layout.addWidget(clear_button)

        layout.addWidget(
            QLabel("Existing Teams / Houses")
        )

        layout.addWidget(self.team_list)
        layout.addWidget(refresh_button)

        self.setLayout(layout)

        self.load_teams()


    # =========================
    # ADD TEAM
    # =========================
    def add_team(self):

        meeting_id = get_active_meeting()

        if not meeting_id:

            QMessageBox.warning(
                self,
                "No Active Meeting",
                "Select an active meeting first."
            )

            return

        name = self.team_name.text().strip()

        if not name:

            QMessageBox.warning(
                self,
                "Missing Name",
                "Enter a team or house name."
            )

            return

        team_type = (
            self.team_type.currentText()
        )

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Prevent duplicate team names
        # inside the same meeting
        cur.execute("""
        SELECT id
        FROM teams
        WHERE meeting_id = ?
        AND LOWER(name) = LOWER(?)
        """, (
            meeting_id,
            name
        ))

        if cur.fetchone():

            conn.close()

            QMessageBox.warning(
                self,
                "Duplicate Team",
                (
                    "A team or house with this "
                    "name already exists."
                )
            )

            return

        cur.execute("""
        INSERT INTO teams
        (
            meeting_id,
            name,
            team_type
        )
        VALUES (?, ?, ?)
        """, (
            meeting_id,
            name,
            team_type
        ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Team Added",
            "Team or house added successfully."
        )

        self.clear_selection()
        self.load_teams()


    # =========================
    # LOAD TEAMS
    # =========================
    def load_teams(self):

        self.team_list.clear()

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            id,
            name,
            team_type
        FROM teams
        WHERE meeting_id = ?
        ORDER BY name
        """, (meeting_id,))

        rows = cur.fetchall()

        conn.close()

        for (
            team_id,
            name,
            team_type
        ) in rows:

            item = QListWidgetItem(
                f"{name} ({team_type})"
            )

            item.setData(
                Qt.UserRole,
                team_id
            )

            self.team_list.addItem(
                item
            )


    # =========================
    # LOAD SELECTED TEAM
    # =========================
    def load_selected_team(
        self,
        item
    ):

        team_id = item.data(
            Qt.UserRole
        )

        if not team_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            name,
            team_type
        FROM teams
        WHERE id = ?
        """, (team_id,))

        team = cur.fetchone()

        conn.close()

        if not team:
            return

        self.selected_team_id = team_id

        name = team[0]
        team_type = team[1]

        self.team_name.setText(
            name
        )

        type_index = (
            self.team_type.findText(
                team_type
            )
        )

        if type_index >= 0:

            self.team_type.setCurrentIndex(
                type_index
            )


    # =========================
    # UPDATE TEAM
    # =========================
    def update_team(self):

        if not self.selected_team_id:

            QMessageBox.warning(
                self,
                "No Team Selected",
                (
                    "Select a team or house "
                    "from the list first."
                )
            )

            return

        name = self.team_name.text().strip()

        if not name:

            QMessageBox.warning(
                self,
                "Missing Name",
                "Enter a team or house name."
            )

            return

        meeting_id = get_active_meeting()

        team_type = (
            self.team_type.currentText()
        )

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Make sure another team does
        # not already use this name.
        cur.execute("""
        SELECT id
        FROM teams
        WHERE meeting_id = ?
        AND LOWER(name) = LOWER(?)
        AND id != ?
        """, (
            meeting_id,
            name,
            self.selected_team_id
        ))

        if cur.fetchone():

            conn.close()

            QMessageBox.warning(
                self,
                "Duplicate Team",
                (
                    "Another team or house "
                    "already uses this name."
                )
            )

            return

        cur.execute("""
        UPDATE teams
        SET
            name = ?,
            team_type = ?
        WHERE id = ?
        """, (
            name,
            team_type,
            self.selected_team_id
        ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Team Updated",
            "Team or house updated successfully."
        )

        self.clear_selection()
        self.load_teams()


    # =========================
    # DELETE TEAM
    # =========================
    def delete_team(self):

        if not self.selected_team_id:

            QMessageBox.warning(
                self,
                "No Team Selected",
                (
                    "Select a team or house "
                    "from the list first."
                )
            )

            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Check whether athletes
        # are still assigned to it.
        cur.execute("""
        SELECT COUNT(*)
        FROM athletes
        WHERE team_id = ?
        """, (
            self.selected_team_id,
        ))

        athlete_count = cur.fetchone()[0]

        if athlete_count > 0:

            conn.close()

            QMessageBox.warning(
                self,
                "Team Cannot Be Deleted",
                (
                    f"{athlete_count} athlete(s) "
                    "are still assigned to this "
                    "team or house.\n\n"
                    "Move or delete those athletes "
                    "before deleting the team."
                )
            )

            return

        name = self.team_name.text().strip()

        answer = QMessageBox.question(
            self,
            "Delete Team",
            (
                f"Are you sure you want to "
                f"delete {name}?"
            ),
            QMessageBox.Yes
            | QMessageBox.No
        )

        if answer != QMessageBox.Yes:

            conn.close()
            return

        cur.execute("""
        DELETE FROM teams
        WHERE id = ?
        """, (
            self.selected_team_id,
        ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Team Deleted",
            "Team or house deleted successfully."
        )

        self.clear_selection()
        self.load_teams()


    # =========================
    # CLEAR SELECTION
    # =========================
    def clear_selection(self):

        self.selected_team_id = None

        self.team_list.clearSelection()

        self.team_name.clear()

        if self.team_type.count() > 0:

            self.team_type.setCurrentIndex(0)

# =========================
# ATHLETE SCREEN
# =========================
class AthleteScreen(QWidget):

    def __init__(self):
        super().__init__()

        self.selected_athlete_id = None

        layout = QVBoxLayout()

        title = QLabel("Athletes")
        title.setAlignment(Qt.AlignCenter)

        self.name_input = QLineEdit()
        self.name_input.setPlaceholderText("Athlete Name")

        self.gender_combo = QComboBox()
        self.gender_combo.addItems([
            "Male",
            "Female"
        ])

        self.age_combo = QComboBox()
        self.age_combo.addItems([
            "U7", "U8", "U9", "U10",
            "U11", "U12", "U13", "U14",
            "U15", "U16", "U17", "U18",
            "U19", "Open"
        ])

        self.team_combo = QComboBox()

        # -------------------------
        # BUTTONS
        # -------------------------
        add_button = QPushButton("Add Athlete")
        add_button.clicked.connect(
            self.save_athlete
        )

        update_button = QPushButton(
            "Update Selected Athlete"
        )
        update_button.clicked.connect(
            self.update_athlete
        )

        delete_button = QPushButton(
            "Delete Selected Athlete"
        )
        delete_button.clicked.connect(
            self.delete_athlete
        )

        withdraw_button = QPushButton(
            "Withdraw Selected Athlete"
        )

        withdraw_button.clicked.connect(
            self.withdraw_athlete
        )

        reactivate_button = QPushButton(
            "Reactivate Selected Athlete"
        )

        reactivate_button.clicked.connect(
            self.reactivate_athlete
        )

        clear_button = QPushButton(
            "Clear Selection"
        )
        clear_button.clicked.connect(
            self.clear_selection
        )

        refresh_button = QPushButton(
            "Refresh Athletes"
        )
        refresh_button.clicked.connect(
            self.load_athletes
        )

        # -------------------------
        # ATHLETE LIST
        # -------------------------
        self.athlete_list = QListWidget()

        self.athlete_list.itemClicked.connect(
            self.load_selected_athlete
        )

        # -------------------------
        # LAYOUT
        # -------------------------
        layout.addWidget(title)

        layout.addWidget(QLabel("Athlete Name"))
        layout.addWidget(self.name_input)

        layout.addWidget(QLabel("Gender"))
        layout.addWidget(self.gender_combo)

        layout.addWidget(QLabel("Age Group"))
        layout.addWidget(self.age_combo)

        layout.addWidget(QLabel("Team / House"))
        layout.addWidget(self.team_combo)

        layout.addWidget(add_button)
        layout.addWidget(update_button)
        layout.addWidget(withdraw_button)
        layout.addWidget(reactivate_button)
        layout.addWidget(delete_button)
        layout.addWidget(clear_button)

        layout.addWidget(
            QLabel("Existing Athletes")
        )

        layout.addWidget(self.athlete_list)
        layout.addWidget(refresh_button)

        self.setLayout(layout)

        self.load_teams()
        self.load_athletes()


    # =========================
    # LOAD TEAMS
    # =========================
    def load_teams(self):

        self.team_combo.clear()

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT id, name
        FROM teams
        WHERE meeting_id = ?
        ORDER BY name
        """, (meeting_id,))

        rows = cur.fetchall()

        conn.close()

        for team_id, name in rows:

            self.team_combo.addItem(
                name,
                team_id
            )


    # =========================
    # ADD ATHLETE
    # =========================
    def save_athlete(self):

        meeting_id = get_active_meeting()

        if not meeting_id:

            QMessageBox.warning(
                self,
                "No Active Meeting",
                "Select an active meeting first."
            )

            return

        name = self.name_input.text().strip()

        if not name:

            QMessageBox.warning(
                self,
                "Missing Name",
                "Enter the athlete's name."
            )

            return

        team_id = self.team_combo.currentData()

        if team_id is None:

            QMessageBox.warning(
                self,
                "No Team",
                "Select a team or house first."
            )

            return

        gender = self.gender_combo.currentText()
        age_group = self.age_combo.currentText()

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        INSERT INTO athletes
        (
            meeting_id,
            full_name,
            gender,
            age_group,
            team_id
        )
        VALUES (?, ?, ?, ?, ?)
        """, (
            meeting_id,
            name,
            gender,
            age_group,
            team_id
        ))

        conn.commit()
        conn.close()

        self.clear_selection()
        self.load_athletes()

        QMessageBox.information(
            self,
            "Athlete Added",
            "Athlete added successfully."
        )


    # =========================
    # LOAD ATHLETES
    # =========================
    def load_athletes(self):

        self.athlete_list.clear()

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            a.id,
            a.full_name,
            a.gender,
            a.age_group,
            a.team_id,
            t.name,
            a.status
        FROM athletes a

        LEFT JOIN teams t
            ON a.team_id = t.id

        WHERE a.meeting_id = ?

        ORDER BY
            a.status,
            a.full_name
        """, (meeting_id,))

        rows = cur.fetchall()

        conn.close()

        for (
            athlete_id,
            name,
            gender,
            age_group,
            team_id,
            team_name,
            status
        ) in rows:

            if not team_name:
                team_name = "No Team"

            item = QListWidgetItem(
                f"{name} | {gender} | "
                f"{age_group} | {team_name} | "
                f"{status}"
            )

            item.setData(
                Qt.UserRole,
                athlete_id
            )

            self.athlete_list.addItem(
                item
            )

    # =========================
    # LOAD SELECTED ATHLETE
    # =========================
    def load_selected_athlete(
        self,
        item
    ):

        athlete_id = item.data(
            Qt.UserRole
        )

        if not athlete_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            full_name,
            gender,
            age_group,
            team_id
        FROM athletes
        WHERE id = ?
        """, (athlete_id,))

        athlete = cur.fetchone()

        conn.close()

        if not athlete:
            return

        name = athlete[0]
        gender = athlete[1]
        age_group = athlete[2]
        team_id = athlete[3]

        self.selected_athlete_id = (
            athlete_id
        )

        self.name_input.setText(
            name
        )

        gender_index = (
            self.gender_combo.findText(
                gender
            )
        )

        if gender_index >= 0:
            self.gender_combo.setCurrentIndex(
                gender_index
            )

        age_index = (
            self.age_combo.findText(
                age_group
            )
        )

        if age_index >= 0:
            self.age_combo.setCurrentIndex(
                age_index
            )

        team_index = (
            self.team_combo.findData(
                team_id
            )
        )

        if team_index >= 0:
            self.team_combo.setCurrentIndex(
                team_index
            )


    # =========================
    # UPDATE ATHLETE
    # =========================
    def update_athlete(self):

        if not self.selected_athlete_id:

            QMessageBox.warning(
                self,
                "No Athlete Selected",
                "Select an athlete from the list first."
            )

            return

        name = self.name_input.text().strip()

        if not name:

            QMessageBox.warning(
                self,
                "Missing Name",
                "Enter the athlete's name."
            )

            return

        gender = self.gender_combo.currentText()
        age_group = self.age_combo.currentText()
        team_id = self.team_combo.currentData()

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Get previous athlete details
        cur.execute("""
        SELECT
            gender,
            age_group,
            team_id
        FROM athletes
        WHERE id = ?
        """, (
            self.selected_athlete_id,
        ))

        previous = cur.fetchone()

        if not previous:

            conn.close()
            return

        old_gender = previous[0]
        old_age_group = previous[1]

        cur.execute("""
        UPDATE athletes
        SET
            full_name = ?,
            gender = ?,
            age_group = ?,
            team_id = ?
        WHERE id = ?
        """, (
            name,
            gender,
            age_group,
            team_id,
            self.selected_athlete_id
        ))

        # If only their team changes,
        # existing awarded points must move
        # to the new team.
        cur.execute("""
        UPDATE awarded_points
        SET team_id = ?
        WHERE athlete_id = ?
        """, (
            team_id,
            self.selected_athlete_id
        ))

        # If gender or age group changes,
        # existing event entries may now
        # be invalid.
        if (
            gender != old_gender
            or age_group != old_age_group
        ):

            athlete_id = (
                self.selected_athlete_id
            )

            cur.execute("""
            DELETE FROM awarded_points
            WHERE athlete_id = ?
            """, (athlete_id,))

            cur.execute("""
            DELETE FROM final_results
            WHERE athlete_id = ?
            """, (athlete_id,))

            cur.execute("""
            DELETE FROM finals
            WHERE athlete_id = ?
            """, (athlete_id,))

            cur.execute("""
            DELETE FROM results
            WHERE athlete_id = ?
            """, (athlete_id,))

            cur.execute("""
            DELETE FROM heat_entries
            WHERE athlete_id = ?
            """, (athlete_id,))

            # Remove measured field attempts
            cur.execute("""
            DELETE FROM field_attempts
            WHERE athlete_id = ?
            """, (athlete_id,))

            # Remove High Jump attempts
            cur.execute("""
            DELETE FROM high_jump_attempts
            WHERE athlete_id = ?
            """, (athlete_id,))

            cur.execute("""
            DELETE FROM event_entries
            WHERE athlete_id = ?
            """, (athlete_id,))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Athlete Updated",
            "Athlete updated successfully."
        )

        self.clear_selection()
        self.load_athletes()

    # =========================
    # WITHDRAW ATHLETE
    # =========================
    def withdraw_athlete(self):

        if not self.selected_athlete_id:

            QMessageBox.warning(
                self,
                "No Athlete Selected",
                "Select an athlete from the list first."
            )

            return

        name = self.name_input.text().strip()

        answer = QMessageBox.question(
            self,
            "Withdraw Athlete",
            (
                f"Mark {name} as withdrawn?\n\n"
                "Their existing entries, results "
                "and points will be preserved."
            ),
            QMessageBox.Yes
            | QMessageBox.No
        )

        if answer != QMessageBox.Yes:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        UPDATE athletes
        SET status = 'Withdrawn'
        WHERE id = ?
        """, (
            self.selected_athlete_id,
        ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Athlete Withdrawn",
            f"{name} has been marked as withdrawn."
        )

        self.clear_selection()
        self.load_athletes()

    # =========================
    # REACTIVATE ATHLETE
    # =========================
    def reactivate_athlete(self):

        if not self.selected_athlete_id:

            QMessageBox.warning(
                self,
                "No Athlete Selected",
                "Select an athlete from the list first."
            )

            return

        name = self.name_input.text().strip()

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        UPDATE athletes
        SET status = 'Active'
        WHERE id = ?
        """, (
            self.selected_athlete_id,
        ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Athlete Reactivated",
            f"{name} is active again."
        )

        self.clear_selection()
        self.load_athletes()

    # =========================
    # DELETE ATHLETE
    # =========================
    def delete_athlete(self):

        if not self.selected_athlete_id:

            QMessageBox.warning(
                self,
                "No Athlete Selected",
                "Select an athlete from the list first."
            )

            return

        name = self.name_input.text().strip()

        answer = QMessageBox.question(
            self,
            "Delete Athlete",
            (
                f"Are you sure you want to delete "
                f"{name}?\n\n"
                "Their event entries, heat entries, "
                "results, finals and awarded points "
                "will also be removed."
            ),
            QMessageBox.Yes
            | QMessageBox.No
        )

        if answer != QMessageBox.Yes:
            return

        athlete_id = (
            self.selected_athlete_id
        )

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Remove all related athletics data
        cur.execute("""
        DELETE FROM awarded_points
        WHERE athlete_id = ?
        """, (athlete_id,))

        cur.execute("""
        DELETE FROM final_results
        WHERE athlete_id = ?
        """, (athlete_id,))

        cur.execute("""
        DELETE FROM finals
        WHERE athlete_id = ?
        """, (athlete_id,))

        cur.execute("""
        DELETE FROM results
        WHERE athlete_id = ?
        """, (athlete_id,))

        cur.execute("""
        DELETE FROM heat_entries
        WHERE athlete_id = ?
        """, (athlete_id,))

        # Remove measured field attempts
        cur.execute("""
        DELETE FROM field_attempts
        WHERE athlete_id = ?
        """, (athlete_id,))

        # Remove High Jump attempts
        cur.execute("""
        DELETE FROM high_jump_attempts
        WHERE athlete_id = ?
        """, (athlete_id,))

        cur.execute("""
        DELETE FROM event_entries
        WHERE athlete_id = ?
        """, (athlete_id,))

        # Finally delete athlete
        cur.execute("""
        DELETE FROM athletes
        WHERE id = ?
        """, (athlete_id,))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Athlete Deleted",
            "Athlete deleted successfully."
        )

        self.clear_selection()
        self.load_athletes()


    # =========================
    # CLEAR SELECTION
    # =========================
    def clear_selection(self):

        self.selected_athlete_id = None

        self.athlete_list.clearSelection()

        self.name_input.clear()

        if self.gender_combo.count() > 0:
            self.gender_combo.setCurrentIndex(0)

        if self.age_combo.count() > 0:
            self.age_combo.setCurrentIndex(0)

        if self.team_combo.count() > 0:
            self.team_combo.setCurrentIndex(0)

# =========================
# EVENT SCREEN
# =========================
class EventScreen(QWidget):

    def __init__(self):
        super().__init__()

        self.selected_event_id = None

        layout = QVBoxLayout()

        title = QLabel("Events")
        title.setAlignment(Qt.AlignCenter)

        self.event_name = QLineEdit()
        self.event_name.setPlaceholderText("Event Name")

        self.gender_combo = QComboBox()
        self.gender_combo.addItems([
            "Male",
            "Female"
        ])

        self.age_combo = QComboBox()
        self.age_combo.addItems([
            "U7", "U8", "U9", "U10",
            "U11", "U12", "U13", "U14",
            "U15", "U16", "U17", "U18",
            "U19", "Open"
        ])

        self.event_type = QComboBox()
        self.event_type.addItems([
            "Track",
            "Field",
            "High Jump"
        ])

        self.format_combo = QComboBox()

        self.format_combo.addItems([
            "Timed Finals",
            "Normal Heats + Final"
        ])

        # Timed Finals is the default.
        self.format_combo.setCurrentIndex(0)

        # -------------------------
        # BUTTONS
        # -------------------------
        add_button = QPushButton(
            "Add Event"
        )
        add_button.clicked.connect(
            self.save_event
        )

        update_button = QPushButton(
            "Update Selected Event"
        )
        update_button.clicked.connect(
            self.update_event
        )

        delete_button = QPushButton(
            "Delete Selected Event"
        )
        delete_button.clicked.connect(
            self.delete_event
        )

        clear_button = QPushButton(
            "Clear Selection"
        )
        clear_button.clicked.connect(
            self.clear_selection
        )

        refresh_button = QPushButton(
            "Refresh Events"
        )
        refresh_button.clicked.connect(
            self.load_events
        )

        # -------------------------
        # EVENT LIST
        # -------------------------
        self.event_list = QListWidget()

        self.event_list.itemClicked.connect(
            self.load_selected_event
        )

        # -------------------------
        # LAYOUT
        # -------------------------
        layout.addWidget(title)

        layout.addWidget(
            QLabel("Event Name")
        )
        layout.addWidget(self.event_name)

        layout.addWidget(
            QLabel("Gender")
        )
        layout.addWidget(self.gender_combo)

        layout.addWidget(
            QLabel("Age Group")
        )
        layout.addWidget(self.age_combo)

        layout.addWidget(
            QLabel("Event Type")
        )
        layout.addWidget(self.event_type)

        layout.addWidget(
            QLabel("Competition Format")
        )
        layout.addWidget(self.format_combo)

        layout.addWidget(add_button)
        layout.addWidget(update_button)
        layout.addWidget(delete_button)
        layout.addWidget(clear_button)

        layout.addWidget(
            QLabel("Existing Events")
        )

        layout.addWidget(self.event_list)
        layout.addWidget(refresh_button)

        self.setLayout(layout)

        self.load_events()


    # =========================
    # ADD EVENT
    # =========================
    def save_event(self):

        meeting_id = get_active_meeting()

        if not meeting_id:

            QMessageBox.warning(
                self,
                "No Active Meeting",
                "Select an active meeting first."
            )

            return

        name = self.event_name.text().strip()

        if not name:

            QMessageBox.warning(
                self,
                "Missing Event Name",
                "Enter an event name."
            )

            return

        event_type = (
            self.event_type.currentText()
        )

        gender = (
            self.gender_combo.currentText()
        )

        age_group = (
            self.age_combo.currentText()
        )

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Prevent duplicate events
        cur.execute("""
        SELECT id
        FROM events
        WHERE meeting_id = ?
        AND LOWER(event_name) = LOWER(?)
        AND gender = ?
        AND age_group = ?
        """, (
            meeting_id,
            name,
            gender,
            age_group
        ))

        if cur.fetchone():

            conn.close()

            QMessageBox.warning(
                self,
                "Duplicate Event",
                (
                    "This event already exists "
                    "for this gender and age group."
                )
            )

            return

        cur.execute("""
        INSERT INTO events
        (
            meeting_id,
            event_name,
            event_type,
            gender,
            age_group,
            competition_format
        )
        VALUES (?, ?, ?, ?, ?, ?)
        """, (
            meeting_id,
            name,
            event_type,
            gender,
            age_group,
            competition_format
        ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Event Added",
            "Event added successfully."
        )

        self.clear_selection()
        self.load_events()


    # =========================
    # LOAD EVENTS
    # =========================
    def load_events(self):

        self.event_list.clear()

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            id,
            event_name,
            gender,
            age_group,
            event_type
        FROM events
        WHERE meeting_id = ?
        ORDER BY
            event_name,
            gender,
            age_group
        """, (meeting_id,))

        rows = cur.fetchall()

        conn.close()

        for (
            event_id,
            name,
            gender,
            age_group,
            event_type
        ) in rows:

            item = QListWidgetItem(
                f"{name} "
                f"({gender} {age_group}) "
                f"- {event_type}"
            )

            item.setData(
                Qt.UserRole,
                event_id
            )

            self.event_list.addItem(
                item
            )


    # =========================
    # LOAD SELECTED EVENT
    # =========================
    def load_selected_event(
        self,
        item
    ):

        event_id = item.data(
            Qt.UserRole
        )

        if not event_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            event_name,
            event_type,
            gender,
            age_group
        FROM events
        WHERE id = ?
        """, (event_id,))

        event = cur.fetchone()

        conn.close()

        if not event:
            return

        name = event[0]
        event_type = event[1]
        gender = event[2]
        age_group = event[3]

        self.selected_event_id = (
            event_id
        )

        self.event_name.setText(
            name
        )

        gender_index = (
            self.gender_combo.findText(
                gender
            )
        )

        if gender_index >= 0:

            self.gender_combo.setCurrentIndex(
                gender_index
            )

        age_index = (
            self.age_combo.findText(
                age_group
            )
        )

        if age_index >= 0:

            self.age_combo.setCurrentIndex(
                age_index
            )

        type_index = (
            self.event_type.findText(
                event_type
            )
        )

        if type_index >= 0:

            self.event_type.setCurrentIndex(
                type_index
            )


    # =========================
    # UPDATE EVENT
    # =========================
    def update_event(self):

        if not self.selected_event_id:

            QMessageBox.warning(
                self,
                "No Event Selected",
                "Select an event from the list first."
            )

            return

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        name = self.event_name.text().strip()

        if not name:

            QMessageBox.warning(
                self,
                "Missing Event Name",
                "Enter an event name."
            )

            return

        event_type = (
            self.event_type.currentText()
        )

        competition_format = (
            self.format_combo.currentText()
        )

        gender = (
            self.gender_combo.currentText()
        )

        age_group = (
            self.age_combo.currentText()
        )

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Get current event details
        cur.execute("""
        SELECT
            event_name,
            event_type,
            gender,
            age_group
        FROM events
        WHERE id = ?
        """, (
            self.selected_event_id,
        ))

        previous = cur.fetchone()

        if not previous:

            conn.close()
            return

        old_name = previous[0]
        old_event_type = previous[1]
        old_gender = previous[2]
        old_age_group = previous[3]

        # Check for duplicate event
        cur.execute("""
        SELECT id
        FROM events
        WHERE meeting_id = ?
        AND LOWER(event_name) = LOWER(?)
        AND gender = ?
        AND age_group = ?
        AND id != ?
        """, (
            meeting_id,
            name,
            gender,
            age_group,
            self.selected_event_id
        ))

        if cur.fetchone():

            conn.close()

            QMessageBox.warning(
                self,
                "Duplicate Event",
                (
                    "Another event already exists "
                    "with this name, gender "
                    "and age group."
                )
            )

            return

        # Changing only the event name is safe.
        #
        # Changing gender, age group or type
        # can invalidate existing competition data.
        structure_changed = (
            event_type != old_event_type
            or gender != old_gender
            or age_group != old_age_group
        )

        if structure_changed:

            answer = QMessageBox.question(
                self,
                "Event Structure Changed",
                (
                    "You changed the event's gender, "
                    "age group or event type.\n\n"
                    "Existing entries, heats, results, "
                    "finals and points for this event "
                    "must be cleared.\n\n"
                    "Continue with this change?"
                ),
                QMessageBox.Yes
                | QMessageBox.No
            )

            if answer != QMessageBox.Yes:

                conn.close()
                return

            event_id = (
                self.selected_event_id
            )

            # Remove points
            cur.execute("""
            DELETE FROM awarded_points
            WHERE event_id = ?
            """, (event_id,))

            # Remove final results
            cur.execute("""
            DELETE FROM final_results
            WHERE event_id = ?
            """, (event_id,))

                        # Remove measured field attempts
            cur.execute("""
            DELETE FROM field_attempts
            WHERE event_id = ?
            """, (event_id,))

            # Remove High Jump attempts
            cur.execute("""
            DELETE FROM high_jump_attempts
            WHERE event_id = ?
            """, (event_id,))

            # Remove High Jump heights
            cur.execute("""
            DELETE FROM high_jump_heights
            WHERE event_id = ?
            """, (event_id,))

            # Remove finalists
            cur.execute("""
            DELETE FROM finals
            WHERE event_id = ?
            """, (event_id,))

            # Remove heat results
            cur.execute("""
            DELETE FROM results
            WHERE heat_id IN (
                SELECT id
                FROM heats
                WHERE event_id = ?
            )
            """, (event_id,))

            # Remove heat entries
            cur.execute("""
            DELETE FROM heat_entries
            WHERE heat_id IN (
                SELECT id
                FROM heats
                WHERE event_id = ?
            )
            """, (event_id,))

            # Remove heats
            cur.execute("""
            DELETE FROM heats
            WHERE event_id = ?
            """, (event_id,))

            # Remove event entries
            cur.execute("""
            DELETE FROM event_entries
            WHERE event_id = ?
            """, (event_id,))

        cur.execute("""
        UPDATE events
        SET
            event_name = ?,
            event_type = ?,
            gender = ?,
            age_group = ?
        WHERE id = ?
        """, (
            name,
            event_type,
            gender,
            age_group,
            self.selected_event_id
        ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Event Updated",
            "Event updated successfully."
        )

        self.clear_selection()
        self.load_events()


    # =========================
    # DELETE EVENT
    # =========================
    def delete_event(self):

        if not self.selected_event_id:

            QMessageBox.warning(
                self,
                "No Event Selected",
                "Select an event from the list first."
            )

            return

        name = self.event_name.text().strip()

        gender = (
            self.gender_combo.currentText()
        )

        age_group = (
            self.age_combo.currentText()
        )

        answer = QMessageBox.question(
            self,
            "Delete Event",
            (
                f"Are you sure you want to delete:\n\n"
                f"{name} ({gender} {age_group})?\n\n"
                "All entries, heats, results, finals "
                "and points for this event will also "
                "be permanently deleted."
            ),
            QMessageBox.Yes
            | QMessageBox.No
        )

        if answer != QMessageBox.Yes:
            return

        event_id = (
            self.selected_event_id
        )

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Remove awarded points
        cur.execute("""
        DELETE FROM awarded_points
        WHERE event_id = ?
        """, (event_id,))

        # Remove final results
        cur.execute("""
        DELETE FROM final_results
        WHERE event_id = ?
        """, (event_id,))

        # Remove measured field attempts
        cur.execute("""
        DELETE FROM field_attempts
        WHERE event_id = ?
        """, (event_id,))

        # Remove High Jump attempts
        cur.execute("""
        DELETE FROM high_jump_attempts
        WHERE event_id = ?
        """, (event_id,))

        # Remove High Jump heights
        cur.execute("""
        DELETE FROM high_jump_heights
        WHERE event_id = ?
        """, (event_id,))

        # Remove finalists
        cur.execute("""
        DELETE FROM finals
        WHERE event_id = ?
        """, (event_id,))

        # Results depend on heat IDs
        cur.execute("""
        DELETE FROM results
        WHERE heat_id IN (
            SELECT id
            FROM heats
            WHERE event_id = ?
        )
        """, (event_id,))

        # Heat entries depend on heat IDs
        cur.execute("""
        DELETE FROM heat_entries
        WHERE heat_id IN (
            SELECT id
            FROM heats
            WHERE event_id = ?
        )
        """, (event_id,))

        # Remove heats
        cur.execute("""
        DELETE FROM heats
        WHERE event_id = ?
        """, (event_id,))

        # Remove event entries
        cur.execute("""
        DELETE FROM event_entries
        WHERE event_id = ?
        """, (event_id,))

        # Finally remove event itself
        cur.execute("""
        DELETE FROM events
        WHERE id = ?
        """, (event_id,))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Event Deleted",
            "Event deleted successfully."
        )

        self.clear_selection()
        self.load_events()


    # =========================
    # CLEAR SELECTION
    # =========================
    def clear_selection(self):

        self.selected_event_id = None

        self.event_list.clearSelection()

        self.event_name.clear()

        if self.gender_combo.count() > 0:
            self.gender_combo.setCurrentIndex(0)

        if self.age_combo.count() > 0:
            self.age_combo.setCurrentIndex(0)

        if self.event_type.count() > 0:
            self.event_type.setCurrentIndex(0)

# =========================
# EVENT ENTRY SCREEN
# =========================
class EventEntryScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        title = QLabel("Event Entries")
        title.setAlignment(Qt.AlignCenter)

        self.event_combo = QComboBox()
        self.athlete_combo = QComboBox()

        add_button = QPushButton("Add Athlete to Event")
        add_button.clicked.connect(self.add_entry)

        remove_button = QPushButton(
            "Remove Selected Entry"
        )
        remove_button.clicked.connect(
            self.remove_entry
        )

        withdraw_entry_button = QPushButton(
            "Withdraw Selected Entry"
        )

        withdraw_entry_button.clicked.connect(
            self.withdraw_event_entry
        )

        reactivate_entry_button = QPushButton(
            "Reactivate Selected Entry"
        )

        reactivate_entry_button.clicked.connect(
            self.reactivate_event_entry
        )

        self.entry_list = QListWidget()

        refresh_button = QPushButton("Refresh")
        refresh_button.clicked.connect(self.load_entries)

        generate_heats_button = QPushButton(
            "Generate Heats / Direct Final"
        )
        generate_heats_button.clicked.connect(
            self.generate_heats
        )

        # =========================
        # HEAT / LANE EDITOR
        # =========================
        self.heat_combo = QComboBox()

        load_heat_button = QPushButton(
            "Load Heat for Editing"
        )
        load_heat_button.clicked.connect(
            self.load_heat_for_edit
        )

        self.heat_table = QTableWidget()
        self.heat_table.setColumnCount(2)

        self.heat_table.setHorizontalHeaderLabels([
            "Lane",
            "Athlete"
        ])

        save_lanes_button = QPushButton(
            "Save Lane Changes"
        )
        save_lanes_button.clicked.connect(
            self.save_lane_changes
        )

                # =========================
        # MOVE ATHLETE BETWEEN HEATS
        # =========================
        self.move_heat_combo = QComboBox()

        self.move_lane_combo = QComboBox()

        move_athlete_button = QPushButton(
            "Move Selected Athlete"
        )

        move_athlete_button.clicked.connect(
            self.move_athlete_between_heats
        )

        reload_button = QPushButton("Reload Data")
        reload_button.clicked.connect(
            self.refresh_screen
        )

        layout.addWidget(title)
        layout.addWidget(QLabel("Select Event"))
        layout.addWidget(self.event_combo)

        layout.addWidget(QLabel("Select Athlete"))
        layout.addWidget(self.athlete_combo)

        layout.addWidget(add_button)
        layout.addWidget(remove_button)
        layout.addWidget(withdraw_entry_button)
        layout.addWidget(reactivate_entry_button)
        layout.addWidget(self.entry_list)

        layout.addWidget(refresh_button)
        layout.addWidget(reload_button)
        layout.addWidget(generate_heats_button)

        layout.addWidget(
            QLabel("Edit Generated Heat / Lanes")
        )

        layout.addWidget(self.heat_combo)
        layout.addWidget(load_heat_button)
        layout.addWidget(self.heat_table)
        layout.addWidget(save_lanes_button)

        layout.addWidget(
            QLabel(
                "Move Selected Athlete "
                "to Another Heat"
            )
        )

        layout.addWidget(
            QLabel("Destination Heat")
        )

        layout.addWidget(
            self.move_heat_combo
        )

        layout.addWidget(
            QLabel("Destination Lane")
        )

        layout.addWidget(
            self.move_lane_combo
        )

        layout.addWidget(
            move_athlete_button
        )

        self.setLayout(layout)

        self.load_events()
        self.load_athletes()
        self.load_entries()

        self.filter_athletes_by_event()

        self.event_combo.currentIndexChanged.connect(
            self.filter_athletes_by_event
        )

        self.load_heat_choices()

        self.heat_combo.currentIndexChanged.connect(
            self.load_heat_for_edit
        )

        self.heat_combo.currentIndexChanged.connect(
            self.load_destination_heats
        )

        self.move_heat_combo.currentIndexChanged.connect(
            self.load_destination_lanes
        )

        self.load_heat_for_edit()

        self.load_destination_heats()

    def load_athletes(self):

        self.filter_athletes_by_event()

    def add_entry(self):

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        event_id = self.event_combo.currentData()
        athlete_id = self.athlete_combo.currentData()

        if not event_id or not athlete_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Event details
        cur.execute("""
        SELECT gender, age_group
        FROM events
        WHERE id = ?
        """, (event_id,))

        event_gender, event_age = cur.fetchone()

        # Athlete details
        cur.execute("""
        SELECT gender, age_group
        FROM athletes
        WHERE id = ?
        """, (athlete_id,))

        ath_gender, ath_age = cur.fetchone()

        # Validation
        if event_gender != ath_gender or event_age != ath_age:

            QMessageBox.warning(
                self,
                "Invalid Entry",
                "Athlete does not match this event."
            )

            conn.close()
            return

        cur.execute("""
        SELECT id
        FROM event_entries
        WHERE meeting_id = ?
        AND event_id = ?
        AND athlete_id = ?
        """, (
            meeting_id,
            event_id,
            athlete_id
        ))

        if cur.fetchone():

            QMessageBox.warning(
                self,
                "Duplicate Entry",
                "This athlete is already entered in this event."
            )

            conn.close()
            return

        # Add entry
        cur.execute("""
        INSERT INTO event_entries
        (meeting_id, event_id, athlete_id)
        VALUES (?, ?, ?)
        """, (
            meeting_id,
            event_id,
            athlete_id
        ))

        conn.commit()
        conn.close()

        self.load_entries()

    def load_entries(self):

        self.entry_list.clear()

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            e.event_name,
            e.gender,
            e.age_group,
            a.full_name,
            ee.event_id,
            ee.athlete_id,
            ee.status,
            a.status
        FROM event_entries ee

        JOIN events e
            ON ee.event_id = e.id

        JOIN athletes a
            ON ee.athlete_id = a.id

        WHERE ee.meeting_id = ?

        ORDER BY
            e.event_name,
            e.gender,
            e.age_group,
            ee.status,
            a.full_name
        """, (
            meeting_id,
        ))

        rows = cur.fetchall()

        conn.close()

        current_event = ""

        for (
            event,
            gender,
            age_group,
            athlete,
            event_id,
            athlete_id,
            entry_status,
            athlete_status
        ) in rows:

            heading = (
                f"{event} "
                f"({gender} {age_group})"
            )

            if heading != current_event:

                self.entry_list.addItem("")

                heading_item = QListWidgetItem(
                    heading
                )

                self.entry_list.addItem(
                    heading_item
                )

                current_event = heading

            # =========================
            # DISPLAY STATUS
            # =========================
            status_text = ""

            if athlete_status != "Active":

                status_text = (
                    f"  [ATHLETE {athlete_status.upper()}]"
                )

            elif entry_status != "Active":

                status_text = (
                    f"  [{entry_status.upper()}]"
                )

            athlete_item = QListWidgetItem(
                f"    {athlete}"
                f"{status_text}"
            )

            athlete_item.setData(
                Qt.UserRole,
                (
                    event_id,
                    athlete_id
                )
            )

            self.entry_list.addItem(
                athlete_item
            )

    # =========================
    # SELECTED EVENT ENTRY
    # =========================
    def get_selected_event_entry(
        self
    ):

        item = (
            self.entry_list.currentItem()
        )

        if not item:

            QMessageBox.warning(
                self,
                "No Entry Selected",
                (
                    "Select an athlete underneath "
                    "an event first."
                )
            )

            return None

        entry_data = item.data(
            Qt.UserRole
        )

        if not entry_data:

            QMessageBox.warning(
                self,
                "Invalid Selection",
                (
                    "Select an athlete entry, "
                    "not an event heading."
                )
            )

            return None

        try:

            event_id, athlete_id = (
                entry_data
            )

        except (
            TypeError,
            ValueError
        ):

            QMessageBox.warning(
                self,
                "Invalid Selection",
                "Select an athlete entry."
            )

            return None

        return (
            event_id,
            athlete_id
        )


    # =========================
    # WITHDRAW EVENT ENTRY
    # =========================
    def withdraw_event_entry(self):

        selected = (
            self.get_selected_event_entry()
        )

        if not selected:
            return

        event_id, athlete_id = selected

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # =========================
        # ENTRY DETAILS
        # =========================
        cur.execute("""
        SELECT
            ee.status,
            a.full_name,
            e.event_name,
            e.gender,
            e.age_group
        FROM event_entries ee

        JOIN athletes a
            ON ee.athlete_id = a.id

        JOIN events e
            ON ee.event_id = e.id

        WHERE ee.meeting_id = ?
        AND ee.event_id = ?
        AND ee.athlete_id = ?
        """, (
            meeting_id,
            event_id,
            athlete_id
        ))

        details = cur.fetchone()

        if not details:

            conn.close()
            return

        current_status = details[0]
        athlete_name = details[1]
        event_name = details[2]
        gender = details[3]
        age_group = details[4]

        if current_status == "Withdrawn":

            conn.close()

            QMessageBox.information(
                self,
                "Already Withdrawn",
                (
                    f"{athlete_name} is already "
                    "withdrawn from this event."
                )
            )

            return

        # =========================
        # CHECK RECORDED RESULTS
        # =========================
        cur.execute("""
        SELECT COUNT(*)
        FROM results r

        JOIN heats h
            ON r.heat_id = h.id

        WHERE h.event_id = ?
        AND r.athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        heat_results = (
            cur.fetchone()[0]
        )

        cur.execute("""
        SELECT COUNT(*)
        FROM final_results
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        final_results = (
            cur.fetchone()[0]
        )

        if (
            heat_results > 0
            or final_results > 0
        ):

            conn.close()

            QMessageBox.warning(
                self,
                "Competition Already Started",
                (
                    f"{athlete_name} already has "
                    "competition data for:\n\n"
                    f"{event_name} "
                    f"({gender} {age_group})\n\n"
                    "Do not withdraw the event entry "
                    "at this stage.\n\n"
                    "If the athlete did not start, "
                    "enter DNS in Results or Finals."
                )
            )

            return

        # =========================
        # CHECK GENERATED STRUCTURE
        # =========================
        cur.execute("""
        SELECT COUNT(*)
        FROM heats
        WHERE meeting_id = ?
        AND event_id = ?
        """, (
            meeting_id,
            event_id
        ))

        heat_count = (
            cur.fetchone()[0]
        )

        cur.execute("""
        SELECT COUNT(*)
        FROM finals
        WHERE event_id = ?
        """, (
            event_id,
        ))

        finalist_count = (
            cur.fetchone()[0]
        )

        if (
            heat_count > 0
            or finalist_count > 0
        ):

            answer = QMessageBox.question(
                self,
                "Event Already Generated",
                (
                    f"{athlete_name} can be marked "
                    "as withdrawn from this event.\n\n"
                    "However, heats or a direct final "
                    "have already been generated.\n\n"
                    "After withdrawing the athlete, "
                    "regenerate the heats/direct final "
                    "so the athlete is removed from "
                    "the competition structure.\n\n"
                    "Continue?"
                ),
                QMessageBox.Yes
                | QMessageBox.No
            )

            if answer != QMessageBox.Yes:

                conn.close()
                return

        else:

            answer = QMessageBox.question(
                self,
                "Withdraw Event Entry",
                (
                    f"Withdraw {athlete_name} from:\n\n"
                    f"{event_name} "
                    f"({gender} {age_group})?\n\n"
                    "This affects only this event."
                ),
                QMessageBox.Yes
                | QMessageBox.No
            )

            if answer != QMessageBox.Yes:

                conn.close()
                return

        # =========================
        # WITHDRAW
        # =========================
        cur.execute("""
        UPDATE event_entries
        SET status = 'Withdrawn'
        WHERE meeting_id = ?
        AND event_id = ?
        AND athlete_id = ?
        """, (
            meeting_id,
            event_id,
            athlete_id
        ))

        conn.commit()
        conn.close()

        self.load_entries()

        QMessageBox.information(
            self,
            "Entry Withdrawn",
            (
                f"{athlete_name} has been withdrawn "
                "from this event only."
            )
        )


    # =========================
    # REACTIVATE EVENT ENTRY
    # =========================
    def reactivate_event_entry(self):

        selected = (
            self.get_selected_event_entry()
        )

        if not selected:
            return

        event_id, athlete_id = selected

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            ee.status,
            a.full_name,
            a.status,
            e.event_name
        FROM event_entries ee

        JOIN athletes a
            ON ee.athlete_id = a.id

        JOIN events e
            ON ee.event_id = e.id

        WHERE ee.meeting_id = ?
        AND ee.event_id = ?
        AND ee.athlete_id = ?
        """, (
            meeting_id,
            event_id,
            athlete_id
        ))

        details = cur.fetchone()

        if not details:

            conn.close()
            return

        entry_status = details[0]
        athlete_name = details[1]
        athlete_status = details[2]
        event_name = details[3]

        if entry_status == "Active":

            conn.close()

            QMessageBox.information(
                self,
                "Already Active",
                (
                    f"{athlete_name}'s entry "
                    "is already active."
                )
            )

            return

        # Global withdrawal takes priority
        if athlete_status != "Active":

            conn.close()

            QMessageBox.warning(
                self,
                "Athlete Withdrawn",
                (
                    f"{athlete_name} is currently "
                    "withdrawn as an athlete.\n\n"
                    "Reactivate the athlete on the "
                    "Athletes screen first."
                )
            )

            return

        cur.execute("""
        SELECT COUNT(*)
        FROM heats
        WHERE meeting_id = ?
        AND event_id = ?
        """, (
            meeting_id,
            event_id
        ))

        heat_count = (
            cur.fetchone()[0]
        )

        cur.execute("""
        SELECT COUNT(*)
        FROM finals
        WHERE event_id = ?
        """, (
            event_id,
        ))

        finalist_count = (
            cur.fetchone()[0]
        )

        cur.execute("""
        UPDATE event_entries
        SET status = 'Active'
        WHERE meeting_id = ?
        AND event_id = ?
        AND athlete_id = ?
        """, (
            meeting_id,
            event_id,
            athlete_id
        ))

        conn.commit()
        conn.close()

        self.load_entries()

        if (
            heat_count > 0
            or finalist_count > 0
        ):

            QMessageBox.information(
                self,
                "Entry Reactivated",
                (
                    f"{athlete_name} has been "
                    f"reactivated for {event_name}.\n\n"
                    "Heats or a direct final already "
                    "exist, so regenerate them if "
                    "this athlete must be included."
                )
            )

        else:

            QMessageBox.information(
                self,
                "Entry Reactivated",
                (
                    f"{athlete_name} has been "
                    f"reactivated for {event_name}."
                )
            )
            
    # =========================
    # REMOVE EVENT ENTRY
    # =========================
    def remove_entry(self):

        item = self.entry_list.currentItem()

        if not item:

            QMessageBox.warning(
                self,
                "No Entry Selected",
                (
                    "Select an athlete underneath "
                    "an event first."
                )
            )

            return

        entry_data = item.data(
            Qt.UserRole
        )

        # Headings and blank lines
        # do not contain entry data.
        if not entry_data:

            QMessageBox.warning(
                self,
                "Invalid Selection",
                (
                    "Select an athlete entry, "
                    "not an event heading or heat."
                )
            )

            return

        try:
            event_id, athlete_id = entry_data

        except (TypeError, ValueError):

            QMessageBox.warning(
                self,
                "Invalid Selection",
                (
                    "Select an athlete entry "
                    "from the event list."
                )
            )

            return

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # -------------------------
        # GET NAMES
        # -------------------------
        cur.execute("""
        SELECT
            a.full_name,
            e.event_name,
            e.gender,
            e.age_group,
            e.event_type
        FROM athletes a,
             events e
        WHERE a.id = ?
        AND e.id = ?
        """, (
            athlete_id,
            event_id
        ))

        details = cur.fetchone()

        if not details:

            conn.close()
            return

        athlete_name = details[0]
        event_name = details[1]
        gender = details[2]
        age_group = details[3]
        event_type = details[4]

        # -------------------------
        # CHECK LINKED DATA
        # -------------------------
        cur.execute("""
        SELECT COUNT(*)
        FROM heat_entries he

        JOIN heats h
            ON he.heat_id = h.id

        WHERE h.event_id = ?
        AND he.athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        heat_entry_count = (
            cur.fetchone()[0]
        )

        cur.execute("""
        SELECT COUNT(*)
        FROM results r

        JOIN heats h
            ON r.heat_id = h.id

        WHERE h.event_id = ?
        AND r.athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        result_count = (
            cur.fetchone()[0]
        )

        cur.execute("""
        SELECT COUNT(*)
        FROM finals
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        finalist_count = (
            cur.fetchone()[0]
        )

        cur.execute("""
        SELECT COUNT(*)
        FROM final_results
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        final_result_count = (
            cur.fetchone()[0]
        )

        # -------------------------
        # FIELD ATTEMPTS
        # -------------------------
        cur.execute("""
        SELECT COUNT(*)
        FROM field_attempts
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        field_attempt_count = (
            cur.fetchone()[0]
        )

        # -------------------------
        # HIGH JUMP ATTEMPTS
        # -------------------------
        cur.execute("""
        SELECT COUNT(*)
        FROM high_jump_attempts
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        high_jump_attempt_count = (
            cur.fetchone()[0]
        )

        has_competition_data = (
            heat_entry_count > 0
            or result_count > 0
            or finalist_count > 0
            or final_result_count > 0
            or field_attempt_count > 0
            or high_jump_attempt_count > 0
        )

        # -------------------------
        # CONFIRM REMOVAL
        # -------------------------
        if has_competition_data:

            message = (
                f"{athlete_name} already has "
                f"generated competition data for:\n\n"
                f"{event_name} "
                f"({gender} {age_group})\n\n"
                "Removing this entry will also "
                "remove this athlete's heat entry, "
                "heat result, finalist record, "
                "final result and points for this "
                "event.\n\n"
                "Other athletes will not be deleted.\n\n"
                "Continue?"
            )

        else:

            message = (
                f"Remove {athlete_name} from:\n\n"
                f"{event_name} "
                f"({gender} {age_group})?"
            )

        answer = QMessageBox.question(
            self,
            "Remove Event Entry",
            message,
            QMessageBox.Yes
            | QMessageBox.No
        )

        if answer != QMessageBox.Yes:

            conn.close()
            return

        # -------------------------
        # REMEMBER AFFECTED HEATS
        # -------------------------
        cur.execute("""
        SELECT he.heat_id
        FROM heat_entries he

        JOIN heats h
            ON he.heat_id = h.id

        WHERE h.event_id = ?
        AND he.athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        affected_heats = [
            row[0]
            for row in cur.fetchall()
        ]

        # -------------------------
        # DELETE ATHLETE DATA
        # -------------------------

        # Points
        cur.execute("""
        DELETE FROM awarded_points
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        # Final result
        cur.execute("""
        DELETE FROM final_results
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        # Finalist
        cur.execute("""
        DELETE FROM finals
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        # Heat result
        cur.execute("""
        DELETE FROM results
        WHERE athlete_id = ?
        AND heat_id IN (
            SELECT id
            FROM heats
            WHERE event_id = ?
        )
        """, (
            athlete_id,
            event_id
        ))

                # Heat entry
        cur.execute("""
        DELETE FROM heat_entries
        WHERE athlete_id = ?
        AND heat_id IN (
            SELECT id
            FROM heats
            WHERE event_id = ?
        )
        """, (
            athlete_id,
            event_id
        ))

        # Measured field attempts
        cur.execute("""
        DELETE FROM field_attempts
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        # High Jump attempts
        cur.execute("""
        DELETE FROM high_jump_attempts
        WHERE event_id = ?
        AND athlete_id = ?
        """, (
            event_id,
            athlete_id
        ))

        # Event entry
        cur.execute("""
        DELETE FROM event_entries
        WHERE meeting_id = ?
        AND event_id = ?
        AND athlete_id = ?
        """, (
            meeting_id,
            event_id,
            athlete_id
        ))

        # =========================
        # NON-TRACK EVENTS
        # =========================
        # Field and High Jump use
        # specialised ranking rules.
        # Do not use the generic
        # final recalculation below.
        if event_type != "Track":

            # Invalidate official results
            # for the whole event.
            cur.execute("""
            DELETE FROM final_results
            WHERE event_id = ?
            """, (
                event_id,
            ))

            cur.execute("""
            DELETE FROM awarded_points
            WHERE event_id = ?
            """, (
                event_id,
            ))

            conn.commit()
            conn.close()

            if event_type == "High Jump":

                follow_up = (
                    "Open High Jump Results and "
                    "finalize the event again."
                )

            else:

                follow_up = (
                    "Open Field Results and save "
                    "the results again."
                )

            QMessageBox.information(
                self,
                "Entry Removed",
                (
                    f"{athlete_name} has been "
                    f"removed from {event_name}.\n\n"
                    "The official results and points "
                    "for this event were cleared so "
                    "they can be recalculated safely.\n\n"
                    f"{follow_up}"
                )
            )

            self.refresh_screen()

            return

        # -------------------------
        # RECALCULATE HEAT POSITIONS
        # -------------------------
        for heat_id in affected_heats:

            cur.execute("""
            UPDATE results
            SET position = NULL
            WHERE heat_id = ?
            """, (heat_id,))

            cur.execute("""
            SELECT e.event_type
            FROM heats h

            JOIN events e
                ON h.event_id = e.id

            WHERE h.id = ?
            """, (heat_id,))

            event_type_row = (
                cur.fetchone()
            )

            if not event_type_row:
                continue

            event_type = (
                event_type_row[0]
            )

            cur.execute("""
            SELECT
                athlete_id,
                performance
            FROM results
            WHERE heat_id = ?
            AND performance != ''
            """, (heat_id,))

            heat_results = []

            for (
                remaining_athlete_id,
                performance
            ) in cur.fetchall():

                try:

                    value = float(
                        performance
                    )

                    heat_results.append(
                        (
                            remaining_athlete_id,
                            value
                        )
                    )

                except (
                    ValueError,
                    TypeError
                ):
                    continue

            if event_type == "Track":

                heat_results.sort(
                    key=lambda x: x[1]
                )

            else:

                heat_results.sort(
                    key=lambda x: x[1],
                    reverse=True
                )

            # =========================
            # ASSIGN POSITIONS
            # WITH TIES
            # =========================
            previous_value = None
            previous_position = None

            for index, (
                remaining_athlete_id,
                value
            ) in enumerate(
                heat_results,
                start=1
            ):

                if (
                    previous_value is not None
                    and value == previous_value
                ):

                    position = (
                        previous_position
                    )

                else:

                    # Competition ranking:
                    # 1, 2, 2, 4
                    position = index

                cur.execute("""
                UPDATE results
                SET position = ?
                WHERE heat_id = ?
                AND athlete_id = ?
                """, (
                    position,
                    heat_id,
                    remaining_athlete_id
                ))

                previous_value = value
                previous_position = position

        # -------------------------
        # RE-SEED REMAINING FINALISTS
        # -------------------------
        cur.execute("""
        SELECT athlete_id
        FROM finals
        WHERE event_id = ?
        ORDER BY seed_position
        """, (event_id,))

        remaining_finalists = [
            row[0]
            for row in cur.fetchall()
        ]

        seed = 1

        for finalist_id in remaining_finalists:

            cur.execute("""
            UPDATE finals
            SET seed_position = ?
            WHERE event_id = ?
            AND athlete_id = ?
            """, (
                seed,
                event_id,
                finalist_id
            ))

            seed += 1

        # -------------------------
        # RECALCULATE FINAL POSITIONS
        # -------------------------
        cur.execute("""
        UPDATE final_results
        SET position = NULL
        WHERE event_id = ?
        """, (event_id,))

        cur.execute("""
        SELECT event_type
        FROM events
        WHERE id = ?
        """, (event_id,))

        event_type_row = (
            cur.fetchone()
        )

        if event_type_row:

            event_type = (
                event_type_row[0]
            )

            cur.execute("""
            SELECT
                athlete_id,
                performance
            FROM final_results
            WHERE event_id = ?
            AND performance != ''
            """, (event_id,))

            final_results = []

            for (
                remaining_athlete_id,
                performance
            ) in cur.fetchall():

                try:

                    value = float(
                        performance
                    )

                    final_results.append(
                        (
                            remaining_athlete_id,
                            value
                        )
                    )

                except (
                    ValueError,
                    TypeError
                ):
                    continue

            if event_type == "Track":

                final_results.sort(
                    key=lambda x: x[1]
                )

            else:

                final_results.sort(
                    key=lambda x: x[1],
                    reverse=True
                )

            # =========================
            # ASSIGN FINAL POSITIONS
            # WITH TIES
            # =========================
            previous_value = None
            previous_position = None

            for index, (
                remaining_athlete_id,
                value
            ) in enumerate(
                final_results,
                start=1
            ):

                if (
                    previous_value is not None
                    and value == previous_value
                ):

                    position = (
                        previous_position
                    )

                else:

                    position = index

                cur.execute("""
                UPDATE final_results
                SET position = ?
                WHERE event_id = ?
                AND athlete_id = ?
                """, (
                    position,
                    event_id,
                    remaining_athlete_id
                ))

                previous_value = value
                previous_position = position

        # -------------------------
        # REBUILD POINTS
        # -------------------------
        cur.execute("""
        DELETE FROM awarded_points
        WHERE event_id = ?
        """, (event_id,))

        cur.execute("""
        SELECT
            fr.athlete_id,
            fr.position,
            a.team_id
        FROM final_results fr

        JOIN athletes a
            ON fr.athlete_id = a.id

        WHERE fr.event_id = ?
        AND fr.position IS NOT NULL
        """, (event_id,))

        remaining_results = (
            cur.fetchall()
        )

        for (
            remaining_athlete_id,
            position,
            team_id
        ) in remaining_results:

            cur.execute("""
            SELECT points
            FROM points_config
            WHERE meeting_id = ?
            AND position = ?
            """, (
                meeting_id,
                position
            ))

            points_row = (
                cur.fetchone()
            )

            if not points_row:
                continue

            points = points_row[0]

            cur.execute("""
            INSERT INTO awarded_points
            (
                meeting_id,
                event_id,
                athlete_id,
                team_id,
                position,
                points
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """, (
                meeting_id,
                event_id,
                remaining_athlete_id,
                team_id,
                position,
                points
            ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Entry Removed",
            (
                f"{athlete_name} has been "
                f"removed from {event_name}."
            )
        )

        self.refresh_screen()
                                    
    def generate_heats(self):

        meeting_id = get_active_meeting()

        if not meeting_id:

            QMessageBox.warning(
                self,
                "No Active Meeting",
                "Select an active meeting first."
            )

            return

        event_id = (
            self.event_combo.currentData()
        )

        if not event_id:

            QMessageBox.warning(
                self,
                "No Event Selected",
                "Select an event first."
            )

            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # =========================
        # EVENT DETAILS
        # =========================
        cur.execute("""
        SELECT
            event_name,
            gender,
            age_group,
            event_type
        FROM events
        WHERE id = ?
        AND meeting_id = ?
        """, (
            event_id,
            meeting_id
        ))

        event = cur.fetchone()

        if not event:

            conn.close()

            QMessageBox.warning(
                self,
                "Event Not Found",
                "The selected event could not be found."
            )

            return

        event_name = event[0]
        gender = event[1]
        age_group = event[2]
        event_type = event[3]

        # =========================
        # NON-TRACK EVENTS
        # =========================
        if event_type != "Track":

            conn.close()

            if event_type == "High Jump":

                message_title = (
                    "High Jump Event"
                )

                message_text = (
                    f"{event_name} "
                    f"({gender} {age_group}) is "
                    "a High Jump event.\n\n"
                    "High Jump does not use "
                    "track heats or lanes.\n\n"
                    "High Jump uses its own "
                    "height-by-height results system."
                )

            else:

                message_title = (
                    "Field Event"
                )

                message_text = (
                    f"{event_name} "
                    f"({gender} {age_group}) is "
                    "a field event.\n\n"
                    "Field events do not use "
                    "track heats or lanes.\n\n"
                    "Enter the performances on "
                    "the Field Results screen."
                )

            QMessageBox.information(
                self,
                message_title,
                message_text
            )

            return
        
        # =========================
        # LOAD EVENT ENTRIES
        # =========================
        cur.execute("""
        SELECT
            ee.athlete_id
        FROM event_entries ee

        JOIN athletes a
            ON ee.athlete_id = a.id

        WHERE ee.meeting_id = ?
        AND ee.event_id = ?
        AND ee.status = 'Active'
        AND a.status = 'Active'

        ORDER BY ee.athlete_id
        """, (
            meeting_id,
            event_id
        ))

        athletes = [
            row[0]
            for row in cur.fetchall()
        ]

        if not athletes:

            conn.close()

            QMessageBox.warning(
                self,
                "No Entries",
                (
                   "No active athletes are entered for this event."
                )
            )

            return

                # =========================
        # DIRECT FINAL OPTION
        # =========================
        direct_final = False

        if len(athletes) <= 8:

            answer = QMessageBox.question(
                self,
                "Direct Final Available",
                (
                    f"{event_name} "
                    f"({gender} {age_group}) has "
                    f"{len(athletes)} athlete(s).\n\n"
                    "Because there are 8 or fewer "
                    "athletes, heats are not required.\n\n"
                    "YES = Create a Direct Final\n"
                    "NO = Create a normal heat\n"
                    "CANCEL = Do nothing"
                ),
                QMessageBox.Yes
                | QMessageBox.No
                | QMessageBox.Cancel
            )

            if answer == QMessageBox.Cancel:

                conn.close()
                return

            if answer == QMessageBox.Yes:

                direct_final = True

        # =========================
        # CHECK EXISTING HEATS
        # =========================
        cur.execute("""
        SELECT COUNT(*)
        FROM heats
        WHERE meeting_id = ?
        AND event_id = ?
        """, (
            meeting_id,
            event_id
        ))

        existing_heats = (
            cur.fetchone()[0]
        )

        # =========================
        # CHECK EXISTING RESULTS
        # =========================
        cur.execute("""
        SELECT COUNT(*)
        FROM results r

        JOIN heats h
            ON r.heat_id = h.id

        WHERE h.meeting_id = ?
        AND h.event_id = ?
        """, (
            meeting_id,
            event_id
        ))

        existing_results = (
            cur.fetchone()[0]
        )

        # =========================
        # CHECK FINALISTS
        # =========================
        cur.execute("""
        SELECT COUNT(*)
        FROM finals
        WHERE event_id = ?
        """, (
            event_id,
        ))

        existing_finalists = (
            cur.fetchone()[0]
        )

        # =========================
        # CHECK FINAL RESULTS
        # =========================
        cur.execute("""
        SELECT COUNT(*)
        FROM final_results
        WHERE event_id = ?
        """, (
            event_id,
        ))

        existing_final_results = (
            cur.fetchone()[0]
        )

        # =========================
        # CHECK POINTS
        # =========================
        cur.execute("""
        SELECT COUNT(*)
        FROM awarded_points
        WHERE event_id = ?
        """, (
            event_id,
        ))

        existing_points = (
            cur.fetchone()[0]
        )

        competition_data_exists = (
            existing_results > 0
            or existing_finalists > 0
            or existing_final_results > 0
            or existing_points > 0
        )

        # =========================
        # CONFIRM REGENERATION
        # =========================
        if (
            existing_heats > 0
            or competition_data_exists
        ):
            if competition_data_exists:

                message = (
                    f"{event_name} "
                    f"({gender} {age_group}) already "
                    "contains generated heats and "
                    "competition data.\n\n"
                    "Regenerating the heats will "
                    "permanently remove:\n\n"
                    f"• {existing_results} heat result(s)\n"
                    f"• {existing_finalists} finalist(s)\n"
                    f"• {existing_final_results} "
                    "final result(s)\n"
                    f"• {existing_points} "
                    "awarded point record(s)\n\n"
                    "The athlete event entries will "
                    "NOT be deleted.\n\n"
                    "The heats and lane assignments "
                    "will then be generated again.\n\n"
                    "Continue?"
                )

            else:

                message = (
                    f"{event_name} "
                    f"({gender} {age_group}) already "
                    "has generated heats.\n\n"
                    "Regenerating them will replace "
                    "the current heat and lane "
                    "assignments.\n\n"
                    "Athlete event entries will "
                    "remain unchanged.\n\n"
                    "Continue?"
                )

            answer = QMessageBox.question(
                self,
                "Regenerate Heats",
                message,
                QMessageBox.Yes
                | QMessageBox.No
            )

            if answer != QMessageBox.Yes:

                conn.close()
                return

        # =========================
        # SAFE CLEANUP
        # =========================
        try:

            # -------------------------
            # REMOVE AWARDED POINTS
            # -------------------------
            cur.execute("""
            DELETE FROM awarded_points
            WHERE event_id = ?
            """, (
                event_id,
            ))

            # -------------------------
            # REMOVE FINAL RESULTS
            # -------------------------
            cur.execute("""
            DELETE FROM final_results
            WHERE event_id = ?
            """, (
                event_id,
            ))

            # -------------------------
            # REMOVE FINALISTS
            # -------------------------
            cur.execute("""
            DELETE FROM finals
            WHERE event_id = ?
            """, (
                event_id,
            ))

            # -------------------------
            # REMOVE HEAT RESULTS
            # BEFORE REMOVING HEATS
            # -------------------------
            cur.execute("""
            DELETE FROM results
            WHERE heat_id IN (
                SELECT id
                FROM heats
                WHERE meeting_id = ?
                AND event_id = ?
            )
            """, (
                meeting_id,
                event_id
            ))

            # -------------------------
            # REMOVE HEAT ENTRIES
            # -------------------------
            cur.execute("""
            DELETE FROM heat_entries
            WHERE heat_id IN (
                SELECT id
                FROM heats
                WHERE meeting_id = ?
                AND event_id = ?
            )
            """, (
                meeting_id,
                event_id
            ))

            # -------------------------
            # REMOVE OLD HEATS
            # -------------------------
            cur.execute("""
            DELETE FROM heats
            WHERE meeting_id = ?
            AND event_id = ?
            """, (
                meeting_id,
                event_id
            ))

            # =========================
            # DIRECT FINAL
            # =========================
            if direct_final:

                seed_position = 1

                for athlete_id in athletes:

                    cur.execute("""
                    INSERT INTO finals
                    (
                        event_id,
                        athlete_id,
                        seed_position
                    )
                    VALUES (?, ?, ?)
                    """, (
                        event_id,
                        athlete_id,
                        seed_position
                    ))

                    seed_position += 1

            # =========================
            # NORMAL HEATS
            # =========================
            else:

                MAX_PER_HEAT = 8

                lane_order = [
                    4, 5, 3, 6,
                    2, 7, 1, 8
                ]

                # =========================
                # CALCULATE BALANCED HEATS
                # =========================

                athlete_count = len(
                    athletes
                )

                heat_count = (
                    athlete_count
                    + MAX_PER_HEAT
                    - 1
                ) // MAX_PER_HEAT

                base_heat_size = (
                    athlete_count
                    // heat_count
                )

                extra_athletes = (
                    athlete_count
                    % heat_count
                )

                heat_sizes = []

                for heat_index in range(
                    heat_count
                ):

                    heat_size = (
                        base_heat_size
                    )

                    # Spread extra athletes
                    # across the first heats.
                    if (
                        heat_index
                        < extra_athletes
                    ):

                        heat_size += 1

                    heat_sizes.append(
                        heat_size
                    )

                # =========================
                # CREATE BALANCED HEATS
                # =========================

                athlete_index = 0

                for heat_index, heat_size in enumerate(
                    heat_sizes,
                    start=1
                ):

                    cur.execute("""
                    INSERT INTO heats
                    (
                        meeting_id,
                        event_id,
                        heat_number
                    )
                    VALUES (?, ?, ?)
                    """, (
                        meeting_id,
                        event_id,
                        heat_index
                    ))

                    new_heat_id = (
                        cur.lastrowid
                    )

                    # =========================
                    # ASSIGN LANES
                    # =========================

                    for lane_index in range(
                        heat_size
                    ):

                        athlete_id = (
                            athletes[
                                athlete_index
                            ]
                        )

                        lane = (
                            lane_order[
                                lane_index
                            ]
                        )

                        cur.execute("""
                        INSERT INTO heat_entries
                        (
                            heat_id,
                            athlete_id,
                            lane_number
                        )
                        VALUES (?, ?, ?)
                        """, (
                            new_heat_id,
                            athlete_id,
                            lane
                        ))

                        athlete_index += 1

            # Save all cleanup and newly
            # generated competition data
            conn.commit()

        except sqlite3.Error as error:

            conn.rollback()
            conn.close()

            QMessageBox.critical(
                self,
                "Heat Generation Failed",
                (
                    "The heats could not be "
                    "regenerated.\n\n"
                    "No changes were saved.\n\n"
                    f"Database error:\n{error}"
                )
            )

            return

        conn.close()

        # =========================
        # SUCCESS
        # =========================
        if direct_final:

            QMessageBox.information(
                self,
                "Direct Final Created",
                (
                    f"A direct final has been "
                    f"created for:\n\n"
                    f"{event_name} "
                    f"({gender} {age_group})\n\n"
                    f"Athletes: {len(athletes)}\n\n"
                    "No preliminary heat is required."
                )
            )

        else:

            QMessageBox.information(
                self,
                "Heats Generated",
                (
                    f"Heats successfully generated "
                    f"for:\n\n"
                    f"{event_name} "
                    f"({gender} {age_group})\n\n"
                    f"Athletes: {len(athletes)}"
                )
            )

        self.refresh_screen()

    def refresh_screen(self):

        self.load_events()
        self.load_athletes()

        self.entry_list.clear()

        self.load_entries()
        self.load_heats()

        self.load_heat_choices()
        self.load_heat_for_edit()
        

    def filter_athletes_by_event(self):

        self.athlete_combo.clear()

        meeting_id = get_active_meeting()
        event_id = self.event_combo.currentData()

        if not meeting_id or not event_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Get the selected event's
        # gender and age group
        cur.execute("""
        SELECT
            gender,
            age_group
        FROM events
        WHERE id = ?
        AND meeting_id = ?
        """, (
            event_id,
            meeting_id
        ))

        event = cur.fetchone()

        if not event:

            conn.close()
            return

        gender = event[0]
        age_group = event[1]

        # Only show ACTIVE athletes
        # matching BOTH gender and age group
        cur.execute("""
        SELECT
            id,
            full_name
        FROM athletes
        WHERE meeting_id = ?
        AND gender = ?
        AND age_group = ?
        AND status = 'Active'
        ORDER BY full_name
        """, (
            meeting_id,
            gender,
            age_group
        ))

        athletes = cur.fetchall()

        conn.close()

        for athlete_id, name in athletes:

            self.athlete_combo.addItem(
                name,
                athlete_id
            )

    def load_events(self):

        self.event_combo.clear()

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            id,
            event_name,
            gender,
            age_group
        FROM events
        WHERE meeting_id = ?
        ORDER BY event_name
        """, (meeting_id,))

        rows = cur.fetchall()

        conn.close()

        for event_id, name, gender, age_group in rows:
            self.event_combo.addItem(
                f"{name} ({gender} {age_group})",
                event_id
            )        

    def load_heats(self):

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Get all heats for this meeting
        cur.execute("""
        SELECT
            h.id,
            h.heat_number,
            e.event_name,
            e.gender,
            e.age_group
        FROM heats h
        JOIN events e
            ON h.event_id = e.id
        WHERE h.meeting_id = ?
        ORDER BY
            e.event_name,
            e.gender,
            e.age_group,
            h.heat_number
        """, (meeting_id,))

        heats = cur.fetchall()

        for heat_id, heat_no, event_name, gender, age_group in heats:

            self.entry_list.addItem(
                f"{event_name} ({gender} {age_group}) - Heat {heat_no}"
            )

            cur.execute("""
            SELECT
                a.full_name,
                he.lane_number
            FROM heat_entries he
            JOIN athletes a
                ON he.athlete_id = a.id
            WHERE he.heat_id = ?
            ORDER BY he.lane_number
            """, (heat_id,))

            athletes = cur.fetchall()

            for name, lane in athletes:
                self.entry_list.addItem(
                    f"    Lane {lane}: {name}"
                )

            self.entry_list.addItem("")

        conn.close()

    # =========================
    # LOAD HEAT CHOICES
    # =========================
    def load_heat_choices(self):

        self.heat_combo.clear()

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            h.id,
            e.event_name,
            e.gender,
            e.age_group,
            h.heat_number
        FROM heats h

        JOIN events e
            ON h.event_id = e.id

        WHERE h.meeting_id = ?

        ORDER BY
            e.event_name,
            e.gender,
            e.age_group,
            h.heat_number
        """, (meeting_id,))

        rows = cur.fetchall()

        conn.close()

        for (
            heat_id,
            event_name,
            gender,
            age_group,
            heat_number
        ) in rows:

            self.heat_combo.addItem(
                (
                    f"{event_name} "
                    f"({gender} {age_group}) "
                    f"- Heat {heat_number}"
                ),
                heat_id
            )


    # =========================
    # LOAD HEAT FOR EDITING
    # =========================
    def load_heat_for_edit(self):

        self.heat_table.setRowCount(0)

        heat_id = self.heat_combo.currentData()

        if not heat_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            he.id,
            he.lane_number,
            a.full_name
        FROM heat_entries he

        JOIN athletes a
            ON he.athlete_id = a.id

        WHERE he.heat_id = ?

        ORDER BY he.lane_number
        """, (heat_id,))

        rows = cur.fetchall()

        conn.close()

        self.heat_table.setRowCount(
            len(rows)
        )

        for row_number, (
            heat_entry_id,
            lane,
            athlete_name
        ) in enumerate(rows):

            lane_item = QTableWidgetItem(
                str(lane)
            )

            self.heat_table.setItem(
                row_number,
                0,
                lane_item
            )

            athlete_item = QTableWidgetItem(
                athlete_name
            )

            athlete_item.setData(
                Qt.UserRole,
                heat_entry_id
            )

            athlete_item.setFlags(
                athlete_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.heat_table.setItem(
                row_number,
                1,
                athlete_item
            )

    # =========================
    # LOAD DESTINATION HEATS
    # =========================
    def load_destination_heats(self):

        self.move_heat_combo.blockSignals(
            True
        )

        self.move_heat_combo.clear()

        source_heat_id = (
            self.heat_combo.currentData()
        )

        if not source_heat_id:

            self.move_heat_combo.blockSignals(
                False
            )

            self.move_lane_combo.clear()

            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Find which event the
        # selected heat belongs to
        cur.execute("""
        SELECT event_id
        FROM heats
        WHERE id = ?
        """, (
            source_heat_id,
        ))

        row = cur.fetchone()

        if not row:

            conn.close()

            self.move_heat_combo.blockSignals(
                False
            )

            return

        event_id = row[0]

        # Only allow other heats
        # belonging to the SAME event
        cur.execute("""
        SELECT
            id,
            heat_number
        FROM heats
        WHERE event_id = ?
        AND id != ?
        ORDER BY heat_number
        """, (
            event_id,
            source_heat_id
        ))

        heats = cur.fetchall()

        conn.close()

        if not heats:

            self.move_heat_combo.addItem(
                "No other heat available",
                None
            )

        else:

            for (
                heat_id,
                heat_number
            ) in heats:

                self.move_heat_combo.addItem(
                    f"Heat {heat_number}",
                    heat_id
                )

        self.move_heat_combo.blockSignals(
            False
        )

        self.load_destination_lanes()

            # =========================
    # LOAD FREE DESTINATION LANES
    # =========================
    def load_destination_lanes(self):

        self.move_lane_combo.clear()

        destination_heat_id = (
            self.move_heat_combo.currentData()
        )

        if not destination_heat_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT lane_number
        FROM heat_entries
        WHERE heat_id = ?
        """, (
            destination_heat_id,
        ))

        used_lanes = {
            row[0]
            for row in cur.fetchall()
        }

        conn.close()

        free_lanes = [
            lane
            for lane in range(1, 9)
            if lane not in used_lanes
        ]

        if not free_lanes:

            self.move_lane_combo.addItem(
                "No free lanes",
                None
            )

            return

        for lane in free_lanes:

            self.move_lane_combo.addItem(
                f"Lane {lane}",
                lane
            )

    # =========================
    # RECALCULATE HEAT POSITIONS
    # =========================
    def recalculate_heat_positions(
        self,
        cur,
        heat_id
    ):

        # Clear old positions first
        cur.execute("""
        UPDATE results
        SET position = NULL
        WHERE heat_id = ?
        """, (
            heat_id,
        ))

        # Determine event type
        cur.execute("""
        SELECT e.event_type
        FROM heats h

        JOIN events e
            ON h.event_id = e.id

        WHERE h.id = ?
        """, (
            heat_id,
        ))

        row = cur.fetchone()

        if not row:
            return

        event_type = row[0]

        cur.execute("""
        SELECT
            athlete_id,
            performance
        FROM results
        WHERE heat_id = ?
        AND performance != ''
        """, (
            heat_id,
        ))

        performances = []

        for (
            athlete_id,
            performance
        ) in cur.fetchall():

            try:

                value = float(
                    performance
                )

                performances.append(
                    (
                        athlete_id,
                        value
                    )
                )

            except (
                ValueError,
                TypeError
            ):

                continue

        # Track:
        # lowest time wins
        if event_type == "Track":

            performances.sort(
                key=lambda x: x[1]
            )

        # Field:
        # highest performance wins
        else:

            performances.sort(
                key=lambda x: x[1],
                reverse=True
            )

            # =========================
        # ASSIGN POSITIONS
        # WITH TIES
        # =========================
        previous_value = None
        previous_position = None

        for index, (
            athlete_id,
            value
        ) in enumerate(
            performances,
            start=1
        ):

            if (
                previous_value is not None
                and value == previous_value
            ):

                position = (
                    previous_position
                )

            else:

                # Competition ranking:
                # 1, 2, 2, 4
                position = index

            cur.execute("""
            UPDATE results
            SET position = ?
            WHERE heat_id = ?
            AND athlete_id = ?
            """, (
                position,
                heat_id,
                athlete_id
            ))

            previous_value = value
            previous_position = position
            
     # =========================
    # MOVE ATHLETE BETWEEN HEATS
    # =========================
    def move_athlete_between_heats(self):

        source_heat_id = (
            self.heat_combo.currentData()
        )

        destination_heat_id = (
            self.move_heat_combo.currentData()
        )

        destination_lane = (
            self.move_lane_combo.currentData()
        )

        if not source_heat_id:

            QMessageBox.warning(
                self,
                "No Source Heat",
                "Select a source heat first."
            )

            return

        if not destination_heat_id:

            QMessageBox.warning(
                self,
                "No Destination Heat",
                (
                    "There is no other heat "
                    "available for this event."
                )
            )

            return

        if destination_lane is None:

            QMessageBox.warning(
                self,
                "No Free Lane",
                (
                    "Select an available "
                    "destination lane."
                )
            )

            return

        selected_row = (
            self.heat_table.currentRow()
        )

        if selected_row < 0:

            QMessageBox.warning(
                self,
                "No Athlete Selected",
                (
                    "Click an athlete in the "
                    "heat table first."
                )
            )

            return

        athlete_item = (
            self.heat_table.item(
                selected_row,
                1
            )
        )

        if not athlete_item:

            QMessageBox.warning(
                self,
                "No Athlete Selected",
                (
                    "Select an athlete row "
                    "from the heat table."
                )
            )

            return

        heat_entry_id = (
            athlete_item.data(
                Qt.UserRole
            )
        )

        if not heat_entry_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Get athlete
        cur.execute("""
        SELECT
            he.athlete_id,
            a.full_name
        FROM heat_entries he

        JOIN athletes a
            ON he.athlete_id = a.id

        WHERE he.id = ?
        """, (
            heat_entry_id,
        ))

        athlete = cur.fetchone()

        if not athlete:

            conn.close()
            return

        athlete_id = athlete[0]
        athlete_name = athlete[1]

        # Get source heat details
        cur.execute("""
        SELECT
            event_id,
            heat_number
        FROM heats
        WHERE id = ?
        """, (
            source_heat_id,
        ))

        source_heat = cur.fetchone()

        # Get destination heat details
        cur.execute("""
        SELECT
            event_id,
            heat_number
        FROM heats
        WHERE id = ?
        """, (
            destination_heat_id,
        ))

        destination_heat = (
            cur.fetchone()
        )

        if (
            not source_heat
            or not destination_heat
        ):

            conn.close()
            return

        source_event_id = (
            source_heat[0]
        )

        source_heat_number = (
            source_heat[1]
        )

        destination_event_id = (
            destination_heat[0]
        )

        destination_heat_number = (
            destination_heat[1]
        )

        # Extra safety check:
        # both heats MUST belong
        # to the same event
        if (
            source_event_id
            != destination_event_id
        ):

            conn.close()

            QMessageBox.warning(
                self,
                "Invalid Heat",
                (
                    "Athletes can only be moved "
                    "between heats of the same event."
                )
            )

            return

        # Make sure destination lane
        # is still free
        cur.execute("""
        SELECT id
        FROM heat_entries
        WHERE heat_id = ?
        AND lane_number = ?
        """, (
            destination_heat_id,
            destination_lane
        ))

        if cur.fetchone():

            conn.close()

            QMessageBox.warning(
                self,
                "Lane Already Used",
                (
                    f"Lane {destination_lane} "
                    "is already occupied."
                )
            )

            self.load_destination_lanes()

            return

        # Make sure athlete is not
        # already in destination heat
        cur.execute("""
        SELECT id
        FROM heat_entries
        WHERE heat_id = ?
        AND athlete_id = ?
        """, (
            destination_heat_id,
            athlete_id
        ))

        if cur.fetchone():

            conn.close()

            QMessageBox.warning(
                self,
                "Duplicate Athlete",
                (
                    "This athlete is already "
                    "in the destination heat."
                )
            )

            return

        # Also check for an orphaned
        # result in destination heat
        cur.execute("""
        SELECT id
        FROM results
        WHERE heat_id = ?
        AND athlete_id = ?
        """, (
            destination_heat_id,
            athlete_id
        ))

        if cur.fetchone():

            conn.close()

            QMessageBox.warning(
                self,
                "Existing Result",
                (
                    "A result already exists for "
                    "this athlete in the "
                    "destination heat."
                )
            )

            return

        # Does athlete already have
        # a recorded result?
        cur.execute("""
        SELECT performance
        FROM results
        WHERE heat_id = ?
        AND athlete_id = ?
        """, (
            source_heat_id,
            athlete_id
        ))

        existing_result = (
            cur.fetchone()
        )

        if existing_result:

            performance = (
                existing_result[0]
            )

            answer = QMessageBox.question(
                self,
                "Move Athlete With Result",
                (
                    f"{athlete_name} already has "
                    f"a result of {performance} "
                    f"in Heat "
                    f"{source_heat_number}.\n\n"
                    "The result will move with "
                    "the athlete and positions "
                    "in both heats will be "
                    "recalculated.\n\n"
                    f"Move to Heat "
                    f"{destination_heat_number}, "
                    f"Lane {destination_lane}?"
                ),
                QMessageBox.Yes
                | QMessageBox.No
            )

        else:

            answer = QMessageBox.question(
                self,
                "Move Athlete",
                (
                    f"Move {athlete_name} "
                    f"from Heat "
                    f"{source_heat_number} "
                    f"to Heat "
                    f"{destination_heat_number}, "
                    f"Lane {destination_lane}?"
                ),
                QMessageBox.Yes
                | QMessageBox.No
            )

        if answer != QMessageBox.Yes:

            conn.close()
            return

        # -------------------------
        # MOVE HEAT ENTRY
        # -------------------------
        cur.execute("""
        UPDATE heat_entries
        SET
            heat_id = ?,
            lane_number = ?
        WHERE id = ?
        """, (
            destination_heat_id,
            destination_lane,
            heat_entry_id
        ))

        # -------------------------
        # MOVE RESULT TOO
        # -------------------------
        if existing_result:

            cur.execute("""
            UPDATE results
            SET heat_id = ?
            WHERE heat_id = ?
            AND athlete_id = ?
            """, (
                destination_heat_id,
                source_heat_id,
                athlete_id
            ))

        # Recalculate rankings
        # in BOTH heats
        self.recalculate_heat_positions(
            cur,
            source_heat_id
        )

        self.recalculate_heat_positions(
            cur,
            destination_heat_id
        )

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Athlete Moved",
            (
                f"{athlete_name} moved to "
                f"Heat {destination_heat_number}, "
                f"Lane {destination_lane}."
            )
        )

        # -------------------------
        # REFRESH DISPLAY
        # -------------------------
        self.entry_list.clear()

        self.load_entries()
        self.load_heats()

        self.load_heat_choices()

        source_index = (
            self.heat_combo.findData(
                source_heat_id
            )
        )

        if source_index >= 0:

            self.heat_combo.setCurrentIndex(
                source_index
            )

        self.load_heat_for_edit()
        self.load_destination_heats()           

    # =========================
    # SAVE LANE CHANGES
    # =========================
    def save_lane_changes(self):

        heat_id = self.heat_combo.currentData()

        if not heat_id:

            QMessageBox.warning(
                self,
                "No Heat Selected",
                "Select a heat first."
            )

            return

        updates = []
        used_lanes = set()

        for row in range(
            self.heat_table.rowCount()
        ):

            lane_item = self.heat_table.item(
                row,
                0
            )

            athlete_item = self.heat_table.item(
                row,
                1
            )

            if not lane_item or not athlete_item:
                continue

            lane_text = lane_item.text().strip()

            try:
                lane = int(lane_text)

            except ValueError:

                QMessageBox.warning(
                    self,
                    "Invalid Lane",
                    (
                        f"'{lane_text}' is not "
                        "a valid lane number."
                    )
                )

                return

            if lane < 1 or lane > 8:

                QMessageBox.warning(
                    self,
                    "Invalid Lane",
                    (
                        "Lane numbers must be "
                        "between 1 and 8."
                    )
                )

                return

            if lane in used_lanes:

                QMessageBox.warning(
                    self,
                    "Duplicate Lane",
                    (
                        f"Lane {lane} has been "
                        "assigned more than once."
                    )
                )

                return

            used_lanes.add(lane)

            heat_entry_id = athlete_item.data(
                Qt.UserRole
            )

            updates.append(
                (
                    lane,
                    heat_entry_id
                )
            )

        if not updates:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT COUNT(*)
        FROM results
        WHERE heat_id = ?
        """, (heat_id,))

        result_count = cur.fetchone()[0]

        if result_count > 0:

            answer = QMessageBox.question(
                self,
                "Results Already Exist",
                (
                    "Results have already been "
                    "entered for this heat.\n\n"
                    "Changing lanes will not delete "
                    "the results, but the saved lane "
                    "assignments will change.\n\n"
                    "Continue?"
                ),
                QMessageBox.Yes
                | QMessageBox.No
            )

            if answer != QMessageBox.Yes:

                conn.close()
                return

        for lane, heat_entry_id in updates:

            cur.execute("""
            UPDATE heat_entries
            SET lane_number = ?
            WHERE id = ?
            """, (
                lane,
                heat_entry_id
            ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Lanes Updated",
            "Lane assignments saved successfully."
        )

        self.load_heat_for_edit()
        self.load_entries()
        self.load_heats()

# =========================
# RESULTS SCREEN
# =========================
class ResultsScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        # =========================
        # TITLE
        # =========================
        title = QLabel("Results")
        title.setAlignment(Qt.AlignCenter)

        layout.addWidget(title)

        # =========================
        # HEAT SELECTION
        # =========================
        layout.addWidget(
            QLabel("Select Heat")
        )

        self.heat_combo = QComboBox()

        layout.addWidget(
            self.heat_combo
        )

        load_button = QPushButton(
            "Load Heat"
        )

        load_button.clicked.connect(
            self.load_heat
        )

        layout.addWidget(
            load_button
        )

        # =========================
        # INSTRUCTIONS
        # =========================
        instruction = QLabel(
            "Enter a numeric performance, "
            "or use DNS / DNF / DQ."
        )

        layout.addWidget(
            instruction
        )

        # =========================
        # RESULTS TABLE
        # =========================
        self.table = QTableWidget()

        self.table.setColumnCount(3)

        self.table.setHorizontalHeaderLabels([
            "Athlete",
            "Performance",
            "Position"
        ])

        self.table.setMinimumHeight(300)

        layout.addWidget(
            self.table,
            1
        )

        # =========================
        # SAVE RESULTS
        # =========================
        save_button = QPushButton(
            "Save Results"
        )

        save_button.clicked.connect(
            self.save_results
        )

        layout.addWidget(
            save_button
        )

        # =========================
        # GENERATE FINALISTS
        # =========================

        layout.addWidget(
            QLabel("Final Qualification Method")
        )

        self.qualification_combo = QComboBox()

        self.qualification_combo.addItems([
            "Fastest / Best 8 Overall",
            "Top 2 From Each Heat + Next Best",
            "Manual Selection (Admin Override)"
        ])

        layout.addWidget(
            self.qualification_combo
        )

        generate_finals_button = QPushButton(
            "Generate Finalists"
        )

        generate_finals_button.clicked.connect(
            self.generate_finalists
        )

        layout.addWidget(
            generate_finals_button
        )

        self.setLayout(layout)

        # Load available heats
        self.load_heats()

        # Automatically display
        # the first available heat
        self.load_heat()
        
    def load_heats(self):

        self.heat_combo.clear()

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        meeting_id = get_active_meeting()

        if not meeting_id:
            conn.close()
            return

        cur.execute("""
        SELECT
            h.id,
            e.event_name,
            e.gender,
            e.age_group,
            h.heat_number
        FROM heats h
        JOIN events e
            ON h.event_id = e.id
        WHERE h.meeting_id = ?
        AND e.event_type = 'Track'
        ORDER BY
            e.event_name,
            e.gender,
            e.age_group,
            h.heat_number
        """, (meeting_id,))

        rows = cur.fetchall()

        conn.close()

        for heat_id, event_name, gender, age_group, heat_no in rows:

            self.heat_combo.addItem(
                f"{event_name} ({gender} {age_group}) - Heat {heat_no}",
                heat_id
            )

    def load_heat(self):

        self.table.setRowCount(0)

        heat_id = self.heat_combo.currentData()

        if not heat_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            a.id,
            a.full_name,
            he.lane_number
        FROM heat_entries he

        JOIN athletes a
            ON he.athlete_id = a.id

        WHERE he.heat_id = ?

        ORDER BY he.lane_number
        """, (heat_id,))

        athletes = cur.fetchall()

        cur.execute("""
        SELECT
            athlete_id,
            performance,
            position
        FROM results
        WHERE heat_id = ?
        """, (heat_id,))

        results = {
            r[0]: (r[1], r[2])
            for r in cur.fetchall()
        }

        conn.close()

        self.table.setRowCount(
            len(athletes)
        )

        for row, athlete in enumerate(
            athletes
        ):

            athlete_id = athlete[0]
            athlete_name = athlete[1]
            lane = athlete[2]

            # -------------------------
            # ATHLETE
            # -------------------------
            athlete_item = QTableWidgetItem(
                f"Lane {lane} - {athlete_name}"
            )

            athlete_item.setData(
                Qt.UserRole,
                athlete_id
            )

            athlete_item.setFlags(
                athlete_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.table.setItem(
                row,
                0,
                athlete_item
            )

            performance, position = (
                results.get(
                    athlete_id,
                    ("", "")
                )
            )

            if performance is None:
                performance = ""

            if position is None:
                position = ""

            # -------------------------
            # PERFORMANCE
            # -------------------------
            performance_item = (
                QTableWidgetItem(
                    str(performance)
                )
            )

            self.table.setItem(
                row,
                1,
                performance_item
            )

            # -------------------------
            # POSITION
            # -------------------------
            position_item = (
                QTableWidgetItem(
                    str(position)
                )
            )

            position_item.setFlags(
                position_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.table.setItem(
                row,
                2,
                position_item
            )

    def save_results(self):

        heat_id = (
            self.heat_combo.currentData()
        )

        if not heat_id:

            QMessageBox.warning(
                self,
                "No Heat Selected",
                "Please select a heat first."
            )

            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        valid_codes = {
            "DNS",
            "DNF",
            "DQ"
        }

        for row in range(
            self.table.rowCount()
        ):

            athlete_item = (
                self.table.item(
                    row,
                    0
                )
            )

            if not athlete_item:
                continue

            athlete_id = athlete_item.data(
                Qt.UserRole
            )

            performance_item = (
                self.table.item(
                    row,
                    1
                )
            )

            performance = ""

            if performance_item:

                performance = (
                    performance_item.text()
                    .strip()
                )

            # -------------------------
            # VALIDATE PERFORMANCE
            # -------------------------

            if performance:

                upper_value = (
                    performance.upper()
                )

                # Special race status
                if upper_value in valid_codes:

                    performance = upper_value

                else:

                    try:

                        float(performance)

                    except ValueError:

                        athlete_name = (
                            athlete_item.text()
                        )

                        conn.close()

                        QMessageBox.warning(
                            self,
                            "Invalid Performance",
                            (
                                f"Invalid result for:\n\n"
                                f"{athlete_name}\n\n"
                                "Enter a number or one "
                                "of these codes:\n\n"
                                "DNS = Did Not Start\n"
                                "DNF = Did Not Finish\n"
                                "DQ = Disqualified"
                            )
                        )

                        return

            cur.execute("""
            INSERT INTO results
            (
                heat_id,
                athlete_id,
                performance
            )
            VALUES (?, ?, ?)

            ON CONFLICT(
                heat_id,
                athlete_id
            )

            DO UPDATE SET
                performance =
                    excluded.performance
            """, (
                heat_id,
                athlete_id,
                performance
            ))

        conn.commit()
        conn.close()

        # Recalculate positions
        self.calculate_positions(
            heat_id
        )

        self.load_heat()

        QMessageBox.information(
            self,
            "Saved",
            "Results saved successfully."
        )
        
    def generate_finalists(self):

        heat_id = (
            self.heat_combo.currentData()
        )

        if not heat_id:

            QMessageBox.warning(
                self,
                "No Heat Selected",
                "Select a heat for the event first."
            )

            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # =========================
        # FIND EVENT
        # =========================
        cur.execute("""
        SELECT event_id
        FROM heats
        WHERE id = ?
        """, (
            heat_id,
        ))

        row = cur.fetchone()

        if not row:

            conn.close()
            return

        event_id = row[0]

        # =========================
        # EVENT DETAILS
        # =========================
        cur.execute("""
        SELECT
            event_name,
            event_type,
            gender,
            age_group
        FROM events
        WHERE id = ?
        """, (
            event_id,
        ))

        event = cur.fetchone()

        if not event:

            conn.close()
            return

        event_name = event[0]
        event_type = event[1]
        gender = event[2]
        age_group = event[3]

        # =========================
        # CHECK EXISTING FINAL
        # =========================
        cur.execute("""
        SELECT COUNT(*)
        FROM final_results
        WHERE event_id = ?
        """, (
            event_id,
        ))

        existing_final_results = (
            cur.fetchone()[0]
        )

        if existing_final_results > 0:

            answer = QMessageBox.question(
                self,
                "Final Already Has Results",
                (
                    f"{event_name} "
                    f"({gender} {age_group}) already "
                    "has final results.\n\n"
                    "Generating the finalists again "
                    "will remove the existing final "
                    "results and points.\n\n"
                    "Continue?"
                ),
                QMessageBox.Yes
                | QMessageBox.No
            )

            if answer != QMessageBox.Yes:

                conn.close()
                return

        # =========================
        # LOAD ALL HEAT RESULTS
        # =========================
        cur.execute("""
        SELECT
            h.id,
            h.heat_number,
            r.athlete_id,
            a.full_name,
            r.performance
        FROM results r

        JOIN heats h
            ON r.heat_id = h.id

        JOIN athletes a
            ON r.athlete_id = a.id

        WHERE h.event_id = ?

        ORDER BY
            h.heat_number,
            r.athlete_id
        """, (
            event_id,
        ))

        rows = cur.fetchall()

        heat_results = {}

        for (
            result_heat_id,
            heat_number,
            athlete_id,
            athlete_name,
            performance
        ) in rows:

            try:

                value = float(
                    performance
                )

            except (
                ValueError,
                TypeError
            ):

                # DNS / DNF / DQ / blank
                # cannot qualify.
                continue

            if result_heat_id not in heat_results:

                heat_results[
                    result_heat_id
                ] = []

            heat_results[
                result_heat_id
            ].append(
                (
                    athlete_id,
                    value,
                    heat_number,
                    athlete_name
                )
            )

        if not heat_results:

            conn.close()

            QMessageBox.warning(
                self,
                "No Valid Results",
                (
                    "No numeric performances "
                    "are available for this event."
                )
            )

            return

        # =========================
        # SORT HELPER
        # =========================
        def sort_results(
            results_to_sort
        ):

            if event_type == "Track":

                return sorted(
                    results_to_sort,
                    key=lambda x: x[1]
                )

            return sorted(
                results_to_sort,
                key=lambda x: x[1],
                reverse=True
            )

        # Sort every heat
        for result_heat_id in heat_results:

            heat_results[
                result_heat_id
            ] = sort_results(
                heat_results[
                    result_heat_id
                ]
            )

        # =========================
        # MANUAL CHOICE HELPER
        # =========================
        def choose_candidates(
            candidates,
            places_needed,
            window_title,
            instruction
        ):

            if places_needed <= 0:
                return []

            if len(candidates) <= places_needed:

                return list(
                    candidates
                )

            selected = []
            available = list(
                candidates
            )

            while len(selected) < places_needed:

                options = []

                for candidate in available:

                    athlete_id = candidate[0]
                    performance = candidate[1]
                    heat_number = candidate[2]
                    athlete_name = candidate[3]

                    options.append(
                        (
                            f"{athlete_name} | "
                            f"{performance:g} | "
                            f"Heat {heat_number} | "
                            f"ID {athlete_id}"
                        )
                    )

                choice, ok = (
                    QInputDialog.getItem(
                        self,
                        window_title,
                        (
                            instruction
                            + "\n\n"
                            + f"Choose athlete "
                            f"{len(selected) + 1} "
                            f"of {places_needed}:"
                        ),
                        options,
                        0,
                        False
                    )
                )

                if not ok:

                    return None

                selected_index = (
                    options.index(
                        choice
                    )
                )

                selected.append(
                    available.pop(
                        selected_index
                    )
                )

            return selected

        # =========================
        # CUTOFF TIE HELPER
        # =========================
        def select_with_cutoff_tie(
            sorted_candidates,
            places_needed,
            title,
            description
        ):

            if places_needed <= 0:
                return []

            if (
                len(sorted_candidates)
                <= places_needed
            ):

                return list(
                    sorted_candidates
                )

            cutoff_value = (
                sorted_candidates[
                    places_needed - 1
                ][1]
            )

            guaranteed = []
            tied_at_cutoff = []

            for candidate in sorted_candidates:

                value = candidate[1]

                if event_type == "Track":

                    better = (
                        value < cutoff_value
                    )

                else:

                    better = (
                        value > cutoff_value
                    )

                if better:

                    guaranteed.append(
                        candidate
                    )

                elif value == cutoff_value:

                    tied_at_cutoff.append(
                        candidate
                    )

            remaining_places = (
                places_needed
                - len(guaranteed)
            )

            # No ambiguous tie
            if (
                len(tied_at_cutoff)
                <= remaining_places
            ):

                return (
                    guaranteed
                    + tied_at_cutoff
                )

            # =========================
            # ACTUAL CUTOFF TIE
            # =========================
            tied_names = "\n".join(
                (
                    f"• {candidate[3]} "
                    f"({candidate[1]:g}) "
                    f"- Heat {candidate[2]}"
                )
                for candidate
                in tied_at_cutoff
            )

            QMessageBox.information(
                self,
                "Qualification Tie",
                (
                    f"{description}\n\n"
                    "The following athletes are "
                    "tied for the remaining "
                    "qualification place(s):\n\n"
                    f"{tied_names}\n\n"
                    "You will now choose manually "
                    "which athlete(s) advance."
                )
            )

            manual_choices = (
                choose_candidates(
                    tied_at_cutoff,
                    remaining_places,
                    title,
                    (
                        "Select the athlete who "
                        "should advance from the tie."
                    )
                )
            )

            if manual_choices is None:

                return None

            return (
                guaranteed
                + manual_choices
            )

        qualification_method = (
            self.qualification_combo
            .currentText()
        )

        final_size = 8

        finalists = []

        # =========================
        # ALL VALID RESULTS
        # =========================
        all_results = []

        for heat in heat_results.values():

            all_results.extend(
                heat
            )

        all_results = sort_results(
            all_results
        )

        # =========================
        # METHOD 1
        # FASTEST / BEST 8
        # =========================
        if qualification_method == (
            "Fastest / Best 8 Overall"
        ):

            finalists = (
                select_with_cutoff_tie(
                    all_results,
                    final_size,
                    "Choose Final Qualifier",
                    (
                        "There is a tie at the "
                        "8-athlete qualification "
                        "cutoff."
                    )
                )
            )

            if finalists is None:

                conn.close()
                return

        # =========================
        # METHOD 2
        # TOP 2 EACH HEAT
        # + NEXT BEST
        # =========================
        elif qualification_method == (
            "Top 2 From Each Heat + Next Best"
        ):

            heat_count = len(
                heat_results
            )

            if heat_count > 4:

                conn.close()

                QMessageBox.warning(
                    self,
                    "Too Many Heats",
                    (
                        f"This event has "
                        f"{heat_count} heats.\n\n"
                        "An 8-athlete final cannot "
                        "take the top 2 automatically "
                        "from more than 4 heats.\n\n"
                        "Use another qualification "
                        "method."
                    )
                )

                return

            qualified_ids = set()

            # -------------------------
            # TOP 2 FROM EACH HEAT
            # -------------------------
            for result_heat_id in sorted(
                heat_results.keys()
            ):

                heat = heat_results[
                    result_heat_id
                ]

                places_from_heat = min(
                    2,
                    len(heat)
                )

                automatic = (
                    select_with_cutoff_tie(
                        heat,
                        places_from_heat,
                        "Choose Heat Qualifier",
                        (
                            "There is a tie for "
                            "an automatic Top 2 "
                            "qualification position."
                        )
                    )
                )

                if automatic is None:

                    conn.close()
                    return

                for result in automatic:

                    athlete_id = result[0]

                    if athlete_id not in qualified_ids:

                        finalists.append(
                            result
                        )

                        qualified_ids.add(
                            athlete_id
                        )

            # -------------------------
            # NEXT BEST ATHLETES
            # -------------------------
            remaining = []

            for result in all_results:

                athlete_id = result[0]

                if athlete_id not in qualified_ids:

                    remaining.append(
                        result
                    )

            spaces_left = (
                final_size
                - len(finalists)
            )

            if spaces_left > 0:

                next_best = (
                    select_with_cutoff_tie(
                        remaining,
                        spaces_left,
                        "Choose Final Qualifier",
                        (
                            "There is a tie for "
                            "one of the remaining "
                            "next-best qualification "
                            "positions."
                        )
                    )
                )

                if next_best is None:

                    conn.close()
                    return

                finalists.extend(
                    next_best
                )

        # =========================
        # METHOD 3
        # MANUAL ADMIN OVERRIDE
        # =========================
        else:

            places_to_select = min(
                final_size,
                len(all_results)
            )

            finalists = (
                choose_candidates(
                    all_results,
                    places_to_select,
                    "Manual Finalist Selection",
                    (
                        "Admin Override:\n"
                        "Select the athletes who "
                        "should advance to the final."
                    )
                )
            )

            if finalists is None:

                conn.close()
                return

        if not finalists:

            conn.close()

            QMessageBox.warning(
                self,
                "No Finalists",
                "No finalists were selected."
            )

            return

        # =========================
        # SEED BY PERFORMANCE
        # =========================
        finalists = sort_results(
            finalists
        )

        # =========================
        # CLEAR OLD FINAL DATA
        # =========================

        cur.execute("""
        DELETE FROM awarded_points
        WHERE event_id = ?
        """, (
            event_id,
        ))

        cur.execute("""
        DELETE FROM final_results
        WHERE event_id = ?
        """, (
            event_id,
        ))

        cur.execute("""
        DELETE FROM finals
        WHERE event_id = ?
        """, (
            event_id,
        ))

        # =========================
        # CREATE FINALISTS
        # =========================
        seed = 1

        for (
            athlete_id,
            performance,
            heat_number,
            athlete_name
        ) in finalists:

            cur.execute("""
            INSERT INTO finals
            (
                event_id,
                athlete_id,
                seed_position
            )
            VALUES (?, ?, ?)
            """, (
                event_id,
                athlete_id,
                seed
            ))

            seed += 1

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Finalists Generated",
            (
                f"{len(finalists)} finalists "
                f"created for:\n\n"
                f"{event_name} "
                f"({gender} {age_group})\n\n"
                f"Method:\n"
                f"{qualification_method}"
            )
        )

    def calculate_positions(
        self,
        heat_id
    ):

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # =========================
        # CLEAR OLD POSITIONS
        # =========================
        cur.execute("""
        UPDATE results
        SET position = NULL
        WHERE heat_id = ?
        """, (
            heat_id,
        ))

        # =========================
        # FIND EVENT TYPE
        # =========================
        cur.execute("""
        SELECT e.event_type

        FROM heats h

        JOIN events e
            ON h.event_id = e.id

        WHERE h.id = ?
        """, (
            heat_id,
        ))

        row = cur.fetchone()

        if not row:

            conn.commit()
            conn.close()

            return

        event_type = row[0]

        # =========================
        # LOAD VALID RESULTS
        # =========================
        cur.execute("""
        SELECT
            athlete_id,
            performance
        FROM results
        WHERE heat_id = ?
        """, (
            heat_id,
        ))

        rows = cur.fetchall()

        athletes = []

        for (
            athlete_id,
            performance
        ) in rows:

            try:

                value = float(
                    performance
                )

                athletes.append(
                    (
                        athlete_id,
                        value
                    )
                )

            except (
                ValueError,
                TypeError
            ):

                # Blank / DNS / DNF / DQ
                continue

        # =========================
        # SORT RESULTS
        # =========================
        if event_type == "Track":

            athletes.sort(
                key=lambda x: x[1]
            )

        else:

            athletes.sort(
                key=lambda x: x[1],
                reverse=True
            )

        # =========================
        # ASSIGN POSITIONS
        # WITH TIES
        # =========================

        previous_value = None
        previous_position = None

        for index, (
            athlete_id,
            value
        ) in enumerate(
            athletes,
            start=1
        ):

            # Same performance =
            # same position
            if (
                previous_value is not None
                and value == previous_value
            ):

                position = (
                    previous_position
                )

            else:

                # Competition ranking:
                # 1, 2, 2, 4
                position = index

            cur.execute("""
            UPDATE results
            SET position = ?
            WHERE heat_id = ?
            AND athlete_id = ?
            """, (
                position,
                heat_id,
                athlete_id
            ))

            previous_value = value
            previous_position = position

        conn.commit()
        conn.close()

# =========================
# FINALS SCREEN
# =========================
class FinalsScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        # =========================
        # TITLE
        # =========================
        title = QLabel("Final Results")
        title.setAlignment(Qt.AlignCenter)

        layout.addWidget(title)

        # =========================
        # EVENT SELECTION
        # =========================
        layout.addWidget(
            QLabel("Select Final Event")
        )

        self.event_combo = QComboBox()

        layout.addWidget(
            self.event_combo
        )

        load_button = QPushButton(
            "Load Final"
        )

        load_button.clicked.connect(
            self.load_finals
        )

        layout.addWidget(
            load_button
        )

        # =========================
        # INSTRUCTIONS
        # =========================
        instruction = QLabel(
            "Enter a numeric performance, "
            "or use DNS / DNF / DQ."
        )

        layout.addWidget(
            instruction
        )

        # =========================
        # FINAL TABLE
        # =========================
        self.table = QTableWidget()

        self.table.setColumnCount(4)

        self.table.setHorizontalHeaderLabels([
            "Lane",
            "Athlete",
            "Performance",
            "Position"
        ])

        self.table.setMinimumHeight(300)

        layout.addWidget(
            self.table,
            1
        )

        # =========================
        # SAVE
        # =========================
        save_button = QPushButton(
            "Save Final Results"
        )

        save_button.clicked.connect(
            self.save_final_results
        )

        layout.addWidget(
            save_button
        )

        self.setLayout(layout)

        self.load_events()

        self.event_combo.currentIndexChanged.connect(
            self.load_finals
        )

        self.load_finals()

    def load_events(self):

        self.event_combo.clear()

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            id,
            event_name,
            gender,
            age_group
        FROM events
        WHERE meeting_id = ?
        AND event_type = 'Track'
        ORDER BY
            event_name,
            gender,
            age_group
        """, (meeting_id,))

        rows = cur.fetchall()

        conn.close()

        for event_id, name, gender, age_group in rows:

            self.event_combo.addItem(
                f"{name} ({gender} {age_group})",
                event_id
            )

    def load_finals(self):

        self.table.setRowCount(0)

        event_id = (
            self.event_combo.currentData()
        )

        if not event_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # =========================
        # LOAD FINALISTS
        # =========================
        cur.execute("""
        SELECT
            f.athlete_id,
            a.full_name,
            f.seed_position
        FROM finals f

        JOIN athletes a
            ON f.athlete_id = a.id

        WHERE f.event_id = ?

        ORDER BY f.seed_position
        """, (event_id,))

        finalists = cur.fetchall()

        if not finalists:

            conn.close()

            return

        # =========================
        # LOAD EXISTING RESULTS
        # =========================
        cur.execute("""
        SELECT
            athlete_id,
            lane_number,
            performance,
            position
        FROM final_results
        WHERE event_id = ?
        """, (event_id,))

        saved_results = {
            row[0]: (
                row[1],
                row[2],
                row[3]
            )
            for row in cur.fetchall()
        }

        conn.close()

        # Preferred championship lanes
        lane_order = [
            4, 5, 3, 6,
            2, 7, 1, 8
        ]

        self.table.setRowCount(
            len(finalists)
        )

        for row, finalist in enumerate(
            finalists
        ):

            athlete_id = finalist[0]
            athlete_name = finalist[1]
            seed = finalist[2]

            # Default lane from seed
            if 1 <= seed <= len(
                lane_order
            ):

                default_lane = (
                    lane_order[
                        seed - 1
                    ]
                )

            else:

                default_lane = ""

            saved_lane, performance, position = (
                saved_results.get(
                    athlete_id,
                    (
                        None,
                        "",
                        ""
                    )
                )
            )

            # Preserve saved lane if one exists
            if (
                saved_lane is not None
                and saved_lane != 0
            ):

                lane = saved_lane

            else:

                lane = default_lane

            if performance is None:
                performance = ""

            if position is None:
                position = ""

            # =========================
            # LANE - READ ONLY
            # =========================
            lane_item = QTableWidgetItem(
                str(lane)
            )

            lane_item.setFlags(
                lane_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.table.setItem(
                row,
                0,
                lane_item
            )

            # =========================
            # ATHLETE - READ ONLY
            # =========================
            athlete_item = QTableWidgetItem(
                athlete_name
            )

            athlete_item.setData(
                Qt.UserRole,
                athlete_id
            )

            athlete_item.setFlags(
                athlete_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.table.setItem(
                row,
                1,
                athlete_item
            )

            # =========================
            # PERFORMANCE - EDITABLE
            # =========================
            performance_item = (
                QTableWidgetItem(
                    str(performance)
                )
            )

            self.table.setItem(
                row,
                2,
                performance_item
            )

            # =========================
            # POSITION - READ ONLY
            # =========================
            position_item = (
                QTableWidgetItem(
                    str(position)
                )
            )

            position_item.setFlags(
                position_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.table.setItem(
                row,
                3,
                position_item
            )

    def save_final_results(self):

        event_id = (
            self.event_combo.currentData()
        )

        if not event_id:

            QMessageBox.warning(
                self,
                "No Event Selected",
                "Please select a final event first."
            )

            return

        valid_codes = {
            "DNS",
            "DNF",
            "DQ"
        }

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        for row in range(
            self.table.rowCount()
        ):

            athlete_item = (
                self.table.item(
                    row,
                    1
                )
            )

            if not athlete_item:
                continue

            athlete_id = (
                athlete_item.data(
                    Qt.UserRole
                )
            )

            lane_item = (
                self.table.item(
                    row,
                    0
                )
            )

            performance_item = (
                self.table.item(
                    row,
                    2
                )
            )

            # =========================
            # LANE
            # =========================
            lane = 0

            if lane_item:

                try:

                    lane = int(
                        lane_item.text()
                    )

                except ValueError:

                    lane = 0

            # =========================
            # PERFORMANCE
            # =========================
            performance = ""

            if performance_item:

                performance = (
                    performance_item
                    .text()
                    .strip()
                )

            # =========================
            # VALIDATE
            # =========================
            if performance:

                upper_value = (
                    performance.upper()
                )

                if upper_value in valid_codes:

                    performance = (
                        upper_value
                    )

                else:

                    try:

                        value = float(
                            performance
                        )

                        if value < 0:

                            raise ValueError

                    except ValueError:

                        athlete_name = (
                            athlete_item.text()
                        )

                        conn.close()

                        QMessageBox.warning(
                            self,
                            "Invalid Performance",
                            (
                                f"Invalid result for:\n\n"
                                f"{athlete_name}\n\n"
                                "Enter a positive number "
                                "or one of these codes:\n\n"
                                "DNS = Did Not Start\n"
                                "DNF = Did Not Finish\n"
                                "DQ = Disqualified"
                            )
                        )

                        return

            # =========================
            # SAVE / UPDATE
            # =========================
            cur.execute("""
            INSERT INTO final_results
            (
                event_id,
                athlete_id,
                lane_number,
                performance
            )
            VALUES (?, ?, ?, ?)

            ON CONFLICT(
                event_id,
                athlete_id
            )

            DO UPDATE SET
                lane_number =
                    excluded.lane_number,
                performance =
                    excluded.performance
            """, (
                event_id,
                athlete_id,
                lane,
                performance
            ))

        conn.commit()
        conn.close()

        # Recalculate positions
        self.calculate_final_positions(
            event_id
        )

        # Rebuild team points
        self.award_final_points(
            event_id
        )

        # Refresh table
        self.load_finals()

        QMessageBox.information(
            self,
            "Final Results Saved",
            "Final results saved successfully."
        )

    def calculate_final_positions(
        self,
        event_id
    ):

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # =========================
        # CLEAR OLD POSITIONS
        # =========================
        cur.execute("""
        UPDATE final_results
        SET position = NULL
        WHERE event_id = ?
        """, (
            event_id,
        ))

        # =========================
        # EVENT TYPE
        # =========================
        cur.execute("""
        SELECT event_type
        FROM events
        WHERE id = ?
        """, (
            event_id,
        ))

        row = cur.fetchone()

        if not row:

            conn.commit()
            conn.close()

            return

        event_type = row[0]

        # =========================
        # LOAD VALID RESULTS
        # =========================
        cur.execute("""
        SELECT
            athlete_id,
            performance
        FROM final_results
        WHERE event_id = ?
        """, (
            event_id,
        ))

        rows = cur.fetchall()

        athletes = []

        for (
            athlete_id,
            performance
        ) in rows:

            try:

                value = float(
                    performance
                )

                athletes.append(
                    (
                        athlete_id,
                        value
                    )
                )

            except (
                ValueError,
                TypeError
            ):

                # Blank / DNS / DNF / DQ
                continue

        # =========================
        # SORT RESULTS
        # =========================
        if event_type == "Track":

            athletes.sort(
                key=lambda x: x[1]
            )

        else:

            athletes.sort(
                key=lambda x: x[1],
                reverse=True
            )

        # =========================
        # ASSIGN POSITIONS
        # WITH TIES
        # =========================

        previous_value = None
        previous_position = None

        for index, (
            athlete_id,
            value
        ) in enumerate(
            athletes,
            start=1
        ):

            if (
                previous_value is not None
                and value == previous_value
            ):

                position = (
                    previous_position
                )

            else:

                position = index

            cur.execute("""
            UPDATE final_results
            SET position = ?
            WHERE event_id = ?
            AND athlete_id = ?
            """, (
                position,
                event_id,
                athlete_id
            ))

            previous_value = value
            previous_position = position

        conn.commit()
        conn.close()

    def award_final_points(
        self,
        event_id
    ):

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Find the meeting this event belongs to
        cur.execute("""
        SELECT meeting_id
        FROM events
        WHERE id = ?
        """, (event_id,))

        row = cur.fetchone()

        if not row:
            conn.close()
            return

        meeting_id = row[0]

        # Remove old points for this event.
        # This prevents duplicate or outdated points
        # if final results are edited later.
        cur.execute("""
        DELETE FROM awarded_points
        WHERE event_id = ?
        """, (event_id,))

        # Load final results together with
        # the athlete's team
        cur.execute("""
        SELECT
            fr.athlete_id,
            fr.position,
            a.team_id
        FROM final_results fr
        JOIN athletes a
            ON fr.athlete_id = a.id
        WHERE fr.event_id = ?
        AND fr.position IS NOT NULL
        """, (event_id,))

        results = cur.fetchall()

        for athlete_id, position, team_id in results:

            # Find how many points this position earns
            cur.execute("""
            SELECT points
            FROM points_config
            WHERE meeting_id = ?
            AND position = ?
            """, (
                meeting_id,
                position
            ))

            points_row = cur.fetchone()

            # No points configured for this position
            if not points_row:
                continue

            points = points_row[0]

            cur.execute("""
            INSERT INTO awarded_points
            (
                meeting_id,
                event_id,
                athlete_id,
                team_id,
                position,
                points
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """, (
                meeting_id,
                event_id,
                athlete_id,
                team_id,
                position,
                points
            ))

        conn.commit()
        conn.close()

# =========================
# POINTS SETUP SCREEN
# =========================
class PointsSetupScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        title = QLabel("Points Setup")
        title.setAlignment(Qt.AlignCenter)

        info = QLabel(
            "Set the number of points awarded for each finishing position.\n"
            "These values are saved separately for each athletics meeting."
        )
        info.setAlignment(Qt.AlignCenter)

        self.table = QTableWidget()
        self.table.setRowCount(8)
        self.table.setColumnCount(2)

        self.table.setHorizontalHeaderLabels([
            "Position",
            "Points"
        ])

        # Suggested default scoring system
        default_points = [
            10,
            8,
            6,
            5,
            4,
            3,
            2,
            1
        ]

        for row in range(8):

            position = row + 1

            position_item = QTableWidgetItem(
                str(position)
            )

            # Do not allow position number to be edited
            position_item.setFlags(
                position_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.table.setItem(
                row,
                0,
                position_item
            )

            self.table.setItem(
                row,
                1,
                QTableWidgetItem(
                    str(default_points[row])
                )
            )

        save_button = QPushButton(
            "Save Points System"
        )

        save_button.clicked.connect(
            self.save_points
        )

        reload_button = QPushButton(
            "Reload Points"
        )

        reload_button.clicked.connect(
            self.load_points
        )

        layout.addWidget(title)
        layout.addWidget(info)
        layout.addWidget(self.table)
        layout.addWidget(save_button)
        layout.addWidget(reload_button)
        layout.addStretch()

        self.setLayout(layout)

        self.load_points()


    def load_points(self):

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            position,
            points
        FROM points_config
        WHERE meeting_id = ?
        """, (meeting_id,))

        saved_points = {
            position: points
            for position, points
            in cur.fetchall()
        }

        conn.close()

        for row in range(8):

            position = row + 1

            if position in saved_points:

                points = saved_points[position]

                # Show whole numbers nicely
                if float(points).is_integer():
                    points = int(points)

                self.table.setItem(
                    row,
                    1,
                    QTableWidgetItem(
                        str(points)
                    )
                )


    def save_points(self):

        meeting_id = get_active_meeting()

        if not meeting_id:

            QMessageBox.warning(
                self,
                "No Active Meeting",
                "Select an active meeting first."
            )

            return

        points_to_save = []

        for row in range(
            self.table.rowCount()
        ):

            position = row + 1

            points_item = self.table.item(
                row,
                1
            )

            if not points_item:

                QMessageBox.warning(
                    self,
                    "Missing Points",
                    f"Enter points for position {position}."
                )

                return

            text = points_item.text().strip()

            try:

                points = float(text)

            except ValueError:

                QMessageBox.warning(
                    self,
                    "Invalid Points",
                    f"Position {position} does not contain a valid number."
                )

                return

            if points < 0:

                QMessageBox.warning(
                    self,
                    "Invalid Points",
                    "Points cannot be negative."
                )

                return

            points_to_save.append(
                (position, points)
            )

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        for position, points in points_to_save:

            cur.execute("""
            INSERT INTO points_config
            (
                meeting_id,
                position,
                points
            )
            VALUES (?, ?, ?)

            ON CONFLICT(meeting_id, position)
            DO UPDATE SET
                points = excluded.points
            """, (
                meeting_id,
                position,
                points
            ))

        conn.commit()
        conn.close()

        QMessageBox.information(
            self,
            "Points Saved",
            "Points system saved successfully."
        )

# =========================
# TEAM STANDINGS SCREEN
# =========================
class TeamStandingsScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        title = QLabel("Team Standings")
        title.setAlignment(Qt.AlignCenter)

        self.meeting_label = QLabel(
            "Points totals for the active meeting"
        )
        self.meeting_label.setAlignment(Qt.AlignCenter)

        self.table = QTableWidget()

        self.table.setColumnCount(3)

        self.table.setHorizontalHeaderLabels([
            "Position",
            "Team / House",
            "Points"
        ])

        refresh_button = QPushButton(
            "Refresh Standings"
        )

        refresh_button.clicked.connect(
            self.load_standings
        )

        layout.addWidget(title)
        layout.addWidget(self.meeting_label)
        layout.addWidget(self.table)
        layout.addWidget(refresh_button)
        layout.addStretch()

        self.setLayout(layout)


    def load_standings(self):

        self.table.setRowCount(0)

        meeting_id = get_active_meeting()

        if not meeting_id:

            self.meeting_label.setText(
                "No active meeting selected."
            )

            return

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        # Get meeting name
        cur.execute("""
        SELECT name
        FROM meetings
        WHERE id = ?
        """, (meeting_id,))

        meeting_row = cur.fetchone()

        if meeting_row:
            self.meeting_label.setText(
                f"Standings: {meeting_row[0]}"
            )

        # Load teams that actually have athletes
        # participating in this meeting.
        #
        # Total points are calculated from
        # awarded_points for this meeting only.
        cur.execute("""
        SELECT
            t.id,
            t.name,

            COALESCE(
                (
                    SELECT SUM(ap.points)
                    FROM awarded_points ap
                    WHERE ap.meeting_id = ?
                    AND ap.team_id = t.id
                ),
                0
            ) AS total_points

        FROM teams t

        WHERE EXISTS (
            SELECT 1
            FROM athletes a
            WHERE a.meeting_id = ?
            AND a.team_id = t.id
        )

        ORDER BY
            total_points DESC,
            t.name ASC
        """, (
            meeting_id,
            meeting_id
        ))

        standings = cur.fetchall()

        conn.close()

        self.table.setRowCount(
            len(standings)
        )

        previous_points = None
        previous_position = 0

        for row_index, standing in enumerate(
            standings
        ):

            team_id = standing[0]
            team_name = standing[1]
            points = standing[2]

            # Competition-style ranking:
            # 1, 2, 2, 4 if teams are tied.
            if previous_points is None:
                position = 1

            elif points == previous_points:
                position = previous_position

            else:
                position = row_index + 1

            previous_points = points
            previous_position = position

            # Make whole-number totals look cleaner
            if float(points).is_integer():
                points_display = str(int(points))
            else:
                points_display = str(points)

            self.table.setItem(
                row_index,
                0,
                QTableWidgetItem(
                    str(position)
                )
            )

            self.table.setItem(
                row_index,
                1,
                QTableWidgetItem(
                    team_name
                )
            )

            self.table.setItem(
                row_index,
                2,
                QTableWidgetItem(
                    points_display
                )
            )

# =========================
# USER MODE SCREEN
# =========================
class UserModeScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        title = QLabel("User Mode")

        title.setAlignment(Qt.AlignCenter)

        admin_button = QPushButton("Administrator")
        operator_button = QPushButton("Meet Operator")

        admin_button.clicked.connect(
            lambda: self.set_mode("Administrator")
        )

        operator_button.clicked.connect(
            lambda: self.set_mode("Meet Operator")
        )

        self.current_mode = QLabel()

        layout.addWidget(title)
        layout.addWidget(self.current_mode)
        layout.addWidget(admin_button)
        layout.addWidget(operator_button)

        self.setLayout(layout)

        self.refresh()

    def set_mode(self, mode):

        set_user_mode(mode)

        self.refresh()

    def refresh(self):

        self.current_mode.setText(
            f"Current Mode: {get_user_mode()}"
        )

# =========================
# FIELD RESULTS SCREEN
# =========================
class FieldResultsScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        title = QLabel("Field Results")
        title.setAlignment(Qt.AlignCenter)

        layout.addWidget(title)

        layout.addWidget(
            QLabel("Select Field Event")
        )

        self.event_combo = QComboBox()

        layout.addWidget(
            self.event_combo
        )

        load_button = QPushButton(
            "Load Field Event"
        )

        load_button.clicked.connect(
            self.load_event
        )

        layout.addWidget(
            load_button
        )

        instruction = QLabel(
            "Enter a distance for each attempt. "
            "Use X for a foul and - for a pass. "
            "Attempts 4-6 are normally used for "
            "the final three attempts."
        )

        layout.addWidget(
            instruction
        )

        self.table = QTableWidget()
        self.table.setColumnCount(9)

        self.table.setHorizontalHeaderLabels([
            "Athlete",
            "Attempt 1",
            "Attempt 2",
            "Attempt 3",
            "Attempt 4",
            "Attempt 5",
            "Attempt 6",
            "Best",
            "Position"
        ])

        self.table.setMinimumHeight(
            400
        )

        layout.addWidget(
            self.table,
            1
        )

        save_button = QPushButton(
            "Save Field Results"
        )

        save_button.clicked.connect(
            self.save_results
        )

        layout.addWidget(
            save_button
        )

        self.setLayout(layout)

        self.load_events()

        self.event_combo.currentIndexChanged.connect(
            self.load_event
        )

        self.load_event()


    # =========================
    # LOAD FIELD EVENTS
    # =========================
    def load_events(self):

        self.event_combo.clear()

        meeting_id = get_active_meeting()

        if not meeting_id:
            return

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        cur.execute("""
        SELECT
            id,
            event_name,
            gender,
            age_group
        FROM events
        WHERE meeting_id = ?
        AND event_type = 'Field'
        ORDER BY
            event_name,
            gender,
            age_group
        """, (
            meeting_id,
        ))

        rows = cur.fetchall()

        conn.close()

        for (
            event_id,
            event_name,
            gender,
            age_group
        ) in rows:

            self.event_combo.addItem(
                (
                    f"{event_name} "
                    f"({gender} {age_group})"
                ),
                event_id
            )

    # =========================
    # DETERMINE FINAL-ROUND
    # QUALIFIERS
    # =========================
    def determine_final_round_qualifiers(
        self,
        athlete_ids,
        attempts
    ):

        # =========================
        # 8 OR FEWER ATHLETES
        # =========================
        if len(athlete_ids) <= 8:

            return set(
                athlete_ids
            ), True

        # =========================
        # FIRST 3 MUST BE COMPLETE
        # =========================
        for athlete_id in athlete_ids:

            for attempt_number in range(
                1,
                4
            ):

                value = attempts.get(
                    (
                        athlete_id,
                        attempt_number
                    ),
                    ""
                )

                if (
                    value is None
                    or not str(value).strip()
                ):

                    return set(), False

        rankings = []

        # =========================
        # RANK FIRST 3 ATTEMPTS
        # =========================
        for athlete_id in athlete_ids:

            numeric_attempts = []

            for attempt_number in range(
                1,
                4
            ):

                value = attempts.get(
                    (
                        athlete_id,
                        attempt_number
                    ),
                    ""
                )

                try:

                    numeric_value = float(
                        value
                    )

                except (
                    ValueError,
                    TypeError
                ):

                    continue

                numeric_attempts.append(
                    numeric_value
                )

            # No valid mark means the
            # athlete cannot advance.
            if not numeric_attempts:
                continue

            ordered = sorted(
                numeric_attempts,
                reverse=True
            )

            # Best mark wins.
            # Second-best and third-best
            # are used as tie-breakers.
            tie_key = tuple(
                ordered
                + [-1.0] * (
                    3 - len(ordered)
                )
            )

            rankings.append(
                (
                    athlete_id,
                    tie_key
                )
            )

        rankings.sort(
            key=lambda item: item[1],
            reverse=True
        )

        # Fewer than 8 athletes with
        # valid marks: all advance.
        if len(rankings) <= 8:

            return {
                athlete_id
                for athlete_id, tie_key
                in rankings
            }, True

        # =========================
        # TOP 8 CUTOFF
        # =========================
        cutoff_key = rankings[7][1]

        qualifiers = {
            athlete_id
            for athlete_id, tie_key
            in rankings
            if tie_key >= cutoff_key
        }

        return qualifiers, True

    # =========================
    # LOAD EVENT
    # =========================

    def load_event(self):

        self.table.setRowCount(0)

        event_id = (
            self.event_combo.currentData()
        )

        meeting_id = get_active_meeting()

        if (
            not meeting_id
            or not event_id
        ):
            return

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        cur.execute("""
        SELECT
            a.id,
            a.full_name
        FROM event_entries ee

        JOIN athletes a
            ON ee.athlete_id = a.id

        WHERE ee.meeting_id = ?
        AND ee.event_id = ?
        AND ee.status = 'Active'
        AND a.status = 'Active'

        ORDER BY a.full_name
        """, (
            meeting_id,
            event_id
        ))

        athletes = cur.fetchall()

        cur.execute("""
        SELECT
            athlete_id,
            attempt_number,
            performance
        FROM field_attempts
        WHERE event_id = ?
        """, (
            event_id,
        ))

        saved_attempts = {}

        for (
            athlete_id,
            attempt_number,
            performance
        ) in cur.fetchall():

            saved_attempts[
                (
                    athlete_id,
                    attempt_number
                )
            ] = performance

        cur.execute("""
        SELECT
            athlete_id,
            performance,
            position
        FROM final_results
        WHERE event_id = ?
        """, (
            event_id,
        ))

        saved_results = {
            row[0]: (
                row[1],
                row[2]
            )
            for row in cur.fetchall()
        }

        conn.close()

        athlete_ids = [
            athlete_id
            for athlete_id, athlete_name
            in athletes
        ]

        qualifiers, preliminary_complete = (
            self.determine_final_round_qualifiers(
                athlete_ids,
                saved_attempts
            )
        )

        self.table.setRowCount(
            len(athletes)
        )

        for row, (
            athlete_id,
            athlete_name
        ) in enumerate(
            athletes
        ):

            athlete_item = (
                QTableWidgetItem(
                    athlete_name
                )
            )

            athlete_item.setData(
                Qt.UserRole,
                athlete_id
            )

            athlete_item.setFlags(
                athlete_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.table.setItem(
                row,
                0,
                athlete_item
            )

            for attempt_number in range(
                1,
                7
            ):

                value = saved_attempts.get(
                    (
                        athlete_id,
                        attempt_number
                    ),
                    ""
                )

                # =========================
                # LOCK ATTEMPTS 4-6
                # =========================
                final_round_locked = (
                    attempt_number >= 4
                    and len(athletes) > 8
                    and (
                        not preliminary_complete
                        or athlete_id
                        not in qualifiers
                    )
                )

                if final_round_locked:

                    value = ""

                item = QTableWidgetItem(
                    str(
                        value
                        if value is not None
                        else ""
                    )
                )

                if final_round_locked:

                    item.setFlags(
                        item.flags()
                        & ~Qt.ItemIsEditable
                    )

                self.table.setItem(
                    row,
                    attempt_number,
                    item
                )

            best, position = (
                saved_results.get(
                    athlete_id,
                    (
                        "",
                        ""
                    )
                )
            )

            if best is None:
                best = ""

            if position is None:
                position = ""

            best_item = QTableWidgetItem(
                str(best)
            )

            best_item.setFlags(
                best_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.table.setItem(
                row,
                7,
                best_item
            )

            position_item = (
                QTableWidgetItem(
                    str(position)
                )
            )

            position_item.setFlags(
                position_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.table.setItem(
                row,
                8,
                position_item
            )

    # =========================
    # SAVE FIELD RESULTS
    # =========================

    def save_results(self):

        event_id = (
            self.event_combo.currentData()
        )

        meeting_id = get_active_meeting()

        if (
            not meeting_id
            or not event_id
        ):

            QMessageBox.warning(
                self,
                "No Event Selected",
                "Select a field event first."
            )

            return

        valid_codes = {
            "X",
            "-"
        }

        # =========================
        # READ + VALIDATE TABLE
        # =========================
        athlete_ids = []
        table_attempts = {}

        for row in range(
            self.table.rowCount()
        ):

            athlete_item = (
                self.table.item(
                    row,
                    0
                )
            )

            if not athlete_item:
                continue

            athlete_id = (
                athlete_item.data(
                    Qt.UserRole
                )
            )

            athlete_name = (
                athlete_item.text()
            )

            athlete_ids.append(
                athlete_id
            )

            for attempt_number in range(
                1,
                7
            ):

                item = self.table.item(
                    row,
                    attempt_number
                )

                performance = ""

                if item:

                    performance = (
                        item.text()
                        .strip()
                    )

                if not performance:

                    table_attempts[
                        (
                            athlete_id,
                            attempt_number
                        )
                    ] = ""

                    continue

                upper_value = (
                    performance.upper()
                )

                if upper_value in valid_codes:

                    table_attempts[
                        (
                            athlete_id,
                            attempt_number
                        )
                    ] = upper_value

                    continue

                try:

                    numeric_value = float(
                        performance
                    )

                    if numeric_value < 0:
                        raise ValueError

                except ValueError:

                    QMessageBox.warning(
                        self,
                        "Invalid Attempt",
                        (
                            f"Invalid attempt for:\n\n"
                            f"{athlete_name}\n\n"
                            "Enter a positive number, "
                            "X for a foul, or - for a pass."
                        )
                    )

                    return

                table_attempts[
                    (
                        athlete_id,
                        attempt_number
                    )
                ] = performance

        # =========================
        # DETERMINE FINAL-ROUND
        # QUALIFIERS
        # =========================
        qualifiers, preliminary_complete = (
            self.determine_final_round_qualifiers(
                athlete_ids,
                table_attempts
            )
        )

        if len(athlete_ids) <= 8:

            final_round_allowed = set(
                athlete_ids
            )

        elif preliminary_complete:

            final_round_allowed = set(
                qualifiers
            )

        else:

            final_round_allowed = set()

        # =========================
        # CHECK IF EVENT COMPLETE
        # =========================
        if len(athlete_ids) <= 8:

            event_complete = (
                bool(athlete_ids)
                and all(
                    str(
                        table_attempts.get(
                            (
                                athlete_id,
                                attempt_number
                            ),
                            ""
                        )
                    ).strip()
                    for athlete_id
                    in athlete_ids
                    for attempt_number
                    in range(1, 7)
                )
            )

        else:

            event_complete = (
                preliminary_complete
                and all(
                    str(
                        table_attempts.get(
                            (
                                athlete_id,
                                attempt_number
                            ),
                            ""
                        )
                    ).strip()
                    for athlete_id
                    in final_round_allowed
                    for attempt_number
                    in range(4, 7)
                )
            )

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        try:

            for athlete_id in athlete_ids:

                numeric_attempts = []

                for attempt_number in range(
                    1,
                    7
                ):

                    # =========================
                    # BLOCK NON-QUALIFIERS
                    # FROM ATTEMPTS 4-6
                    # =========================
                    if (
                        attempt_number >= 4
                        and athlete_id
                        not in final_round_allowed
                    ):

                        cur.execute("""
                        DELETE FROM field_attempts
                        WHERE event_id = ?
                        AND athlete_id = ?
                        AND attempt_number = ?
                        """, (
                            event_id,
                            athlete_id,
                            attempt_number
                        ))

                        continue

                    performance = (
                        table_attempts.get(
                            (
                                athlete_id,
                                attempt_number
                            ),
                            ""
                        )
                    )

                    if not performance:

                        cur.execute("""
                        DELETE FROM field_attempts
                        WHERE event_id = ?
                        AND athlete_id = ?
                        AND attempt_number = ?
                        """, (
                            event_id,
                            athlete_id,
                            attempt_number
                        ))

                        continue

                    try:

                        numeric_value = float(
                            performance
                        )

                        numeric_attempts.append(
                            numeric_value
                        )

                    except (
                        ValueError,
                        TypeError
                    ):

                        pass

                    cur.execute("""
                    INSERT INTO field_attempts
                    (
                        meeting_id,
                        event_id,
                        athlete_id,
                        attempt_number,
                        performance
                    )
                    VALUES (?, ?, ?, ?, ?)

                    ON CONFLICT(
                        event_id,
                        athlete_id,
                        attempt_number
                    )

                    DO UPDATE SET
                        performance =
                            excluded.performance
                    """, (
                        meeting_id,
                        event_id,
                        athlete_id,
                        attempt_number,
                        performance
                    ))

                # =========================
                # BEST PERFORMANCE
                # =========================
                if numeric_attempts:

                    best = max(
                        numeric_attempts
                    )

                    best_text = (
                        f"{best:g}"
                    )

                else:

                    best_text = ""

                cur.execute("""
                INSERT INTO final_results
                (
                    event_id,
                    athlete_id,
                    lane_number,
                    performance
                )
                VALUES (?, ?, 0, ?)

                ON CONFLICT(
                    event_id,
                    athlete_id
                )

                DO UPDATE SET
                    lane_number = 0,
                    performance =
                        excluded.performance
                """, (
                    event_id,
                    athlete_id,
                    best_text
                ))

            conn.commit()

        except sqlite3.Error as error:

            conn.rollback()
            conn.close()

            QMessageBox.critical(
                self,
                "Save Failed",
                (
                    "Field results could not "
                    "be saved.\n\n"
                    f"Database error:\n{error}"
                )
            )

            return

        conn.close()

        self.calculate_positions(
            event_id
        )

        # =========================
        # POINTS ONLY AFTER EVENT
        # IS COMPLETE
        # =========================
        if event_complete:

            self.award_points(
                event_id
            )

        else:

            points_conn = sqlite3.connect(
                DB_NAME
            )

            points_cur = (
                points_conn.cursor()
            )

            points_cur.execute("""
            DELETE FROM awarded_points
            WHERE event_id = ?
            """, (
                event_id,
            ))

            points_conn.commit()
            points_conn.close()

        self.load_event()

        # =========================
        # STATUS MESSAGE
        # =========================
        if event_complete:

            message = (
                "Field event completed.\n\n"
                "Final positions and points "
                "have been updated."
            )

        elif (
            len(athlete_ids) > 8
            and preliminary_complete
        ):

            message = (
                "The first three rounds "
                "have been saved.\n\n"
                f"{len(final_round_allowed)} "
                "athlete(s) may now complete "
                "Attempts 4-6."
            )

        elif len(athlete_ids) > 8:

            message = (
                "Results saved.\n\n"
                "Complete Attempts 1-3 for "
                "every athlete. Attempts 4-6 "
                "will then unlock for the "
                "qualifying athletes."
            )

        else:

            message = (
                "Results saved.\n\n"
                "There are 8 or fewer athletes, "
                "so every athlete may complete "
                "Attempts 1-6."
            )

        QMessageBox.information(
            self,
            "Field Results Saved",
            message
        )

    # =========================
    # CALCULATE FIELD POSITIONS
    # =========================
    def calculate_positions(
        self,
        event_id
    ):

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        cur.execute("""
        UPDATE final_results
        SET position = NULL
        WHERE event_id = ?
        """, (
            event_id,
        ))

        cur.execute("""
        SELECT
            athlete_id,
            performance
        FROM field_attempts
        WHERE event_id = ?
        ORDER BY
            athlete_id,
            attempt_number
        """, (
            event_id,
        ))

        athlete_attempts = {}

        for (
            athlete_id,
            performance
        ) in cur.fetchall():

            try:

                value = float(
                    performance
                )

            except (
                ValueError,
                TypeError
            ):

                continue

            if athlete_id not in (
                athlete_attempts
            ):

                athlete_attempts[
                    athlete_id
                ] = []

            athlete_attempts[
                athlete_id
            ].append(
                value
            )

        rankings = []

        for (
            athlete_id,
            attempts
        ) in athlete_attempts.items():

            ordered = sorted(
                attempts,
                reverse=True
            )

            tie_key = tuple(
                ordered
                + [-1.0] * (
                    6 - len(ordered)
                )
            )

            rankings.append(
                (
                    athlete_id,
                    tie_key
                )
            )

        rankings.sort(
            key=lambda x: x[1],
            reverse=True
        )

        previous_key = None
        previous_position = None

        for index, (
            athlete_id,
            tie_key
        ) in enumerate(
            rankings,
            start=1
        ):

            if (
                previous_key is not None
                and tie_key == previous_key
            ):

                position = (
                    previous_position
                )

            else:

                position = index

            cur.execute("""
            UPDATE final_results
            SET position = ?
            WHERE event_id = ?
            AND athlete_id = ?
            """, (
                position,
                event_id,
                athlete_id
            ))

            previous_key = tie_key
            previous_position = position

        conn.commit()
        conn.close()


    # =========================
    # AWARD FIELD POINTS
    # =========================
    def award_points(
        self,
        event_id
    ):

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        cur.execute("""
        SELECT meeting_id
        FROM events
        WHERE id = ?
        """, (
            event_id,
        ))

        row = cur.fetchone()

        if not row:

            conn.close()
            return

        meeting_id = row[0]

        cur.execute("""
        DELETE FROM awarded_points
        WHERE event_id = ?
        """, (
            event_id,
        ))

        cur.execute("""
        SELECT
            fr.athlete_id,
            fr.position,
            a.team_id
        FROM final_results fr

        JOIN athletes a
            ON fr.athlete_id = a.id

        WHERE fr.event_id = ?
        AND fr.position IS NOT NULL
        """, (
            event_id,
        ))

        results = cur.fetchall()

        for (
            athlete_id,
            position,
            team_id
        ) in results:

            cur.execute("""
            SELECT points
            FROM points_config
            WHERE meeting_id = ?
            AND position = ?
            """, (
                meeting_id,
                position
            ))

            points_row = (
                cur.fetchone()
            )

            if not points_row:
                continue

            points = points_row[0]

            cur.execute("""
            INSERT INTO awarded_points
            (
                meeting_id,
                event_id,
                athlete_id,
                team_id,
                position,
                points
            )
            VALUES (?, ?, ?, ?, ?, ?)

            ON CONFLICT(
                event_id,
                athlete_id
            )

            DO UPDATE SET
                team_id =
                    excluded.team_id,
                position =
                    excluded.position,
                points =
                    excluded.points
            """, (
                meeting_id,
                event_id,
                athlete_id,
                team_id,
                position,
                points
            ))

        conn.commit()
        conn.close()

# =========================
# HIGH JUMP RESULTS SCREEN
# =========================
class HighJumpResultsScreen(QWidget):

    def __init__(self):
        super().__init__()

        layout = QVBoxLayout()

        title = QLabel("High Jump Results")
        title.setAlignment(Qt.AlignCenter)

        self.event_combo = QComboBox()

        refresh_button = QPushButton(
            "Refresh High Jump Events"
        )

        refresh_button.clicked.connect(
            self.load_events
        )

        layout.addWidget(title)

        layout.addWidget(
            QLabel("Select High Jump Event")
        )

        layout.addWidget(
            self.event_combo
        )

        layout.addWidget(
            refresh_button
        )

        # =========================
        # COMPETITION HEIGHTS
        # =========================
        layout.addWidget(
            QLabel(
                "Competition Heights "
                "(metres, separated by commas)"
            )
        )

        self.heights_input = QLineEdit()

        self.heights_input.setPlaceholderText(
            "Example: 1.10, 1.15, 1.20, 1.25"
        )

        layout.addWidget(
            self.heights_input
        )

        save_heights_button = QPushButton(
            "Save Competition Heights"
        )

        save_heights_button.clicked.connect(
            self.save_heights
        )

        layout.addWidget(
            save_heights_button
        )

        # =========================
        # HEIGHT TABLE
        # =========================
        self.heights_table = QTableWidget()

        self.heights_table.setColumnCount(
            2
        )

        self.heights_table.setHorizontalHeaderLabels([
            "Order",
            "Height (m)"
        ])

        layout.addWidget(
            self.heights_table
        )

        # =========================
        # HIGH JUMP ATTEMPT GRID
        # =========================
        layout.addWidget(
            QLabel(
                "Athlete Attempts"
            )
        )

        instruction = QLabel(
            "Use O = cleared, X = failed attempt, "
            "- = pass. Examples: O, XO, XXO, XXX, X-, XX-"
        )

        instruction.setWordWrap(
            True
        )

        layout.addWidget(
            instruction
        )

        self.results_table = QTableWidget()

        layout.addWidget(
            self.results_table,
            1
        )

        save_attempts_button = QPushButton(
            "Save High Jump Attempts"
        )

        save_attempts_button.clicked.connect(
            self.save_attempts
        )

        finalize_button = QPushButton(
            "Finalize High Jump Event"
        )

        finalize_button.clicked.connect(
            self.finalize_event
        )

        layout.addWidget(
            finalize_button
        )

        layout.addWidget(
            save_attempts_button
        )

        self.setLayout(layout)

        self.load_events()

        self.event_combo.currentIndexChanged.connect(
            self.load_heights
        )

        self.load_heights()


    # =========================
    # LOAD HIGH JUMP EVENTS
    # =========================
    def load_events(self):

        current_event_id = (
            self.event_combo.currentData()
        )

        self.event_combo.blockSignals(
            True
        )

        self.event_combo.clear()

        meeting_id = get_active_meeting()

        if meeting_id:

            conn = sqlite3.connect(
                DB_NAME
            )

            cur = conn.cursor()

            cur.execute("""
            SELECT
                id,
                event_name,
                gender,
                age_group
            FROM events
            WHERE meeting_id = ?
            AND event_type = 'High Jump'
            ORDER BY
                event_name,
                gender,
                age_group
            """, (
                meeting_id,
            ))

            rows = cur.fetchall()

            conn.close()

            for (
                event_id,
                event_name,
                gender,
                age_group
            ) in rows:

                self.event_combo.addItem(
                    (
                        f"{event_name} "
                        f"({gender} {age_group})"
                    ),
                    event_id
                )

        if current_event_id is not None:

            index = (
                self.event_combo.findData(
                    current_event_id
                )
            )

            if index >= 0:

                self.event_combo.setCurrentIndex(
                    index
                )

        self.event_combo.blockSignals(
            False
        )

        self.load_heights()


    # =========================
    # LOAD COMPETITION HEIGHTS
    # =========================
    def load_heights(self):

        self.heights_table.setRowCount(
            0
        )

        self.heights_input.clear()

        event_id = (
            self.event_combo.currentData()
        )

        if not event_id:
            return

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        cur.execute("""
        SELECT
            height_order,
            height
        FROM high_jump_heights
        WHERE event_id = ?
        ORDER BY height_order
        """, (
            event_id,
        ))

        rows = cur.fetchall()

        conn.close()

        self.heights_table.setRowCount(
            len(rows)
        )

        height_texts = []

        for row_index, (
            height_order,
            height
        ) in enumerate(rows):

            order_item = QTableWidgetItem(
                str(height_order)
            )

            order_item.setFlags(
                order_item.flags()
                & ~Qt.ItemIsEditable
            )

            height_item = QTableWidgetItem(
                f"{height:.2f}"
            )

            height_item.setFlags(
                height_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.heights_table.setItem(
                row_index,
                0,
                order_item
            )

            self.heights_table.setItem(
                row_index,
                1,
                height_item
            )

            height_texts.append(
                f"{height:.2f}"
            )

        self.heights_input.setText(
            ", ".join(height_texts)
        )

        self.load_attempt_grid()

    # =========================
    # CALCULATE HIGH JUMP
    # RESULTS
    # =========================
    def calculate_high_jump_results(
        self,
        athletes,
        heights,
        saved_attempts
    ):

        results = {}

        for (
            athlete_id,
            athlete_name
        ) in athletes:

            best_height = None
            attempts_at_best = None

            total_failures = 0
            failures_to_best = 0

            consecutive_failures = 0
            eliminated = False

            for (
                height_id,
                height
            ) in heights:

                attempts = (
                    saved_attempts.get(
                        (
                            athlete_id,
                            height_id
                        ),
                        {}
                    )
                )

                code = ""

                for attempt_number in sorted(
                    attempts.keys()
                ):

                    code += str(
                        attempts[
                            attempt_number
                        ]
                    ).upper()

                if not code:
                    continue

                attempts_used = 0

                for symbol in code:

                    if symbol == "X":

                        attempts_used += 1

                        total_failures += 1

                        consecutive_failures += 1

                        if consecutive_failures >= 3:

                            eliminated = True

                    elif symbol == "O":

                        attempts_used += 1

                        best_height = height

                        attempts_at_best = (
                            attempts_used
                        )

                        failures_to_best = (
                            total_failures
                        )

                        consecutive_failures = 0

                        eliminated = False

                    elif symbol == "-":

                        pass

            if eliminated:

                status = "Eliminated"

            else:

                status = "Active"

            results[
                athlete_id
            ] = {
                "best_height": best_height,
                "attempts_at_best": attempts_at_best,
                "failures_to_best": failures_to_best,
                "total_failures": total_failures,
                "status": status,
                "position": None
            }

        # =========================
        # HIGH JUMP RANKING
        # =========================
        ranked = []

        for (
            athlete_id,
            data
        ) in results.items():

            best_height = data[
                "best_height"
            ]

            if best_height is None:
                continue

            attempts_at_best = (
                data[
                    "attempts_at_best"
                ]
            )

            failures_to_best = (
                data[
                    "failures_to_best"
                ]
            )

            ranked.append(
                (
                    athlete_id,
                    best_height,
                    attempts_at_best,
                    failures_to_best
                )
            )

        ranked.sort(
            key=lambda item: (
                -item[1],
                item[2],
                item[3]
            )
        )

        previous_key = None
        previous_position = None

        for index, (
            athlete_id,
            best_height,
            attempts_at_best,
            failures_to_best
        ) in enumerate(
            ranked,
            start=1
        ):

            tie_key = (
                best_height,
                attempts_at_best,
                failures_to_best
            )

            if (
                previous_key is not None
                and tie_key == previous_key
            ):

                position = (
                    previous_position
                )

            else:

                position = index

            results[
                athlete_id
            ][
                "position"
            ] = position

            previous_key = tie_key
            previous_position = position

        return results

    # =========================
    # LOAD HIGH JUMP ATTEMPTS
    # =========================
    def load_attempt_grid(self):

        self.results_table.clear()

        self.results_table.setRowCount(
            0
        )

        event_id = (
            self.event_combo.currentData()
        )

        meeting_id = get_active_meeting()

        self.results_table.setColumnCount(
            1
        )

        self.results_table.setHorizontalHeaderLabels([
            "Athlete"
        ])

        if (
            not meeting_id
            or not event_id
        ):
            return

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        # =========================
        # LOAD HEIGHTS
        # =========================
        cur.execute("""
        SELECT
            id,
            height
        FROM high_jump_heights
        WHERE event_id = ?
        ORDER BY height_order
        """, (
            event_id,
        ))

        heights = cur.fetchall()

        # =========================
        # LOAD ACTIVE ATHLETES
        # =========================
        cur.execute("""
        SELECT
            a.id,
            a.full_name
        FROM event_entries ee

        JOIN athletes a
            ON ee.athlete_id = a.id

        WHERE ee.meeting_id = ?
        AND ee.event_id = ?
        AND ee.status = 'Active'
        AND a.status = 'Active'

        ORDER BY a.full_name
        """, (
            meeting_id,
            event_id
        ))

        athletes = cur.fetchall()

        # =========================
        # LOAD SAVED ATTEMPTS
        # =========================
        cur.execute("""
        SELECT
            athlete_id,
            height_id,
            attempt_number,
            result
        FROM high_jump_attempts
        WHERE event_id = ?
        ORDER BY
            athlete_id,
            height_id,
            attempt_number
        """, (
            event_id,
        ))

        saved_attempts = {}

        for (
            athlete_id,
            height_id,
            attempt_number,
            result
        ) in cur.fetchall():

            key = (
                athlete_id,
                height_id
            )

            if key not in saved_attempts:

                saved_attempts[key] = {}

            saved_attempts[
                key
            ][
                attempt_number
            ] = result

        conn.close()

        # =========================
        # CALCULATE LIVE RESULTS
        # =========================
        calculated = (
            self.calculate_high_jump_results(
                athletes,
                heights,
                saved_attempts
            )
        )

        # Athlete
        # + height columns
        # + Best
        # + Failures
        # + Status
        # + Position
        total_columns = (
            1
            + len(heights)
            + 4
        )

        self.results_table.setColumnCount(
            total_columns
        )

        self.results_table.setHorizontalHeaderItem(
            0,
            QTableWidgetItem("Athlete")
        )

        for column, (
            height_id,
            height
        ) in enumerate(
            heights,
            start=1
        ):

            header_item = QTableWidgetItem(
                f"{height:.2f}"
            )

            header_item.setData(
                Qt.UserRole,
                height_id
            )

            self.results_table.setHorizontalHeaderItem(
                column,
                header_item
            )

        best_column = (
            len(heights) + 1
        )

        failures_column = (
            len(heights) + 2
        )

        status_column = (
            len(heights) + 3
        )

        position_column = (
            len(heights) + 4
        )

        self.results_table.setHorizontalHeaderItem(
            best_column,
            QTableWidgetItem("Best")
        )

        self.results_table.setHorizontalHeaderItem(
            failures_column,
            QTableWidgetItem("Failures")
        )

        self.results_table.setHorizontalHeaderItem(
            status_column,
            QTableWidgetItem("Status")
        )

        self.results_table.setHorizontalHeaderItem(
            position_column,
            QTableWidgetItem("Position")
        )

        self.results_table.setRowCount(
            len(athletes)
        )

        for row, (
            athlete_id,
            athlete_name
        ) in enumerate(
            athletes
        ):

            athlete_item = QTableWidgetItem(
                athlete_name
            )

            athlete_item.setData(
                Qt.UserRole,
                athlete_id
            )

            athlete_item.setFlags(
                athlete_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.results_table.setItem(
                row,
                0,
                athlete_item
            )

            for column, (
                height_id,
                height
            ) in enumerate(
                heights,
                start=1
            ):

                attempts = (
                    saved_attempts.get(
                        (
                            athlete_id,
                            height_id
                        ),
                        {}
                    )
                )

                code = ""

                for attempt_number in sorted(
                    attempts.keys()
                ):

                    code += str(
                        attempts[
                            attempt_number
                        ]
                    )

                attempt_item = QTableWidgetItem(
                    code
                )

                self.results_table.setItem(
                    row,
                    column,
                    attempt_item
                )

            result = calculated.get(
                athlete_id,
                {}
            )

            best_height = result.get(
                "best_height"
            )

            if best_height is None:

                best_text = ""

            else:

                best_text = (
                    f"{best_height:.2f}"
                )

            best_item = QTableWidgetItem(
                best_text
            )

            best_item.setFlags(
                best_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.results_table.setItem(
                row,
                best_column,
                best_item
            )

            failures_item = QTableWidgetItem(
                str(
                    result.get(
                        "total_failures",
                        0
                    )
                )
            )

            failures_item.setFlags(
                failures_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.results_table.setItem(
                row,
                failures_column,
                failures_item
            )

            status_item = QTableWidgetItem(
                result.get(
                    "status",
                    "Active"
                )
            )

            status_item.setFlags(
                status_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.results_table.setItem(
                row,
                status_column,
                status_item
            )

            position = result.get(
                "position"
            )

            position_item = QTableWidgetItem(
                (
                    str(position)
                    if position is not None
                    else ""
                )
            )

            position_item.setFlags(
                position_item.flags()
                & ~Qt.ItemIsEditable
            )

            self.results_table.setItem(
                row,
                position_column,
                position_item
            )

    # =========================
    # SAVE HIGH JUMP ATTEMPTS
    # =========================
    def save_attempts(self):

        meeting_id = get_active_meeting()

        event_id = (
            self.event_combo.currentData()
        )

        if (
            not meeting_id
            or not event_id
        ):

            QMessageBox.warning(
                self,
                "No High Jump Event",
                "Select a High Jump event first."
            )

            return

        if self.results_table.columnCount() <= 1:

            QMessageBox.warning(
                self,
                "No Heights",
                (
                    "Save the competition heights "
                    "before entering attempts."
                )
            )

            return

        # =========================
        # ALLOWED CELL CODES
        # =========================
        valid_codes = {
            "",
            "O",
            "XO",
            "XXO",
            "X",
            "XX",
            "XXX",
            "-",
            "X-",
            "XX-"
        }

        values_to_save = []

        # =========================
        # VALIDATE TABLE
        # =========================
        for row in range(
            self.results_table.rowCount()
        ):

            athlete_item = (
                self.results_table.item(
                    row,
                    0
                )
            )

            if not athlete_item:
                continue

            athlete_id = (
                athlete_item.data(
                    Qt.UserRole
                )
            )

            athlete_name = (
                athlete_item.text()
            )

            # Only the columns that have
            # a height ID are attempt columns.
            for column in range(
                1,
                self.results_table.columnCount()
            ):

                header_item = (
                    self.results_table.horizontalHeaderItem(
                        column
                    )
                )

                if not header_item:
                    continue

                height_id = (
                    header_item.data(
                        Qt.UserRole
                    )
                )

                if height_id is None:
                    continue

                height_text = (
                    header_item.text()
                )

                item = (
                    self.results_table.item(
                        row,
                        column
                    )
                )

                code = ""

                if item:

                    code = (
                        item.text()
                        .strip()
                        .upper()
                        .replace(" ", "")
                    )

                if code not in valid_codes:

                    QMessageBox.warning(
                        self,
                        "Invalid High Jump Result",
                        (
                            f"Invalid result for:\n\n"
                            f"{athlete_name}\n"
                            f"Height: {height_text} m\n\n"
                            "Allowed examples:\n"
                            "O\n"
                            "XO\n"
                            "XXO\n"
                            "X\n"
                            "XX\n"
                            "XXX\n"
                            "-\n"
                            "X-\n"
                            "XX-"
                        )
                    )

                    return

                values_to_save.append(
                    (
                        athlete_id,
                        height_id,
                        code
                    )
                )

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        try:

            for (
                athlete_id,
                height_id,
                code
            ) in values_to_save:

                # Remove the old sequence for
                # this athlete at this height.
                cur.execute("""
                DELETE FROM high_jump_attempts
                WHERE event_id = ?
                AND athlete_id = ?
                AND height_id = ?
                """, (
                    event_id,
                    athlete_id,
                    height_id
                ))

                if not code:
                    continue

                # Store each symbol as its own
                # numbered attempt.
                for (
                    attempt_number,
                    result
                ) in enumerate(
                    code,
                    start=1
                ):

                    cur.execute("""
                    INSERT INTO high_jump_attempts
                    (
                        meeting_id,
                        event_id,
                        athlete_id,
                        height_id,
                        attempt_number,
                        result
                    )
                    VALUES (?, ?, ?, ?, ?, ?)
                    """, (
                        meeting_id,
                        event_id,
                        athlete_id,
                        height_id,
                        attempt_number,
                        result
                    ))

            conn.commit()

        except sqlite3.Error as error:

            conn.rollback()
            conn.close()

            QMessageBox.critical(
                self,
                "Save Failed",
                (
                    "High Jump attempts could "
                    "not be saved.\n\n"
                    f"Database error:\n{error}"
                )
            )

            return

        conn.close()

        self.load_attempt_grid()

        QMessageBox.information(
            self,
            "High Jump Attempts Saved",
            "High Jump attempts saved successfully."
        )

    # =========================
    # FINALIZE HIGH JUMP EVENT
    # =========================
    def finalize_event(self):

        meeting_id = get_active_meeting()

        event_id = (
            self.event_combo.currentData()
        )

        if (
            not meeting_id
            or not event_id
        ):

            QMessageBox.warning(
                self,
                "No High Jump Event",
                "Select a High Jump event first."
            )

            return

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        # =========================
        # LOAD HEIGHTS
        # =========================
        cur.execute("""
        SELECT
            id,
            height
        FROM high_jump_heights
        WHERE event_id = ?
        ORDER BY height_order
        """, (
            event_id,
        ))

        heights = cur.fetchall()

        if not heights:

            conn.close()

            QMessageBox.warning(
                self,
                "No Heights",
                (
                    "Competition heights have not "
                    "been set for this event."
                )
            )

            return

        # =========================
        # LOAD ACTIVE ENTRIES
        # =========================
        cur.execute("""
        SELECT
            a.id,
            a.full_name
        FROM event_entries ee

        JOIN athletes a
            ON ee.athlete_id = a.id

        WHERE ee.meeting_id = ?
        AND ee.event_id = ?
        AND ee.status = 'Active'
        AND a.status = 'Active'

        ORDER BY a.full_name
        """, (
            meeting_id,
            event_id
        ))

        athletes = cur.fetchall()

        if not athletes:

            conn.close()

            QMessageBox.warning(
                self,
                "No Athletes",
                (
                    "There are no active athletes "
                    "in this High Jump event."
                )
            )

            return

        # =========================
        # LOAD ATTEMPTS
        # =========================
        cur.execute("""
        SELECT
            athlete_id,
            height_id,
            attempt_number,
            result
        FROM high_jump_attempts
        WHERE event_id = ?
        ORDER BY
            athlete_id,
            height_id,
            attempt_number
        """, (
            event_id,
        ))

        saved_attempts = {}

        for (
            athlete_id,
            height_id,
            attempt_number,
            result
        ) in cur.fetchall():

            key = (
                athlete_id,
                height_id
            )

            if key not in saved_attempts:

                saved_attempts[
                    key
                ] = {}

            saved_attempts[
                key
            ][
                attempt_number
            ] = result

        conn.close()

        if not saved_attempts:

            QMessageBox.warning(
                self,
                "No Results",
                (
                    "No High Jump attempts have "
                    "been recorded yet."
                )
            )

            return

        calculated = (
            self.calculate_high_jump_results(
                athletes,
                heights,
                saved_attempts
            )
        )

        # =========================
        # CHECK ACTIVE ATHLETES
        # =========================
        active_count = sum(
            1
            for athlete_id, data
            in calculated.items()
            if data.get("status") == "Active"
        )

        if active_count > 1:

            answer = QMessageBox.question(
                self,
                "Competition Still Active",
                (
                    f"{active_count} athletes are "
                    "still marked as Active.\n\n"
                    "This normally means the "
                    "competition may not be finished.\n\n"
                    "Finalize the current results anyway?"
                ),
                QMessageBox.Yes
                | QMessageBox.No
            )

            if answer != QMessageBox.Yes:
                return

        else:

            answer = QMessageBox.question(
                self,
                "Finalize High Jump",
                (
                    "Save the current High Jump "
                    "positions as the official results "
                    "and award points?\n\n"
                    "You can finalize again later if "
                    "a correction is required."
                ),
                QMessageBox.Yes
                | QMessageBox.No
            )

            if answer != QMessageBox.Yes:
                return

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        try:

            # =========================
            # REBUILD OFFICIAL RESULTS
            # =========================
            cur.execute("""
            DELETE FROM final_results
            WHERE event_id = ?
            """, (
                event_id,
            ))

            cur.execute("""
            DELETE FROM awarded_points
            WHERE event_id = ?
            """, (
                event_id,
            ))

            for (
                athlete_id,
                athlete_name
            ) in athletes:

                data = calculated.get(
                    athlete_id,
                    {}
                )

                best_height = data.get(
                    "best_height"
                )

                position = data.get(
                    "position"
                )

                if best_height is None:

                    performance = "NM"

                else:

                    performance = (
                        f"{best_height:.2f}"
                    )

                cur.execute("""
                INSERT INTO final_results
                (
                    event_id,
                    athlete_id,
                    lane_number,
                    performance,
                    position
                )
                VALUES (?, ?, 0, ?, ?)
                """, (
                    event_id,
                    athlete_id,
                    performance,
                    position
                ))

            # =========================
            # AWARD POINTS
            # =========================
            cur.execute("""
            SELECT
                fr.athlete_id,
                fr.position,
                a.team_id
            FROM final_results fr

            JOIN athletes a
                ON fr.athlete_id = a.id

            WHERE fr.event_id = ?
            AND fr.position IS NOT NULL
            """, (
                event_id,
            ))

            official_results = (
                cur.fetchall()
            )

            for (
                athlete_id,
                position,
                team_id
            ) in official_results:

                cur.execute("""
                SELECT points
                FROM points_config
                WHERE meeting_id = ?
                AND position = ?
                """, (
                    meeting_id,
                    position
                ))

                points_row = (
                    cur.fetchone()
                )

                if not points_row:
                    continue

                points = points_row[0]

                cur.execute("""
                INSERT INTO awarded_points
                (
                    meeting_id,
                    event_id,
                    athlete_id,
                    team_id,
                    position,
                    points
                )
                VALUES (?, ?, ?, ?, ?, ?)

                ON CONFLICT(
                    event_id,
                    athlete_id
                )

                DO UPDATE SET
                    team_id =
                        excluded.team_id,
                    position =
                        excluded.position,
                    points =
                        excluded.points
                """, (
                    meeting_id,
                    event_id,
                    athlete_id,
                    team_id,
                    position,
                    points
                ))

            conn.commit()

        except sqlite3.Error as error:

            conn.rollback()
            conn.close()

            QMessageBox.critical(
                self,
                "Finalize Failed",
                (
                    "The High Jump event could "
                    "not be finalized.\n\n"
                    f"Database error:\n{error}"
                )
            )

            return

        conn.close()

        self.load_attempt_grid()

        QMessageBox.information(
            self,
            "High Jump Finalized",
            (
                "The official High Jump results "
                "have been saved and the points "
                "have been updated."
            )
        )

    # =========================
    # SAVE COMPETITION HEIGHTS
    # =========================
    def save_heights(self):

        meeting_id = get_active_meeting()

        event_id = (
            self.event_combo.currentData()
        )

        if (
            not meeting_id
            or not event_id
        ):

            QMessageBox.warning(
                self,
                "No High Jump Event",
                "Select a High Jump event first."
            )

            return

        raw_text = (
            self.heights_input.text()
            .strip()
        )

        if not raw_text:

            QMessageBox.warning(
                self,
                "No Heights",
                (
                    "Enter the competition heights "
                    "separated by commas."
                )
            )

            return

        parts = [
            part.strip()
            for part in raw_text.split(",")
            if part.strip()
        ]

        heights = []

        try:

            for part in parts:

                height = float(
                    part
                )

                if height <= 0:
                    raise ValueError

                heights.append(
                    height
                )

        except ValueError:

            QMessageBox.warning(
                self,
                "Invalid Height",
                (
                    "Every height must be a positive "
                    "number.\n\n"
                    "Example:\n"
                    "1.10, 1.15, 1.20, 1.25"
                )
            )

            return

        # =========================
        # DUPLICATE HEIGHT CHECK
        # =========================
        if len(heights) != len(
            set(heights)
        ):

            QMessageBox.warning(
                self,
                "Duplicate Height",
                (
                    "Each competition height "
                    "must be unique."
                )
            )

            return

        # =========================
        # ASCENDING ORDER CHECK
        # =========================
        if heights != sorted(
            heights
        ):

            QMessageBox.warning(
                self,
                "Height Order",
                (
                    "Enter the heights from "
                    "lowest to highest."
                )
            )

            return

        conn = sqlite3.connect(
            DB_NAME
        )

        cur = conn.cursor()

        # =========================
        # LOCK HEIGHTS AFTER
        # ATTEMPTS HAVE STARTED
        # =========================
        cur.execute("""
        SELECT COUNT(*)
        FROM high_jump_attempts
        WHERE event_id = ?
        """, (
            event_id,
        ))

        attempt_count = (
            cur.fetchone()[0]
        )

        if attempt_count > 0:

            conn.close()

            QMessageBox.warning(
                self,
                "Heights Locked",
                (
                    "High Jump attempts have already "
                    "been recorded for this event.\n\n"
                    "The competition heights can no "
                    "longer be changed."
                )
            )

            return

        try:

            cur.execute("""
            DELETE FROM high_jump_heights
            WHERE event_id = ?
            """, (
                event_id,
            ))

            for (
                height_order,
                height
            ) in enumerate(
                heights,
                start=1
            ):

                cur.execute("""
                INSERT INTO high_jump_heights
                (
                    meeting_id,
                    event_id,
                    height_order,
                    height
                )
                VALUES (?, ?, ?, ?)
                """, (
                    meeting_id,
                    event_id,
                    height_order,
                    height
                ))

            conn.commit()

        except sqlite3.Error as error:

            conn.rollback()
            conn.close()

            QMessageBox.critical(
                self,
                "Save Failed",
                (
                    "The High Jump heights could "
                    "not be saved.\n\n"
                    f"Database error:\n{error}"
                )
            )

            return

        conn.close()

        self.load_heights()

        QMessageBox.information(
            self,
            "Heights Saved",
            (
                f"{len(heights)} competition "
                "height(s) saved successfully."
            )
        )

# =========================
# DASHBOARD SCREEN
# =========================
class DashboardScreen(QWidget):

    def __init__(self, stack):
        super().__init__()

        self.stack = stack

        layout = QVBoxLayout()

        title = QLabel("MY ATHLETICS")
        title.setAlignment(Qt.AlignCenter)

        subtitle = QLabel("Version 1.0 Foundation")
        subtitle.setAlignment(Qt.AlignCenter)

        create_button = QPushButton("Create Meeting")
        create_button.clicked.connect(
            lambda: self.stack.setCurrentIndex(1)
        )

        load_button = QPushButton("Load Meeting")
        load_button.clicked.connect(
            lambda: self.stack.setCurrentIndex(2)
        )

        teams_button = QPushButton("Teams")
        teams_button.clicked.connect(
            self.open_teams
        )

        mode_button = QPushButton("User Mode")
        mode_button.clicked.connect(
            lambda: self.stack.setCurrentIndex(4)
        )

        athletes_button = QPushButton("Athletes")
        athletes_button.clicked.connect(
            self.open_athletes
        )

        events_button = QPushButton("Events")
        events_button.clicked.connect(
            self.open_events
        )

        event_entry_button = QPushButton("Event Entries")
        event_entry_button.clicked.connect(
            self.open_event_entries
        )

        results_button = QPushButton("Results")
        results_button.clicked.connect(
        self.open_results
        )

        field_results_button = QPushButton(
            "Field Results"
        )

        field_results_button.clicked.connect(
            self.open_field_results
        )

        high_jump_button = QPushButton(
            "High Jump Results"
        )

        high_jump_button.clicked.connect(
            self.open_high_jump_results
        )

        finals_button = QPushButton("Finals")
        finals_button.clicked.connect(
            lambda: self.stack.setCurrentIndex(9)
        )

        points_button = QPushButton("Points Setup")
        points_button.clicked.connect(
            self.open_points
        )

        standings_button = QPushButton(
            "Team Standings"
        )

        standings_button.clicked.connect(
            self.open_standings
        )

        backup_button = QPushButton(
            "Backup Database"
        )

        backup_button.clicked.connect(
            self.backup_database
        )

        restore_button = QPushButton(
            "Restore Backup"
        )

        restore_button.clicked.connect(
            self.restore_database
        )

        template_button = QPushButton(
            "Download Entry Template"
        )

        template_button.clicked.connect(
            self.download_entry_template
        )

        import_button = QPushButton(
            "Import Entry Files"
        )

        import_button.clicked.connect(
            self.import_entry_files
        )

        generate_all_button = QPushButton(
            "Generate All Heats / Finals"
        )

        generate_all_button.clicked.connect(
            self.generate_all_heats
        )

        layout.addStretch()
        layout.addWidget(title)
        layout.addWidget(subtitle)
        layout.addWidget(create_button)
        layout.addWidget(load_button)
        layout.addWidget(teams_button)
        layout.addWidget(mode_button)
        layout.addWidget(athletes_button)
        layout.addWidget(events_button)
        layout.addWidget(event_entry_button)
        layout.addWidget(results_button)
        layout.addWidget(field_results_button)
        layout.addWidget(high_jump_button)
        layout.addWidget(finals_button)
        layout.addWidget(points_button)
        layout.addWidget(standings_button)
        layout.addWidget(backup_button)
        layout.addWidget(restore_button)
        layout.addWidget(template_button)
        layout.addWidget(import_button)
        layout.addWidget(generate_all_button)
        layout.addStretch()

        self.setLayout(layout)

    def open_events(self):

        event_screen = self.stack.widget(6)

        event_screen.load_events()

        self.stack.setCurrentIndex(6)

    def open_event_entries(self):
        self.stack.widget(7).refresh_screen()
        self.stack.setCurrentIndex(7)

    def open_results(self):

        self.stack.widget(8).load_heats()

        self.stack.setCurrentIndex(8)

    def open_field_results(self):

        field_screen = (
            self.stack.widget(12)
        )

        field_screen.load_events()
        field_screen.load_event()

        self.stack.setCurrentIndex(
            12
        )


    def open_high_jump_results(self):

        high_jump_screen = (
            self.stack.widget(13)
        )

        high_jump_screen.load_events()
        high_jump_screen.load_heights()

        self.stack.setCurrentIndex(
            13
        )

    def open_points(self):

        self.stack.widget(10).load_points()

        self.stack.setCurrentIndex(10)

    def open_teams(self):

        team_screen = self.stack.widget(3)

        team_screen.load_teams()

        self.stack.setCurrentIndex(3)

    def open_athletes(self):

        athlete_screen = self.stack.widget(5)

        athlete_screen.load_teams()
        athlete_screen.load_athletes()

        self.stack.setCurrentIndex(5)

    def open_standings(self):

        self.stack.widget(11).load_standings()

        self.stack.setCurrentIndex(11)

    # =========================
    # GENERATE ALL HEATS
    # AND DIRECT FINALS
    # =========================
    def generate_all_heats(self):

        meeting_id = get_active_meeting()

        if not meeting_id:

            QMessageBox.warning(
                self,
                "No Active Meeting",
                "Select an active meeting first."
            )

            return

        # =========================
        # PREVIEW WHAT WILL HAPPEN
        # =========================

        conn = sqlite3.connect(DB_NAME)
        cur = conn.cursor()

        cur.execute("""
        SELECT
            id,
            event_name,
            gender,
            age_group,
            event_type
        FROM events
        WHERE meeting_id = ?
        ORDER BY
            age_group,
            gender,
            event_name
        """, (
            meeting_id,
        ))

        events = cur.fetchall()

        if not events:

            conn.close()

            QMessageBox.warning(
                self,
                "No Events",
                (
                    "There are no events in the "
                    "active meeting."
                )
            )

            return

        track_events_ready = 0
        non_track_events = 0
        no_entry_events = 0
        existing_data_events = 0

        # =========================
        # SCAN EVENTS FIRST
        # =========================

        for (
            event_id,
            event_name,
            gender,
            age_group,
            event_type
        ) in events:

            if event_type != "Track":

                non_track_events += 1
                continue

            # Active entries only.
            cur.execute("""
            SELECT COUNT(*)
            FROM event_entries ee

            JOIN athletes a
                ON ee.athlete_id = a.id

            WHERE ee.meeting_id = ?
            AND ee.event_id = ?
            AND ee.status = 'Active'
            AND a.status = 'Active'
            """, (
                meeting_id,
                event_id
            ))

            entry_count = (
                cur.fetchone()[0]
            )

            if entry_count == 0:

                no_entry_events += 1
                continue

            # Existing heats.
            cur.execute("""
            SELECT COUNT(*)
            FROM heats
            WHERE meeting_id = ?
            AND event_id = ?
            """, (
                meeting_id,
                event_id
            ))

            existing_heats = (
                cur.fetchone()[0]
            )

            # Existing heat results.
            cur.execute("""
            SELECT COUNT(*)
            FROM results r

            JOIN heats h
                ON r.heat_id = h.id

            WHERE h.meeting_id = ?
            AND h.event_id = ?
            """, (
                meeting_id,
                event_id
            ))

            existing_results = (
                cur.fetchone()[0]
            )

            # Existing finalists.
            cur.execute("""
            SELECT COUNT(*)
            FROM finals
            WHERE event_id = ?
            """, (
                event_id,
            ))

            existing_finalists = (
                cur.fetchone()[0]
            )

            # Existing final results.
            cur.execute("""
            SELECT COUNT(*)
            FROM final_results
            WHERE event_id = ?
            """, (
                event_id,
            ))

            existing_final_results = (
                cur.fetchone()[0]
            )

            # Existing points.
            cur.execute("""
            SELECT COUNT(*)
            FROM awarded_points
            WHERE event_id = ?
            """, (
                event_id,
            ))

            existing_points = (
                cur.fetchone()[0]
            )

            if (
                existing_heats > 0
                or existing_results > 0
                or existing_finalists > 0
                or existing_final_results > 0
                or existing_points > 0
            ):

                existing_data_events += 1
                continue

            track_events_ready += 1

        conn.close()

        if track_events_ready == 0:

            QMessageBox.information(
                self,
                "Nothing to Generate",
                (
                    "No new track events require "
                    "heat or final generation.\n\n"
                    f"Non-track events skipped: "
                    f"{non_track_events}\n"
                    f"Events with no entries: "
                    f"{no_entry_events}\n"
                    f"Already generated: "
                    f"{existing_data_events}"
                )
            )

            return

        answer = QMessageBox.question(
            self,
            "Generate All Heats / Finals",
            (
                f"{track_events_ready} track "
                "event(s) are ready.\n\n"
                "AthletiDesk will automatically:\n\n"
                "• Create a Direct Final for "
                "events with 8 or fewer athletes.\n"
                "• Create balanced heats for "
                "events with more than 8 athletes.\n"
                "• Leave Field and High Jump "
                "events unchanged.\n"
                "• Skip events that already contain "
                "heats, finals, results or points.\n\n"
                "Continue?"
            ),
            QMessageBox.Yes
            | QMessageBox.No
        )

        if answer != QMessageBox.Yes:
            return

        # =========================
        # GENERATE EVERYTHING
        # =========================

        conn = sqlite3.connect(DB_NAME)

        direct_final_events = 0
        heat_events = 0
        total_heats_created = 0
        total_athletes_processed = 0

        skipped_non_track = 0
        skipped_no_entries = 0
        skipped_existing = 0

        try:

            cur = conn.cursor()

            cur.execute(
                "BEGIN"
            )

            cur.execute("""
            SELECT
                id,
                event_name,
                gender,
                age_group,
                event_type
            FROM events
            WHERE meeting_id = ?
            ORDER BY
                age_group,
                gender,
                event_name
            """, (
                meeting_id,
            ))

            events = cur.fetchall()

            for (
                event_id,
                event_name,
                gender,
                age_group,
                event_type
            ) in events:

                # =====================
                # FIELD / HIGH JUMP
                # =====================

                if event_type != "Track":

                    skipped_non_track += 1
                    continue

                # =====================
                # ACTIVE ATHLETES
                # =====================

                cur.execute("""
                SELECT
                    ee.athlete_id
                FROM event_entries ee

                JOIN athletes a
                    ON ee.athlete_id = a.id

                WHERE ee.meeting_id = ?
                AND ee.event_id = ?
                AND ee.status = 'Active'
                AND a.status = 'Active'

                ORDER BY ee.athlete_id
                """, (
                    meeting_id,
                    event_id
                ))

                athletes = [
                    row[0]
                    for row in cur.fetchall()
                ]

                if not athletes:

                    skipped_no_entries += 1
                    continue

                # =====================
                # SAFETY CHECK
                # =====================

                cur.execute("""
                SELECT COUNT(*)
                FROM heats
                WHERE meeting_id = ?
                AND event_id = ?
                """, (
                    meeting_id,
                    event_id
                ))

                existing_heats = (
                    cur.fetchone()[0]
                )

                cur.execute("""
                SELECT COUNT(*)
                FROM results r

                JOIN heats h
                    ON r.heat_id = h.id

                WHERE h.meeting_id = ?
                AND h.event_id = ?
                """, (
                    meeting_id,
                    event_id
                ))

                existing_results = (
                    cur.fetchone()[0]
                )

                cur.execute("""
                SELECT COUNT(*)
                FROM finals
                WHERE event_id = ?
                """, (
                    event_id,
                ))

                existing_finalists = (
                    cur.fetchone()[0]
                )

                cur.execute("""
                SELECT COUNT(*)
                FROM final_results
                WHERE event_id = ?
                """, (
                    event_id,
                ))

                existing_final_results = (
                    cur.fetchone()[0]
                )

                cur.execute("""
                SELECT COUNT(*)
                FROM awarded_points
                WHERE event_id = ?
                """, (
                    event_id,
                ))

                existing_points = (
                    cur.fetchone()[0]
                )

                if (
                    existing_heats > 0
                    or existing_results > 0
                    or existing_finalists > 0
                    or existing_final_results > 0
                    or existing_points > 0
                ):

                    skipped_existing += 1
                    continue

                athlete_count = len(
                    athletes
                )

                total_athletes_processed += (
                    athlete_count
                )

                # =====================
                # DIRECT FINAL
                # 8 OR FEWER
                # =====================

                if athlete_count <= 8:

                    seed_position = 1

                    for athlete_id in athletes:

                        cur.execute("""
                        INSERT INTO finals
                        (
                            event_id,
                            athlete_id,
                            seed_position
                        )
                        VALUES (?, ?, ?)
                        """, (
                            event_id,
                            athlete_id,
                            seed_position
                        ))

                        seed_position += 1

                    direct_final_events += 1

                    continue

                # =====================
                # NORMAL HEATS
                # MORE THAN 8
                # =====================

                MAX_PER_HEAT = 8

                lane_order = [
                    4, 5, 3, 6,
                    2, 7, 1, 8
                ]

                heat_count = (
                    athlete_count
                    + MAX_PER_HEAT
                    - 1
                ) // MAX_PER_HEAT

                base_heat_size = (
                    athlete_count
                    // heat_count
                )

                extra_athletes = (
                    athlete_count
                    % heat_count
                )

                heat_sizes = []

                for heat_index in range(
                    heat_count
                ):

                    heat_size = (
                        base_heat_size
                    )

                    if (
                        heat_index
                        < extra_athletes
                    ):

                        heat_size += 1

                    heat_sizes.append(
                        heat_size
                    )

                athlete_index = 0

                for (
                    heat_number,
                    heat_size
                ) in enumerate(
                    heat_sizes,
                    start=1
                ):

                    cur.execute("""
                    INSERT INTO heats
                    (
                        meeting_id,
                        event_id,
                        heat_number
                    )
                    VALUES (?, ?, ?)
                    """, (
                        meeting_id,
                        event_id,
                        heat_number
                    ))

                    heat_id = (
                        cur.lastrowid
                    )

                    total_heats_created += 1

                    for lane_index in range(
                        heat_size
                    ):

                        athlete_id = (
                            athletes[
                                athlete_index
                            ]
                        )

                        lane_number = (
                            lane_order[
                                lane_index
                            ]
                        )

                        cur.execute("""
                        INSERT INTO heat_entries
                        (
                            heat_id,
                            athlete_id,
                            lane_number
                        )
                        VALUES (?, ?, ?)
                        """, (
                            heat_id,
                            athlete_id,
                            lane_number
                        ))

                        athlete_index += 1

                heat_events += 1

            conn.commit()

        except Exception as error:

            conn.rollback()

            QMessageBox.critical(
                self,
                "Generation Failed",
                (
                    "AthletiDesk could not "
                    "generate all heats.\n\n"
                    "The entire operation was "
                    "rolled back, so no partial "
                    "generation was saved.\n\n"
                    f"Reason:\n{error}"
                )
            )

            return

        finally:

            conn.close()

        # =========================
        # SUCCESS SUMMARY
        # =========================

        QMessageBox.information(
            self,
            "Generation Complete",
            (
                "All eligible track events "
                "have been processed.\n\n"
                f"Direct finals created: "
                f"{direct_final_events}\n"
                f"Events using heats: "
                f"{heat_events}\n"
                f"Total heats created: "
                f"{total_heats_created}\n"
                f"Athletes processed: "
                f"{total_athletes_processed}\n\n"
                f"Field / High Jump skipped: "
                f"{skipped_non_track}\n"
                f"No-entry events skipped: "
                f"{skipped_no_entries}\n"
                f"Already generated skipped: "
                f"{skipped_existing}"
            )
        )

    # =========================
    # BACKUP DATABASE
    # =========================
    def backup_database(self):

        timestamp = datetime.now().strftime(
            "%Y-%m-%d_%H%M%S"
        )

        default_name = (
            f"My_Athletics_Backup_"
            f"{timestamp}.db"
        )

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Save Athletics Backup",
            default_name,
            "SQLite Database (*.db)"
        )

        # User cancelled.
        if not file_path:
            return

        if not file_path.lower().endswith(
            ".db"
        ):

            file_path += ".db"

        # Do not allow the live database
        # itself to be selected as backup.
        live_database = os.path.abspath(
            DB_NAME
        )

        backup_database = os.path.abspath(
            file_path
        )

        if (
            live_database
            == backup_database
        ):

            QMessageBox.warning(
                self,
                "Invalid Backup Location",
                (
                    "The backup cannot overwrite "
                    "the database currently being "
                    "used by the application.\n\n"
                    "Choose another filename or "
                    "location."
                )
            )

            return

        source_conn = None
        backup_conn = None

        try:

            source_conn = sqlite3.connect(
                DB_NAME
            )

            backup_conn = sqlite3.connect(
                file_path
            )

            # SQLite's backup system makes
            # a safe, complete copy even while
            # the application database is open.
            source_conn.backup(
                backup_conn
            )

            backup_conn.commit()

        except sqlite3.Error as error:

            QMessageBox.critical(
                self,
                "Backup Failed",
                (
                    "The database could not "
                    "be backed up.\n\n"
                    f"Database error:\n{error}"
                )
            )

            return

        finally:

            if backup_conn is not None:
                backup_conn.close()

            if source_conn is not None:
                source_conn.close()

        QMessageBox.information(
            self,
            "Backup Complete",
            (
                "The athletics database was "
                "backed up successfully.\n\n"
                f"Saved as:\n{file_path}"
            )
        )

    # =========================
    # RESTORE DATABASE
    # =========================
    def restore_database(self):

        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Athletics Backup",
            "",
            "SQLite Database (*.db)"
        )

        # User cancelled.
        if not file_path:
            return

        live_database = os.path.abspath(
            DB_NAME
        )

        selected_database = os.path.abspath(
            file_path
        )

        if (
            live_database
            == selected_database
        ):

            QMessageBox.warning(
                self,
                "Invalid Restore File",
                (
                    "You selected the database "
                    "currently being used by the app.\n\n"
                    "Choose a separate backup file."
                )
            )

            return

        # =========================
        # VALIDATE BACKUP FILE
        # =========================
        source_conn = None

        try:

            source_conn = sqlite3.connect(
                file_path
            )

            cur = source_conn.cursor()

            integrity_row = cur.execute(
                "PRAGMA integrity_check"
            ).fetchone()

            if (
                not integrity_row
                or integrity_row[0].lower()
                != "ok"
            ):

                source_conn.close()
                source_conn = None

                QMessageBox.critical(
                    self,
                    "Invalid Backup",
                    (
                        "The selected database "
                        "failed the SQLite integrity "
                        "check and will not be restored."
                    )
                )

                return

            cur.execute("""
            SELECT name
            FROM sqlite_master
            WHERE type = 'table'
            """)

            existing_tables = {
                row[0]
                for row in cur.fetchall()
            }

            required_tables = {
                "meetings",
                "teams",
                "athletes",
                "events",
                "event_entries"
            }

            missing_tables = (
                required_tables
                - existing_tables
            )

            if missing_tables:

                source_conn.close()
                source_conn = None

                QMessageBox.critical(
                    self,
                    "Invalid Athletics Backup",
                    (
                        "The selected file does not "
                        "appear to be a valid "
                        "My Athletics database.\n\n"
                        "Required tables are missing."
                    )
                )

                return

        except sqlite3.Error as error:

            if source_conn is not None:
                source_conn.close()

            QMessageBox.critical(
                self,
                "Invalid Backup",
                (
                    "The selected file could not "
                    "be read as an SQLite database.\n\n"
                    f"Database error:\n{error}"
                )
            )

            return

        finally:

            if source_conn is not None:
                source_conn.close()

        # =========================
        # CONFIRM RESTORE
        # =========================
        answer = QMessageBox.question(
            self,
            "Restore Backup",
            (
                "Restoring this backup will replace "
                "the athletics data currently stored "
                "in the application.\n\n"
                "Before restoring, the app will "
                "automatically create a safety backup "
                "of the current database.\n\n"
                "After the restore, the application "
                "will close and should be opened again.\n\n"
                "Continue?"
            ),
            QMessageBox.Yes
            | QMessageBox.No
        )

        if answer != QMessageBox.Yes:
            return

        # =========================
        # CREATE SAFETY BACKUP
        # =========================
        timestamp = datetime.now().strftime(
            "%Y-%m-%d_%H%M%S"
        )

        app_folder = os.path.dirname(
            os.path.abspath(DB_NAME)
        )

        backup_folder = os.path.join(
            app_folder,
            "Backups"
        )

        try:

            os.makedirs(
                backup_folder,
                exist_ok=True
            )

        except OSError as error:

            QMessageBox.critical(
                self,
                "Restore Cancelled",
                (
                    "The safety-backup folder "
                    "could not be created.\n\n"
                    "The restore has been cancelled.\n\n"
                    f"Error:\n{error}"
                )
            )

            return

        safety_path = os.path.join(
            backup_folder,
            (
                "Before_Restore_"
                f"{timestamp}.db"
            )
        )

        current_conn = None
        safety_conn = None

        try:

            current_conn = sqlite3.connect(
                DB_NAME
            )

            safety_conn = sqlite3.connect(
                safety_path
            )

            current_conn.backup(
                safety_conn
            )

            safety_conn.commit()

        except sqlite3.Error as error:

            QMessageBox.critical(
                self,
                "Restore Cancelled",
                (
                    "A safety backup of the "
                    "current database could not "
                    "be created.\n\n"
                    "Nothing has been restored.\n\n"
                    f"Database error:\n{error}"
                )
            )

            return

        finally:

            if safety_conn is not None:
                safety_conn.close()

            if current_conn is not None:
                current_conn.close()

        # =========================
        # RESTORE SELECTED BACKUP
        # =========================
        source_conn = None
        destination_conn = None

        try:

            source_conn = sqlite3.connect(
                file_path
            )

            destination_conn = sqlite3.connect(
                DB_NAME
            )

            source_conn.backup(
                destination_conn
            )

            destination_conn.commit()

            # Check restored database.
            integrity_row = (
                destination_conn.execute(
                    "PRAGMA integrity_check"
                ).fetchone()
            )

            if (
                not integrity_row
                or integrity_row[0].lower()
                != "ok"
            ):

                raise sqlite3.DatabaseError(
                    "Restored database failed "
                    "integrity check."
                )

        except sqlite3.Error as error:

            if destination_conn is not None:
                destination_conn.close()
                destination_conn = None

            if source_conn is not None:
                source_conn.close()
                source_conn = None

            # =========================
            # AUTOMATIC RECOVERY
            # =========================
            recovery_source = None
            recovery_destination = None

            try:

                recovery_source = sqlite3.connect(
                    safety_path
                )

                recovery_destination = sqlite3.connect(
                    DB_NAME
                )

                recovery_source.backup(
                    recovery_destination
                )

                recovery_destination.commit()

                recovery_message = (
                    "The restore failed, but the "
                    "previous database was recovered "
                    "successfully from the automatic "
                    "safety backup."
                )

            except sqlite3.Error:

                recovery_message = (
                    "The restore failed and automatic "
                    "recovery also failed.\n\n"
                    "Your safety backup is located at:\n"
                    f"{safety_path}"
                )

            finally:

                if recovery_destination is not None:
                    recovery_destination.close()

                if recovery_source is not None:
                    recovery_source.close()

            QMessageBox.critical(
                self,
                "Restore Failed",
                (
                    f"{recovery_message}\n\n"
                    f"Database error:\n{error}"
                )
            )

            return

        finally:

            if destination_conn is not None:
                destination_conn.close()

            if source_conn is not None:
                source_conn.close()

        QMessageBox.information(
            self,
            "Restore Complete",
            (
                "The athletics backup was "
                "restored successfully.\n\n"
                "A safety copy of the previous "
                "database was saved as:\n\n"
                f"{safety_path}\n\n"
                "The application will now close.\n"
                "Open it again to use the "
                "restored data."
            )
        )

        QApplication.quit()


    # =========================
    # DOWNLOAD ENTRY TEMPLATE
    # =========================
    def download_entry_template(self):

        template_path = resource_path(
            os.path.join(
                "assets",
                "templates",
                "AthletiDesk_Entry_Template.xlsx"
            )
        )

        # Make sure the bundled
        # master template exists.
        if not os.path.exists(
            template_path
        ):

            QMessageBox.critical(
                self,
                "Template Not Found",
                (
                    "The AthletiDesk entry template "
                    "could not be found.\n\n"
                    "Expected location:\n"
                    f"{template_path}"
                )
            )

            return

        default_name = (
            "AthletiDesk_Entry_Template.xlsx"
        )

        save_path, _ = (
            QFileDialog.getSaveFileName(
                self,
                "Save Entry Template",
                default_name,
                "Excel Workbook (*.xlsx)"
            )
        )

        # User cancelled.
        if not save_path:
            return

        if not save_path.lower().endswith(
            ".xlsx"
        ):

            save_path += ".xlsx"

        # Prevent saving over the
        # master template itself.
        if (
            os.path.abspath(save_path)
            == os.path.abspath(template_path)
        ):

            QMessageBox.warning(
                self,
                "Invalid Save Location",
                (
                    "The master AthletiDesk "
                    "template cannot be "
                    "overwritten.\n\n"
                    "Choose another location."
                )
            )

            return

        try:

            shutil.copy2(
                template_path,
                save_path
            )

        except OSError as error:

            QMessageBox.critical(
                self,
                "Template Export Failed",
                (
                    "The entry template could "
                    "not be saved.\n\n"
                    f"Error:\n{error}"
                )
            )

            return

        answer = QMessageBox.question(
            self,
            "Template Saved",
            (
                "The AthletiDesk entry template "
                "was saved successfully.\n\n"
                f"Saved as:\n{save_path}\n\n"
                "Would you like to open it now?"
            ),
            QMessageBox.Yes
            | QMessageBox.No
        )

        if answer == QMessageBox.Yes:

            try:

                os.startfile(
                    save_path
                )

            except OSError as error:

                QMessageBox.warning(
                    self,
                    "Could Not Open File",
                    (
                        "The template was saved "
                        "successfully, but Windows "
                        "could not open it.\n\n"
                        f"Error:\n{error}"
                    )
                )

    # =========================
    # IMPORT ENTRY FILES
    # VALIDATION + PREVIEW
    # =========================
    def import_entry_files(self):

        meeting_id = get_active_meeting()

        if not meeting_id:

            QMessageBox.warning(
                self,
                "No Active Meeting",
                (
                    "Select an active meeting "
                    "before importing entry files."
                )
            )

            return

        file_paths, _ = (
            QFileDialog.getOpenFileNames(
                self,
                "Select AthletiDesk Entry Files",
                "",
                "Excel Workbooks (*.xlsx)"
            )
        )

        if not file_paths:
            return

        # =========================
        # TEMPLATE DEFINITIONS
        # =========================

        required_sheets = [
            "General",
            "Male entries",
            "Female entries"
        ]

        expected_headers = [
            "Entry Number",
            "Name",
            "Surname",
            "Date of Birth",
            "Age Group",
            "Item 1",
            "Item 2",
            "Item 3",
            "Item 4",
            "Item 5",
            "Item 6",
            "Item 7",
            "Item 8",
            "Item 9",
            "Relay"
        ]

        age_groups = [
            "U7",
            "U8",
            "U9",
            "U10",
            "U11",
            "U12",
            "U13",
            "U14",
            "U15",
            "U16",
            "U17",
            "U18",
            "U19"
        ]

        full_event_catalogue = [
            "60m",
            "80m",
            "100m",
            "150m",
            "200m",
            "300m",
            "400m",
            "600m",
            "800m",
            "1000m",
            "1200m",
            "1500m",
            "2000m",
            "3000m",
            "5000m",
            "10000m",
            "Long Jump",
            "High Jump",
            "Triple Jump",
            "Pole Vault",
            "Shot Put",
            "Discus",
            "Javelin",
            "Hammer Throw",
            "70m Hurdles",
            "75m Hurdles",
            "80m Hurdles",
            "90m Hurdles",
            "100m Hurdles",
            "110m Hurdles",
            "150m Hurdles",
            "200m Hurdles",
            "300m Hurdles",
            "400m Hurdles",
            "1000m Race Walk",
            "1500m Race Walk",
            "3000m Race Walk",
            "5000m Race Walk",
            "10km Race Walk",
            "20km Race Walk",
            "50km Race Walk",
            "1000m Steeplechase",
            "1500m Steeplechase",
            "2000m Steeplechase",
            "3000m Steeplechase",
            "Turbo Javelin",
            "Cricket Ball",
            "5km Road",
            "10km Road",
            "Half Marathon",
            "Marathon",
            "1 Mile"
        ]

        male_events = {
            "U7": [
                "60m",
                "80m"
            ],

            "U8": [
                "60m",
                "80m"
            ],

            "U9": [
                "60m",
                "80m",
                "600m"
            ],

            "U10": [
                "80m",
                "100m",
                "1200m",
                "70m Hurdles",
                "High Jump",
                "Long Jump",
                "Shot Put"
            ],

            "U11": [
                "80m",
                "100m",
                "1200m",
                "70m Hurdles",
                "High Jump",
                "Long Jump",
                "Shot Put"
            ],

            "U12": [
                "100m",
                "150m",
                "1200m",
                "75m Hurdles",
                "High Jump",
                "Long Jump",
                "Shot Put",
                "Discus",
                "Javelin"
            ],

            "U13": [
                "100m",
                "200m",
                "800m",
                "1500m",
                "80m Hurdles",
                "High Jump",
                "Long Jump",
                "Javelin"
            ]
        }

        female_events = {
            "U7": [
                "60m",
                "80m"
            ],

            "U8": [
                "60m",
                "80m"
            ],

            "U9": [
                "60m",
                "80m",
                "600m"
            ],

            "U10": [
                "80m",
                "100m",
                "1200m",
                "70m Hurdles",
                "High Jump",
                "Long Jump",
                "Shot Put"
            ],

            "U11": [
                "80m",
                "100m",
                "1200m",
                "70m Hurdles",
                "High Jump",
                "Long Jump",
                "Shot Put"
            ],

            "U12": [
                "100m",
                "150m",
                "1200m",
                "75m Hurdles",
                "High Jump",
                "Long Jump",
                "Shot Put",
                "Discus",
                "Javelin"
            ],

            "U13": [
                "100m",
                "200m",
                "800m",
                "1500m",
                "75m Hurdles",
                "High Jump",
                "Long Jump",
                "Javelin"
            ]
        }

        # We have not yet locked down
        # high-school restrictions.
        # For now use the full catalogue.
        for age_group in [
            "U14",
            "U15",
            "U16",
            "U17",
            "U18",
            "U19"
        ]:

            male_events[
                age_group
            ] = full_event_catalogue

            female_events[
                age_group
            ] = full_event_catalogue

        helper_columns = [
            27,  # AA
            28,  # AB
            29,  # AC
            30,  # AD
            31,  # AE
            32,  # AF
            33,  # AG
            34,  # AH
            35,  # AI
            36,  # AJ
            37,  # AK
            38,  # AL
            39   # AM
        ]

        errors = []
        warnings = []

        athlete_records = []

        files_passed = 0
        total_rows = 0
        valid_rows = 0

        male_count = 0
        female_count = 0

        event_selection_count = 0
        relay_count = 0

        team_names = set()

        # =========================
        # HELPERS
        # =========================

        def clean_text(value):

            if value is None:
                return ""

            return str(value).strip()

        def normalise_date(value):

            if value is None:
                return None

            # Excel date / Python date
            if (
                hasattr(
                    value,
                    "strftime"
                )
                and not isinstance(
                    value,
                    str
                )
            ):

                try:

                    return value.strftime(
                        "%Y-%m-%d"
                    )

                except ValueError:
                    return None

            text = clean_text(
                value
            )

            if not text:
                return None

            formats = [
                "%d/%m/%Y",
                "%Y-%m-%d",
                "%d-%m-%Y",
                "%d.%m.%Y"
            ]

            for date_format in formats:

                try:

                    return (
                        datetime.strptime(
                            text,
                            date_format
                        ).strftime(
                            "%Y-%m-%d"
                        )
                    )

                except ValueError:
                    continue

            return None

        # =========================
        # PROCESS FILES
        # =========================

        for file_path in file_paths:

            workbook = None

            file_name = os.path.basename(
                file_path
            )

            file_errors_before = len(
                errors
            )

            try:

                workbook = load_workbook(
                    file_path,
                    data_only=False,
                    read_only=True
                )

            except (
                InvalidFileException,
                BadZipFile,
                OSError,
                ValueError
            ) as error:

                errors.append(
                    (
                        f"{file_name}: "
                        "could not be read as "
                        "an AthletiDesk workbook. "
                        f"{error}"
                    )
                )

                continue

            try:

                # =====================
                # REQUIRED SHEETS
                # =====================

                missing_sheets = []

                for sheet_name in required_sheets:

                    if (
                        sheet_name
                        not in workbook.sheetnames
                    ):

                        missing_sheets.append(
                            sheet_name
                        )

                if missing_sheets:

                    for sheet_name in missing_sheets:

                        errors.append(
                            (
                                f"{file_name}: "
                                "missing required sheet "
                                f"'{sheet_name}'."
                            )
                        )

                    continue

                general = workbook[
                    "General"
                ]

                # =====================
                # GENERAL STRUCTURE
                # =====================

                general_checks = {
                    "A1": (
                        "ATHLETICS ENTRY TEMPLATE"
                    ),
                    "A4": "Meeting Name",
                    "A5": "Meeting Date",
                    "A6": (
                        "Team / House / "
                        "School Name"
                    ),
                    "A7": (
                        "Team Code / "
                        "Abbreviation"
                    ),
                    "A8": (
                        "Teacher / Coach / "
                        "Organiser"
                    ),
                    "A9": "Contact Number",
                    "A10": "Email"
                }

                structure_failed = False

                for (
                    cell_reference,
                    expected_value
                ) in general_checks.items():

                    if (
                        general[
                            cell_reference
                        ].value
                        != expected_value
                    ):

                        errors.append(
                            (
                                f"{file_name}: "
                                "General sheet structure "
                                "was changed at "
                                f"{cell_reference}."
                            )
                        )

                        structure_failed = True

                # =====================
                # TEAM INFORMATION
                # =====================

                team_name = clean_text(
                    general["B6"].value
                )

                if not team_name:

                    errors.append(
                        (
                            f"{file_name}: "
                            "Team / House / School "
                            "Name is missing on the "
                            "General sheet."
                        )
                    )

                else:

                    team_names.add(
                        team_name
                    )

                # =====================
                # ENTRY SHEETS
                # =====================

                entry_sheet_details = [
                    (
                        "Male entries",
                        "Male",
                        "MALE ATHLETE ENTRIES",
                        male_events
                    ),
                    (
                        "Female entries",
                        "Female",
                        "FEMALE ATHLETE ENTRIES",
                        female_events
                    )
                ]

                for (
                    sheet_name,
                    gender,
                    expected_title,
                    event_rules
                ) in entry_sheet_details:

                    sheet = workbook[
                        sheet_name
                    ]

                    # -----------------
                    # TITLE
                    # -----------------

                    if (
                        sheet["A1"].value
                        != expected_title
                    ):

                        errors.append(
                            (
                                f"{file_name} - "
                                f"{sheet_name}: "
                                "sheet title was changed."
                            )
                        )

                        structure_failed = True

                    # -----------------
                    # HEADINGS
                    # -----------------

                    actual_headers = [
                        sheet.cell(
                            row=4,
                            column=column
                        ).value
                        for column in range(
                            1,
                            16
                        )
                    ]

                    if (
                        actual_headers
                        != expected_headers
                    ):

                        errors.append(
                            (
                                f"{file_name} - "
                                f"{sheet_name}: "
                                "column headings "
                                "were changed."
                            )
                        )

                        structure_failed = True

                    # -----------------
                    # SYSTEM MARKER
                    # -----------------

                    if (
                        sheet["AA1"].value
                        != (
                            "SYSTEM EVENT LISTS "
                            "- DO NOT EDIT"
                        )
                    ):

                        errors.append(
                            (
                                f"{file_name} - "
                                f"{sheet_name}: "
                                "system event data "
                                "was changed."
                            )
                        )

                        structure_failed = True

                    # -----------------
                    # AGE HEADINGS
                    # -----------------

                    system_age_groups = [
                        clean_text(
                            sheet.cell(
                                row=2,
                                column=column
                            ).value
                        )
                        for column in (
                            helper_columns
                        )
                    ]

                    if (
                        system_age_groups
                        != age_groups
                    ):

                        errors.append(
                            (
                                f"{file_name} - "
                                f"{sheet_name}: "
                                "system age-group data "
                                "was changed."
                            )
                        )

                        structure_failed = True

                    # -----------------
                    # SYSTEM EVENT LISTS
                    # -----------------

                    for (
                        index,
                        age_group
                    ) in enumerate(
                        age_groups
                    ):

                        expected_events = (
                            event_rules[
                                age_group
                            ]
                        )

                        column_number = (
                            helper_columns[
                                index
                            ]
                        )

                        actual_events = []

                        for event_index in range(
                            len(
                                expected_events
                            )
                        ):

                            value = clean_text(
                                sheet.cell(
                                    row=(
                                        3
                                        + event_index
                                    ),
                                    column=(
                                        column_number
                                    )
                                ).value
                            )

                            actual_events.append(
                                value
                            )

                        if (
                            actual_events
                            != expected_events
                        ):

                            errors.append(
                                (
                                    f"{file_name} - "
                                    f"{sheet_name}: "
                                    f"{age_group} system "
                                    "event list was changed."
                                )
                            )

                            structure_failed = True

                # Do not trust athlete data
                # if the workbook structure
                # itself has been altered.
                if structure_failed:
                    continue

                # =====================
                # READ ATHLETE DATA
                # =====================

                for (
                    sheet_name,
                    gender,
                    expected_title,
                    event_rules
                ) in entry_sheet_details:

                    sheet = workbook[
                        sheet_name
                    ]

                    # Template contains
                    # entry rows 5 - 204.
                    for row_number in range(
                        5,
                        205
                    ):

                        row_values = [
                            sheet.cell(
                                row=row_number,
                                column=column
                            ).value
                            for column in range(
                                1,
                                16
                            )
                        ]

                        # Completely empty row.
                        if all(
                            clean_text(value) == ""
                            for value in row_values
                        ):

                            continue

                        total_rows += 1

                        row_error_count = len(
                            errors
                        )

                        entry_number = clean_text(
                            row_values[0]
                        )

                        first_name = clean_text(
                            row_values[1]
                        )

                        surname = clean_text(
                            row_values[2]
                        )

                        date_of_birth = (
                            normalise_date(
                                row_values[3]
                            )
                        )

                        age_group = clean_text(
                            row_values[4]
                        )

                        raw_items = (
                            row_values[
                                5:14
                            ]
                        )

                        relay_value = clean_text(
                            row_values[14]
                        )

                        row_label = (
                            f"{file_name} - "
                            f"{sheet_name}, "
                            f"row {row_number}"
                        )

                        # =================
                        # REQUIRED DETAILS
                        # =================

                        if not first_name:

                            errors.append(
                                (
                                    f"{row_label}: "
                                    "Name is missing."
                                )
                            )

                        if not surname:

                            errors.append(
                                (
                                    f"{row_label}: "
                                    "Surname is missing."
                                )
                            )

                        if not date_of_birth:

                            errors.append(
                                (
                                    f"{row_label}: "
                                    "Date of Birth is "
                                    "missing or invalid."
                                )
                            )

                        if (
                            age_group
                            not in age_groups
                        ):

                            errors.append(
                                (
                                    f"{row_label}: "
                                    "Age Group is missing "
                                    "or invalid."
                                )
                            )

                        # Entry number is useful,
                        # but not required yet.
                        if not entry_number:

                            warnings.append(
                                (
                                    f"{row_label}: "
                                    "Entry Number is blank."
                                )
                            )

                        # =================
                        # ITEM ORDER
                        # =================

                        selected_items = []

                        blank_found = False

                        for (
                            item_index,
                            raw_item
                        ) in enumerate(
                            raw_items,
                            start=1
                        ):

                            item = clean_text(
                                raw_item
                            )

                            if not item:

                                blank_found = True
                                continue

                            if blank_found:

                                errors.append(
                                    (
                                        f"{row_label}: "
                                        f"Item {item_index} "
                                        "contains an event "
                                        "after an earlier "
                                        "Item was left blank."
                                    )
                                )

                            selected_items.append(
                                item
                            )

                        # =================
                        # DUPLICATE ITEMS
                        # =================

                        normalised_items = [
                            item.casefold()
                            for item
                            in selected_items
                        ]

                        if (
                            len(
                                normalised_items
                            )
                            != len(
                                set(
                                    normalised_items
                                )
                            )
                        ):

                            errors.append(
                                (
                                    f"{row_label}: "
                                    "the same event was "
                                    "selected more than once."
                                )
                            )

                        # =================
                        # EVENT ELIGIBILITY
                        # =================

                        if (
                            age_group
                            in event_rules
                        ):

                            allowed_events = (
                                event_rules[
                                    age_group
                                ]
                            )

                            for item in selected_items:

                                if (
                                    item
                                    not in allowed_events
                                ):

                                    errors.append(
                                        (
                                            f"{row_label}: "
                                            f"'{item}' is not "
                                            "allowed for "
                                            f"{gender} "
                                            f"{age_group}."
                                        )
                                    )

                        # =================
                        # RELAY
                        # =================

                        relay_normalised = (
                            relay_value.casefold()
                        )

                        if (
                            relay_normalised
                            not in (
                                "",
                                "yes",
                                "no"
                            )
                        ):

                            errors.append(
                                (
                                    f"{row_label}: "
                                    "Relay must be "
                                    "Yes, No or blank."
                                )
                            )

                        # =================
                        # STORE VALID ROW
                        # =================

                        if (
                            len(errors)
                            == row_error_count
                        ):

                            valid_rows += 1

                            if gender == "Male":
                                male_count += 1
                            else:
                                female_count += 1

                            event_selection_count += (
                                len(
                                    selected_items
                                )
                            )

                            if (
                                relay_normalised
                                == "yes"
                            ):

                                relay_count += 1

                            athlete_records.append({
                                "file": file_name,
                                "sheet": sheet_name,
                                "row": row_number,
                                "entry_number": (
                                    entry_number
                                ),
                                "first_name": (
                                    first_name
                                ),
                                "surname": surname,
                                "dob": date_of_birth,
                                "gender": gender,
                                "age_group": (
                                    age_group
                                ),
                                "team": team_name,
                                "items": (
                                    selected_items
                                ),
                                "relay": (
                                    relay_normalised
                                    == "yes"
                                )
                            })

                if (
                    len(errors)
                    == file_errors_before
                ):

                    files_passed += 1

            finally:

                workbook.close()

        # =========================
        # DUPLICATE ATHLETES
        # =========================

        merged_athletes = {}

        duplicate_records = 0

        for record in athlete_records:

            athlete_key = (
                record[
                    "first_name"
                ].casefold(),
                record[
                    "surname"
                ].casefold(),
                record[
                    "dob"
                ]
            )

            if (
                athlete_key
                not in merged_athletes
            ):

                merged_athletes[
                    athlete_key
                ] = {
                    "record": record,
                    "items": set(
                        record[
                            "items"
                        ]
                    ),
                    "relay": record[
                        "relay"
                    ]
                }

                continue

            duplicate_records += 1

            existing = (
                merged_athletes[
                    athlete_key
                ]
            )

            original = existing[
                "record"
            ]

            # Same person but
            # contradictory information.
            if (
                original[
                    "gender"
                ]
                != record[
                    "gender"
                ]
                or original[
                    "age_group"
                ]
                != record[
                    "age_group"
                ]
                or original[
                    "team"
                ].casefold()
                != record[
                    "team"
                ].casefold()
            ):

                errors.append(
                    (
                        f"{record['first_name']} "
                        f"{record['surname']} "
                        f"({record['dob']}): "
                        "duplicate athlete records "
                        "contain conflicting gender, "
                        "age group or team information."
                    )
                )

                continue

            # Same athlete repeated.
            # Later we will merge the
            # event selections.
            warnings.append(
                (
                    f"{record['first_name']} "
                    f"{record['surname']} "
                    f"({record['dob']}): "
                    "appears more than once. "
                    "The final importer will merge "
                    "these records."
                )
            )

            existing[
                "items"
            ].update(
                record[
                    "items"
                ]
            )

            if record[
                "relay"
            ]:

                existing[
                    "relay"
                ] = True

        unique_athletes = len(
            merged_athletes
        )

        merged_event_entries = sum(
            len(
                athlete_data[
                    "items"
                ]
            )
            for athlete_data
            in merged_athletes.values()
        )

        merged_relay_athletes = sum(
            1
            for athlete_data
            in merged_athletes.values()
            if athlete_data[
                "relay"
            ]
        )

        # =========================
        # PREVIEW MESSAGE
        # =========================

        preview_lines = [
            "ATHLETIDESK IMPORT PREVIEW",
            "",
            (
                f"Files selected: "
                f"{len(file_paths)}"
            ),
            (
                f"Files passing structure check: "
                f"{files_passed}"
            ),
            (
                f"Teams / Houses / Schools: "
                f"{len(team_names)}"
            ),
            "",
            (
                f"Athlete rows found: "
                f"{total_rows}"
            ),
            (
                f"Valid athlete rows: "
                f"{valid_rows}"
            ),
            (
                f"Male entries: "
                f"{male_count}"
            ),
            (
                f"Female entries: "
                f"{female_count}"
            ),
            (
                f"Unique athletes: "
                f"{unique_athletes}"
            ),
            (
                f"Duplicate records: "
                f"{duplicate_records}"
            ),
            "",
            (
                f"Individual event entries: "
                f"{merged_event_entries}"
            ),
            (
                f"Relay athletes: "
                f"{merged_relay_athletes}"
            ),
            "",
            (
                f"Warnings: "
                f"{len(warnings)}"
            ),
            (
                f"Errors: "
                f"{len(errors)}"
            )
        ]

        # =========================
        # ERRORS
        # =========================

        if errors:

            preview_lines.extend([
                "",
                "FIRST ERRORS FOUND:"
            ])

            for error in errors[:15]:

                preview_lines.append(
                    f"• {error}"
                )

            if len(errors) > 15:

                preview_lines.append(
                    (
                        f"• ...and "
                        f"{len(errors) - 15} "
                        "more error(s)."
                    )
                )

            preview_lines.extend([
                "",
                (
                    "No athlete data has "
                    "been imported."
                )
            ])

            QMessageBox.warning(
                self,
                "Import Preview - Problems Found",
                "\n".join(
                    preview_lines
                )
            )

            return

        # =========================
        # WARNINGS
        # =========================

        if warnings:

            preview_lines.extend([
                "",
                "WARNINGS:"
            ])

            for warning in warnings[:10]:

                preview_lines.append(
                    f"• {warning}"
                )

            if len(warnings) > 10:

                preview_lines.append(
                    (
                        f"• ...and "
                        f"{len(warnings) - 10} "
                        "more warning(s)."
                    )
                )

        preview_lines.extend([
            "",
            (
                "All athlete and event data "
                "passed validation."
            ),
            "",
            (
                "No data has been imported yet."
            )
        ])

        # =========================
        # CONFIRM DATABASE IMPORT
        # =========================

        if unique_athletes == 0:

            QMessageBox.warning(
                self,
                "Nothing to Import",
                (
                    "The files passed validation, "
                    "but no athlete records were found."
                )
            )

            return

        answer = QMessageBox.question(
            self,
            "Import Entries",
            (
                "\n".join(preview_lines)
                + "\n\n"
                + "Import the teams, athletes "
                + "and individual event entries "
                + "into the active meeting now?"
            ),
            QMessageBox.Yes
            | QMessageBox.No
        )

        if answer != QMessageBox.Yes:
            return

        # =========================
        # DATABASE TRANSACTION
        # =========================

        conn = sqlite3.connect(
            DB_NAME
        )

        created_teams = 0
        created_athletes = 0
        reused_athletes = 0

        created_events = 0

        created_event_entries = 0
        existing_event_entries = 0

        relay_athletes_not_imported = 0

        try:

            cur = conn.cursor()

            cur.execute(
                "BEGIN"
            )

            # =====================
            # MEETING TYPE
            # =====================

            cur.execute("""
            SELECT team_mode
            FROM meetings
            WHERE id = ?
            """, (
                meeting_id,
            ))

            meeting_row = (
                cur.fetchone()
            )

            if not meeting_row:

                raise RuntimeError(
                    (
                        "The active meeting "
                        "could not be found."
                    )
                )

            team_type = (
                meeting_row[0]
                or "School"
            )

            # =====================
            # EVENT TYPE HELPER
            # =====================

            measured_field_events = {
                "Long Jump",
                "Triple Jump",
                "Shot Put",
                "Discus",
                "Javelin",
                "Hammer Throw",
                "Turbo Javelin",
                "Cricket Ball"
            }

            def get_event_type(
                event_name
            ):

                if (
                    event_name
                    == "High Jump"
                ):

                    return (
                        "High Jump"
                    )

                if (
                    event_name
                    in measured_field_events
                ):

                    return "Field"

                if (
                    event_name
                    == "Pole Vault"
                ):

                    raise RuntimeError(
                        (
                            "Pole Vault is not "
                            "supported by the "
                            "current AthletiDesk "
                            "results system yet."
                        )
                    )

                return "Track"

            # =====================
            # EXISTING TEAMS
            # =====================

            cur.execute("""
            SELECT
                id,
                name
            FROM teams
            WHERE meeting_id = ?
            """, (
                meeting_id,
            ))

            team_ids = {}

            for (
                existing_team_id,
                existing_team_name
            ) in cur.fetchall():

                team_key = (
                    existing_team_name
                    .strip()
                    .casefold()
                )

                if (
                    team_key
                    in team_ids
                    and team_ids[
                        team_key
                    ]
                    != existing_team_id
                ):

                    raise RuntimeError(
                        (
                            "Duplicate teams with "
                            "the same name already "
                            "exist in this meeting: "
                            f"{existing_team_name}"
                        )
                    )

                team_ids[
                    team_key
                ] = (
                    existing_team_id
                )

            # =====================
            # CREATE MISSING TEAMS
            # =====================

            for team_name in sorted(
                team_names
            ):

                team_key = (
                    team_name
                    .strip()
                    .casefold()
                )

                if (
                    team_key
                    in team_ids
                ):

                    continue

                cur.execute("""
                INSERT INTO teams
                (
                    meeting_id,
                    name,
                    team_type
                )
                VALUES (?, ?, ?)
                """, (
                    meeting_id,
                    team_name,
                    team_type
                ))

                team_ids[
                    team_key
                ] = (
                    cur.lastrowid
                )

                created_teams += 1

            # =====================
            # CREATE / REUSE
            # ATHLETES
            # =====================

            athlete_ids = {}

            for athlete_key, (
                athlete_data
            ) in (
                merged_athletes.items()
            ):

                record = (
                    athlete_data[
                        "record"
                    ]
                )

                team_key = (
                    record[
                        "team"
                    ]
                    .strip()
                    .casefold()
                )

                team_id = (
                    team_ids.get(
                        team_key
                    )
                )

                if not team_id:

                    raise RuntimeError(
                        (
                            "Could not find or "
                            "create team: "
                            f"{record['team']}"
                        )
                    )

                first_name = (
                    record[
                        "first_name"
                    ]
                )

                surname = (
                    record[
                        "surname"
                    ]
                )

                dob = (
                    record[
                        "dob"
                    ]
                )

                full_name = (
                    f"{first_name} "
                    f"{surname}"
                ).strip()

                # =================
                # EXISTING ATHLETE
                # =================

                cur.execute("""
                SELECT
                    id,
                    gender,
                    age_group,
                    team_id,
                    entry_number
                FROM athletes
                WHERE meeting_id = ?
                AND LOWER(
                    COALESCE(
                        first_name,
                        ''
                    )
                ) = LOWER(?)
                AND LOWER(
                    COALESCE(
                        surname,
                        ''
                    )
                ) = LOWER(?)
                AND date_of_birth = ?
                """, (
                    meeting_id,
                    first_name,
                    surname,
                    dob
                ))

                existing_rows = (
                    cur.fetchall()
                )

                if (
                    len(
                        existing_rows
                    )
                    > 1
                ):

                    raise RuntimeError(
                        (
                            "More than one existing "
                            "athlete matches:\n"
                            f"{full_name} "
                            f"({dob})"
                        )
                    )

                if existing_rows:

                    (
                        athlete_id,
                        old_gender,
                        old_age_group,
                        old_team_id,
                        old_entry_number
                    ) = existing_rows[0]

                    if (
                        old_gender
                        != record[
                            "gender"
                        ]
                        or old_age_group
                        != record[
                            "age_group"
                        ]
                        or old_team_id
                        != team_id
                    ):

                        raise RuntimeError(
                            (
                                f"{full_name} "
                                f"({dob}) already "
                                "exists, but their "
                                "gender, age group "
                                "or team is different."
                            )
                        )

                    new_entry_number = (
                        record[
                            "entry_number"
                        ]
                    )

                    if (
                        old_entry_number
                        and new_entry_number
                        and str(
                            old_entry_number
                        ).strip()
                        != str(
                            new_entry_number
                        ).strip()
                    ):

                        raise RuntimeError(
                            (
                                f"{full_name} "
                                "already has a "
                                "different Entry "
                                "Number."
                            )
                        )

                    if (
                        not old_entry_number
                        and new_entry_number
                    ):

                        cur.execute("""
                        UPDATE athletes
                        SET entry_number = ?
                        WHERE id = ?
                        """, (
                            new_entry_number,
                            athlete_id
                        ))

                    reused_athletes += 1

                else:

                    # =================
                    # INSERT ATHLETE
                    # =================

                    cur.execute("""
                    INSERT INTO athletes
                    (
                        meeting_id,
                        full_name,
                        gender,
                        age_group,
                        team_id,
                        entry_number,
                        first_name,
                        surname,
                        date_of_birth
                    )
                    VALUES (
                        ?, ?, ?, ?, ?,
                        ?, ?, ?, ?
                    )
                    """, (
                        meeting_id,
                        full_name,
                        record[
                            "gender"
                        ],
                        record[
                            "age_group"
                        ],
                        team_id,
                        (
                            record[
                                "entry_number"
                            ]
                            or None
                        ),
                        first_name,
                        surname,
                        dob
                    ))

                    athlete_id = (
                        cur.lastrowid
                    )

                    created_athletes += 1

                athlete_ids[
                    athlete_key
                ] = athlete_id

            # =========================
            # CREATE / FIND EVENTS
            # AND EVENT ENTRIES
            # =========================

            event_ids = {}

            for athlete_key, (
                athlete_data
            ) in (
                merged_athletes.items()
            ):

                record = (
                    athlete_data[
                        "record"
                    ]
                )

                athlete_id = (
                    athlete_ids[
                        athlete_key
                    ]
                )

                gender = (
                    record[
                        "gender"
                    ]
                )

                age_group = (
                    record[
                        "age_group"
                    ]
                )

                # merged_athletes stores
                # selections as a set.
                selected_items = sorted(
                    athlete_data[
                        "items"
                    ]
                )

                for event_name in (
                    selected_items
                ):

                    expected_type = (
                        get_event_type(
                            event_name
                        )
                    )

                    event_key = (
                        event_name.casefold(),
                        gender,
                        age_group
                    )

                    event_id = (
                        event_ids.get(
                            event_key
                        )
                    )

                    # =================
                    # FIND EVENT
                    # =================

                    if not event_id:

                        cur.execute("""
                        SELECT
                            id,
                            event_type
                        FROM events
                        WHERE meeting_id = ?
                        AND LOWER(
                            event_name
                        ) = LOWER(?)
                        AND gender = ?
                        AND age_group = ?
                        """, (
                            meeting_id,
                            event_name,
                            gender,
                            age_group
                        ))

                        event_rows = (
                            cur.fetchall()
                        )

                        if (
                            len(
                                event_rows
                            )
                            > 1
                        ):

                            raise RuntimeError(
                                (
                                    "More than one "
                                    "matching event exists:\n"
                                    f"{event_name} "
                                    f"({gender} "
                                    f"{age_group})"
                                )
                            )

                        if event_rows:

                            (
                                event_id,
                                existing_type
                            ) = (
                                event_rows[0]
                            )

                            if (
                                existing_type
                                != expected_type
                            ):

                                raise RuntimeError(
                                    (
                                        f"{event_name} "
                                        f"({gender} "
                                        f"{age_group}) "
                                        "already exists as "
                                        f"'{existing_type}', "
                                        "but the importer "
                                        "expects "
                                        f"'{expected_type}'."
                                    )
                                )

                        else:

                            # =================
                            # CREATE EVENT
                            # =================

                            cur.execute("""
                            INSERT INTO events
                            (
                                meeting_id,
                                event_name,
                                event_type,
                                gender,
                                age_group
                            )
                            VALUES (?, ?, ?, ?, ?)
                            """, (
                                meeting_id,
                                event_name,
                                expected_type,
                                gender,
                                age_group
                            ))

                            event_id = (
                                cur.lastrowid
                            )

                            created_events += 1

                        event_ids[
                            event_key
                        ] = event_id

                    # =================
                    # EVENT ENTRY
                    # =================

                    cur.execute("""
                    SELECT id
                    FROM event_entries
                    WHERE meeting_id = ?
                    AND event_id = ?
                    AND athlete_id = ?
                    """, (
                        meeting_id,
                        event_id,
                        athlete_id
                    ))

                    if cur.fetchone():

                        existing_event_entries += 1

                        continue

                    cur.execute("""
                    INSERT INTO event_entries
                    (
                        meeting_id,
                        event_id,
                        athlete_id
                    )
                    VALUES (?, ?, ?)
                    """, (
                        meeting_id,
                        event_id,
                        athlete_id
                    ))

                    created_event_entries += 1

                if (
                    athlete_data[
                        "relay"
                    ]
                ):

                    relay_athletes_not_imported += 1

            # =====================
            # EVERYTHING PASSED
            # =====================

            conn.commit()

        except Exception as error:

            conn.rollback()

            QMessageBox.critical(
                self,
                "Import Failed",
                (
                    "The import could not "
                    "be completed.\n\n"
                    "No changes from this "
                    "import were saved.\n\n"
                    f"Reason:\n{error}"
                )
            )

            return

        finally:

            conn.close()

        # =========================
        # SUCCESS
        # =========================

        success_message = (
            "The entry import "
            "completed successfully.\n\n"
            f"Teams created: "
            f"{created_teams}\n"
            f"Athletes created: "
            f"{created_athletes}\n"
            f"Existing athletes reused: "
            f"{reused_athletes}\n"
            f"Events created: "
            f"{created_events}\n"
            f"Event entries created: "
            f"{created_event_entries}\n"
            f"Existing event entries reused: "
            f"{existing_event_entries}"
        )

        if (
            relay_athletes_not_imported
            > 0
        ):

            success_message += (
                "\n\n"
                f"Relay athletes detected: "
                f"{relay_athletes_not_imported}\n"
                "Relay entries were not "
                "imported yet."
            )

        QMessageBox.information(
            self,
            "Entries Imported",
            success_message
        )

# =========================
# MAIN WINDOWpy main.py
# =========================
class MainWindow(QMainWindow):

    def __init__(self):
        super().__init__()

        self.setWindowTitle("My Athletics")
        self.resize(1000, 700)

        self.stack = QStackedWidget()

        self.dashboard = DashboardScreen(self.stack)
        self.create_meeting = CreateMeetingScreen()
        self.load_meeting = LoadMeetingScreen()
        self.teams = TeamManagementScreen()
        self.user_mode = UserModeScreen()
        self.athletes = AthleteScreen()
        self.events = EventScreen()
        self.event_entry = EventEntryScreen()
        self.results = ResultsScreen()
        self.finals = FinalsScreen()
        self.points_setup = PointsSetupScreen()
        self.team_standings = TeamStandingsScreen()
        self.field_results = FieldResultsScreen()
        self.high_jump_results = HighJumpResultsScreen()

        self.stack.addWidget(self.dashboard)
        self.stack.addWidget(self.create_meeting)
        self.stack.addWidget(self.load_meeting)
        self.stack.addWidget(self.teams)
        self.stack.addWidget(self.user_mode)
        self.stack.addWidget(self.athletes)
        self.stack.addWidget(self.events)
        self.stack.addWidget(self.event_entry)
        self.stack.addWidget(self.results)
        self.stack.addWidget(self.finals)
        self.stack.addWidget(self.points_setup)
        self.stack.addWidget(self.team_standings)
        self.stack.addWidget(self.field_results)
        self.stack.addWidget(self.high_jump_results)

        container = QWidget()
        athletes_button = QPushButton("Athletes")
        main_layout = QHBoxLayout()

        menu = QVBoxLayout()

        home_button = QPushButton("Dashboard")
        home_button.clicked.connect(
            lambda: self.stack.setCurrentIndex(0)
        )

        menu.addWidget(home_button)
        menu.addStretch()

        main_layout.addLayout(menu)
        main_layout.addWidget(self.stack)

        container.setLayout(main_layout)

        self.setCentralWidget(container)

    def refresh_screen(self):

        self.load_events()

        self.filter_athletes_by_event()

        self.entry_list.clear()

        self.load_entries()
        self.load_heats()

        self.load_heat_choices()
        self.load_heat_for_edit()

    def open_finals(self):

        self.finals.load_events()
        self.finals.load_finals()

        self.stack.setCurrentWidget(
            self.finals
        )

# =========================
# PROGRAM START
# =========================
if __name__ == "__main__":

    init_db()

    app = QApplication(sys.argv)

    window = MainWindow()
    window.show()

    sys.exit(app.exec())
    